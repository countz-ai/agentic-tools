#!/usr/bin/env python3
"""Each recipe family's two artifacts: a one-page summary and the family's workbook
(RUN_CONTRACT.md § Every wave).

    uv run --project ${CLAUDE_PLUGIN_ROOT} python3 family_artifacts.py <run_dir>

For every family of the pinned recipe whose checks are all terminal (ok, blocked or
failed), writes under `<run_dir>/out/families/`, named by the family's number in the
recipe and its title (`05-the-books-reconciled-to-the-bank`):

- `<name>.md` — for a reader who wants to know whether the procedure worked and where
  it stands: the family's title, its status (*Complete*, *Complete, with limitations*,
  *Not completed*), then the summary each of its checks wrote
  (`checks/<check>-summary.md`, agents/worker.md § Files), or that check's conclusion
  where it wrote none. The detail is the workbook's.
- `<name>.xlsx` — the family's check tabs (`out/tabs/<check>.xlsx`), every sheet of
  each, in roster order. A family run as one check gets its tab file unchanged; several
  are merged cell by cell, formats, merges, widths and freeze panes kept, and every
  formula's cached result written back (`formula.cache_results`) so a viewer that does
  not recalculate still shows it.

A family is rebuilt only when what it is built from changed (`out/families/.built.json`
holds the fingerprints), so the preview names it again only when it moved. A family with
a check still running, or with no check (out of the run's scope, or dropped by the
plan), gets nothing. Never exits non-zero and never stops the run: a family it cannot
build is skipped with a note on stderr.
"""
from __future__ import annotations

import copy
import hashlib
import json
import pathlib
import re
import shutil
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import recipe_format  # noqa: E402  sibling: the recipe's families and their titles

OUT = ("out", "families")
STATUS = {"ok": "Complete", "blocked": "Complete, with limitations", "failed": "Not completed"}
RANK = {"ok": 0, "blocked": 1, "failed": 2}


def slug(title: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", recipe_format.plain_title(title).lower()).strip("-")
    return s[:60].rstrip("-") or "family"


def _read(path: pathlib.Path) -> bytes:
    return path.read_bytes() if path.is_file() else b""


def family_checks(run: dict) -> dict[str, list[dict]]:
    """{family id: its check rows, in roster order}; an `extract` step has no family."""
    out: dict[str, list[dict]] = {}
    for c in run.get("checks") or []:
        fam = str((c.get("params") or {}).get("family") or "").lower()
        if fam:
            out.setdefault(fam, []).append(c)
    return out


def summary_text(run_dir: pathlib.Path, run: dict, check: dict) -> str:
    """The check's own summary, its leading heading dropped; else its conclusion."""
    md = run_dir / "checks" / f"{check['id']}-summary.md"
    if md.is_file():
        text = md.read_text(encoding="utf-8").strip()
        return re.sub(r"\A#[^\n]*\n+", "", text).strip()
    rows = [d for d in run.get("dispatches") or [] if d.get("check_id") == check["id"]]
    conclusion = next((d.get("conclusion") for d in reversed(rows) if d.get("conclusion")), None)
    if check.get("status") == "failed":
        return ("We could not complete this step. The workbook carries what it reached, "
                "and the run's report states what is missing.")
    return conclusion or "The step recorded no summary."


def render_md(number: int, title: str, status: str, parts: list[str], book: str | None) -> str:
    lines = [f"# {number}. {recipe_format.plain_title(title)}", "", f"**Status:** {status}", ""]
    if len(parts) == 1:
        lines += [parts[0], ""]
    else:
        for i, p in enumerate(parts, 1):
            lines += [f"## Part {i} of {len(parts)}", "", p, ""]
    lines.append(f"The detail is in the workbook `{book}`." if book
                 else "No workbook: the family's steps wrote no schedule.")
    return "\n".join(lines).rstrip() + "\n"


def merge_tabs(tabs: list[pathlib.Path], dest: pathlib.Path) -> None:
    """Every sheet of every tab file into one workbook at `dest`, cached results kept."""
    import openpyxl
    import formula
    out = openpyxl.Workbook()
    out.remove(out.active)
    results: dict[str, dict[str, object]] = {}
    for tab in tabs:
        src = openpyxl.load_workbook(tab)
        cached = openpyxl.load_workbook(tab, data_only=True)
        for ws in src.worksheets:
            name, n = ws.title, 2
            while name in out.sheetnames:
                name, n = f"{ws.title[:28]}~{n}", n + 1
            dst = out.create_sheet(name)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None and not c.has_style:
                        continue
                    d = dst.cell(row=c.row, column=c.column, value=c.value)
                    if c.has_style:
                        d.font, d.fill = copy.copy(c.font), copy.copy(c.fill)
                        d.border, d.alignment = copy.copy(c.border), copy.copy(c.alignment)
                        d.number_format, d.protection = c.number_format, copy.copy(c.protection)
                    if c.hyperlink is not None:
                        d.hyperlink = copy.copy(c.hyperlink)
                    if c.data_type == "f":
                        v = cached[ws.title][c.coordinate].value
                        if v is not None:
                            results.setdefault(name, {})[c.coordinate] = v
            for rng in ws.merged_cells.ranges:
                dst.merge_cells(str(rng))
            for key, dim in ws.column_dimensions.items():
                dst.column_dimensions[key].width = dim.width
                dst.column_dimensions[key].hidden = dim.hidden
            for key, dim in ws.row_dimensions.items():
                if dim.height is not None:
                    dst.row_dimensions[key].height = dim.height
            dst.freeze_panes = ws.freeze_panes
            dst.sheet_view.showGridLines = ws.sheet_view.showGridLines
            if ws.sheet_properties.tabColor is not None:
                dst.sheet_properties.tabColor = copy.copy(ws.sheet_properties.tabColor)
    tmp = dest.with_name(dest.name + ".tmp.xlsx")
    out.save(tmp)
    formula.cache_results(tmp, results)
    tmp.replace(dest)


def build(run_dir: pathlib.Path) -> list[pathlib.Path]:
    run = json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    recipe = pathlib.Path(str((run.get("plan") or {}).get("recipe") or ""))
    if not recipe.is_file():
        return []
    fams = recipe_format.families(recipe.read_text(encoding="utf-8"))
    by_family = family_checks(run)
    out_dir = run_dir.joinpath(*OUT)
    state_path = out_dir / ".built.json"
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        state = {}
    written: list[pathlib.Path] = []
    for number, fam in enumerate(fams, 1):
        checks = by_family.get(fam["id"]) or []
        if not checks or any(c.get("status") not in STATUS for c in checks):
            continue                    # out of scope, dropped, or a step still running
        name = f"{number:02d}-{slug(fam['title'])}"
        tabs = [p for c in checks if (p := run_dir / "out" / "tabs" / f"{c['id']}.xlsx").is_file()]
        h = hashlib.sha256()
        for c in checks:
            h.update(json.dumps([c["id"], c.get("status"), c.get("seq")]).encode())
            h.update(_read(run_dir / "checks" / f"{c['id']}-summary.md"))
        for t in tabs:
            h.update(_read(t))
        fingerprint = h.hexdigest()[:16]
        if state.get(name) == fingerprint and (out_dir / f"{name}.md").is_file():
            continue
        try:
            out_dir.mkdir(parents=True, exist_ok=True)
            book = None
            if tabs:
                dest = out_dir / f"{name}.xlsx"
                if len(tabs) == 1:
                    shutil.copyfile(tabs[0], dest)
                else:
                    merge_tabs(tabs, dest)
                book = dest.name
                written.append(dest)
            worst = max((c["status"] for c in checks), key=RANK.__getitem__)
            md = out_dir / f"{name}.md"
            tmp = md.with_name(md.name + ".tmp")
            tmp.write_text(render_md(number, fam["title"], STATUS[worst],
                                     [summary_text(run_dir, run, c) for c in checks], book),
                           encoding="utf-8")
            tmp.replace(md)
            written.append(md)
            state[name] = fingerprint
        except Exception as exc:  # noqa: BLE001 - one family never stops the others
            print(f"family_artifacts: {name} skipped ({type(exc).__name__}: {exc})",
                  file=sys.stderr)
    if written:
        tmp = state_path.with_name(state_path.name + ".tmp")
        tmp.write_text(json.dumps(state, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        tmp.replace(state_path)
    return written


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: family_artifacts.py <run_dir>", file=sys.stderr)
        return 0
    run_dir = pathlib.Path(sys.argv[1]).resolve()
    try:
        written = build(run_dir)
    except Exception as exc:  # noqa: BLE001 - the run goes on without them
        print(f"family_artifacts: skipped ({type(exc).__name__}: {exc})", file=sys.stderr)
        return 0
    print(f"family_artifacts: {len(written)} file(s) written under {'/'.join(OUT)}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
