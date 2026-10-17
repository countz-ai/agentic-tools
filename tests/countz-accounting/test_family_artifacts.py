#!/usr/bin/env python3
"""Self-test for scripts/family_artifacts.py: each finished family's summary and workbook.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_family_artifacts.py`
from the repository root. Covers the summary's title, status and text (a step's own
summary, or its conclusion), the workbook (a one-check family's tab copied, a split
family's tabs merged with formulas and their cached results kept), a family still
running left alone, and an unchanged family not rewritten.
"""
from __future__ import annotations

import json
import pathlib
import sys
import tempfile

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import family_artifacts  # noqa: E402
import formula  # noqa: E402
import openpyxl  # noqa: E402

RECIPE = """---
name: demo
objective: demo
---

## The families

### D0 — the population, drawn from the balance sheet (kind `analysis`, one check)

### D1 — the books agree internally (kind `tieout`, per GL account)

### D2 — the `bank_ledger` reconciled to the bank (kind `recon`, per bank account)

### D3 — the walk (kind `analysis`, one check)
"""


def tab(path: pathlib.Path, sheet: str, value: int) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"], ws["A2"] = value, "=A1+1"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws.column_dimensions["A"].width = 30
    wb.save(path)
    formula.cache_results(path, {sheet: {"A2": value + 1}})


def main() -> int:
    bad: list[str] = []

    def expect(cond, what):
        if not cond:
            bad.append(what)

    with tempfile.TemporaryDirectory() as td:
        rd = pathlib.Path(td)
        for d in ("checks", "out/tabs", "recipes"):
            (rd / d).mkdir(parents=True)
        (rd / "recipes" / "demo.md").write_text(RECIPE)
        checks = [
            {"id": "d0_population", "params": {"family": "d0"}, "status": "ok", "seq": 2},
            {"id": "d2_large", "params": {"family": "d2"}, "status": "ok", "seq": 3},
            {"id": "d2_rest", "params": {"family": "d2"}, "status": "blocked", "seq": 4},
            {"id": "d3_walk", "params": {"family": "d3"}, "status": "pending", "seq": 5},
            {"id": "extract_gl", "params": {}, "status": "ok", "seq": 1},
        ]
        (rd / "run.json").write_text(json.dumps({
            "plan": {"recipe": str(rd / "recipes" / "demo.md")}, "checks": checks,
            "dispatches": [{"check_id": "d2_rest", "conclusion": "We reconciled 9 of 10 accounts."}]}))
        (rd / "checks" / "d0_population-summary.md").write_text(
            "# Population\n\nWe drew the population from the balance sheet: 4 accounts.\n")
        (rd / "checks" / "d2_large-summary.md").write_text("We reconciled the operating account.\n")
        tab(rd / "out" / "tabs" / "d0_population.xlsx", "d0 Population", 1)
        tab(rd / "out" / "tabs" / "d2_large.xlsx", "d2_large", 10)
        tab(rd / "out" / "tabs" / "d2_rest.xlsx", "d2_rest", 20)

        written = family_artifacts.build(rd)
        fam = rd / "out" / "families"
        names = sorted(p.name for p in written)
        expect(names == ["01-the-population-drawn-from-the-balance-sheet.md",
                         "01-the-population-drawn-from-the-balance-sheet.xlsx",
                         "03-the-bank-ledger-reconciled-to-the-bank.md",
                         "03-the-bank-ledger-reconciled-to-the-bank.xlsx"],
               f"written: {names}")
        expect(not list(fam.glob("04-*")), "a family with a step still running is left alone")

        one = (fam / "01-the-population-drawn-from-the-balance-sheet.md").read_text()
        expect(one.startswith("# 1. The population, drawn from the balance sheet\n"), one)
        expect("**Status:** Complete" in one and "4 accounts" in one, one)
        expect("# Population" not in one, "the step's own heading is dropped")
        expect((fam / "01-the-population-drawn-from-the-balance-sheet.xlsx").read_bytes()
               == (rd / "out" / "tabs" / "d0_population.xlsx").read_bytes(),
               "a one-check family's workbook is its tab, unchanged")

        two = (fam / "03-the-bank-ledger-reconciled-to-the-bank.md").read_text()
        expect("**Status:** Complete, with limitations" in two, "the worst status stands")
        expect("## Part 1 of 2" in two and "operating account" in two
               and "9 of 10 accounts" in two, f"each part, summary or conclusion:\n{two}")
        expect("`" not in two.split("\n", 1)[0] and "_" not in two.split("\n", 1)[0],
               "the title carries no markup")
        book = fam / "03-the-bank-ledger-reconciled-to-the-bank.xlsx"
        wb = openpyxl.load_workbook(book)
        expect(wb.sheetnames == ["d2_large", "d2_rest"], f"sheets: {wb.sheetnames}")
        expect(wb["d2_rest"]["A2"].value == "=A1+1", "the formula is kept")
        expect(wb["d2_rest"]["A1"].font.bold and wb["d2_rest"].column_dimensions["A"].width == 30,
               "formats and widths are kept")
        cached = openpyxl.load_workbook(book, data_only=True)
        expect(cached["d2_rest"]["A2"].value == 21, "the formula's cached result is written back")

        expect(family_artifacts.build(rd) == [], "an unchanged family is not rewritten")
        (rd / "checks" / "d2_large-summary.md").write_text("We reconciled it again.\n")
        again = sorted(p.name for p in family_artifacts.build(rd))
        expect(again == ["03-the-bank-ledger-reconciled-to-the-bank.md",
                         "03-the-bank-ledger-reconciled-to-the-bank.xlsx"],
               f"a changed summary rebuilds its family only: {again}")

    for b in bad:
        print(f"FAIL: {b}")
    print(f"test_family_artifacts: {'ok' if not bad else f'{len(bad)} failure(s)'}")
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
