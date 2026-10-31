#!/usr/bin/env python3
"""Self-test for scripts/reword.py.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_reword.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import hashlib
import pathlib
import sys
import tempfile

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from reword import Refused, apply  # noqa: E402


def build(path: pathlib.Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "r8 Bridge"
    a["B5"] = "DSO as measured (R5)"
    a["C5"] = 60.9
    a["B6"] = "Customer C10007, days beyond terms"
    a["C6"] = 12.0
    a["C7"] = "=C5+C6"
    a["D5"] = "F.r5.dso_3m.2025-06"
    wb.create_sheet("r5 Aging and DSO")["B5"] = "Headline DSO"
    b = wb.create_sheet("Sources")
    b["B2"] = "DSO as measured (R5)"
    b["B3"] = "June 30, 2025, days of DSO"
    wb.save(path)


def cells(path: pathlib.Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    return {(ws.title, c.coordinate): c.value for ws in wb for row in ws.iter_rows() for c in row
            if c.value is not None}


def main() -> int:
    bad = []
    with tempfile.TemporaryDirectory() as td:
        wb = pathlib.Path(td) / "workbook.xlsx"
        build(wb)
        before = cells(wb)
        # a step token dropped on one tab: the other tab's copy of the same string stays
        log = apply(wb, [{"tab": "r8 Bridge", "old": " (R5)", "new": "", "why": "check_report gate 7"},
                         {"tab": "*", "old": ", days", "new": ": days", "why": "check_prose"}])
        after = cells(wb)
        want = dict(before)
        want[("r8 Bridge", "B5")] = "DSO as measured"
        want[("r8 Bridge", "B6")] = "Customer C10007: days beyond terms"
        want[("Sources", "B3")] = "June 30, 2025: days of DSO"
        if after != want:
            bad.append(f"applied edits: {sorted(set(after.items()) ^ set(want.items()))}")
        if [len(r["cells"]) for r in log] != [1, 2]:
            bad.append(f"log cells: {log}")
        refusals = [
            ([{"tab": "r8 Bridge", "old": "60.9", "new": "61.0", "why": "x"}], "a number cell"),
            ([{"tab": "r8 Bridge", "old": "C10007", "new": "C10008", "why": "x"}], "a moved number"),
            ([{"tab": "r8 Bridge", "old": "F.r5.dso_3m.2025-06", "new": "DSO", "why": "x"}], "a dropped id"),
            ([{"tab": "Nope", "old": "a", "new": "b", "why": "x"}], "a missing tab"),
            ([{"tab": "r8 Bridge", "cell": "C7", "old": "C5", "new": "C6", "why": "x"}], "a formula"),
            ([{"tab": "*", "old": "measured", "new": "stated", "why": "x"},
              {"tab": "r8 Bridge", "old": "nothing here", "new": "y", "why": "x"}], "a second edit matching nothing"),
        ]
        for edits, what in refusals:
            h = hashlib.sha256(wb.read_bytes()).hexdigest()
            try:
                apply(wb, edits)
                bad.append(f"{what}: applied, want refused")
            except Refused:
                if hashlib.sha256(wb.read_bytes()).hexdigest() != h:
                    bad.append(f"{what}: refused but the workbook changed")
    for b in bad:
        print(f"reword: {b}")
    print("reword: ok" if not bad else "reword: self-check FAILED")
    return 0 if not bad else 1


if __name__ == "__main__":
    raise SystemExit(main())
