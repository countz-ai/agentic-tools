#!/usr/bin/env python3
"""Self-test for the report deck: build_report.py renders a deck from a spec-conformant
workbook, check_report.py passes it, and each refuses what it exists to refuse.

Builds, in a temporary directory, a run whose workbook is written to WORKBOOK.md +
WORKBOOK_STYLE.md (the kit of § 8) and passes link_workbook.py + check_workbook.py, so
the deck is tested over a workbook the plugin itself would seal. Then:

  1. a good report.yaml builds, and the deck passes check_report.py; on the deck as
     stored, a page's headline and its message are separate shapes, a page written with
     no message carries none, and every table reads its first column left and every
     other column right;
  2. a figure typed into a sentence builds (the builder cannot know) and is refused by
     the gate, named;
  3. a reference the workbook cannot resolve is refused by the builder, named;
  4. a deck whose figure was edited after the build is refused by the gate;
  4a. a chart is DRAWN, not placed as a chart part — the deck holds no `ppt/charts/`
     part — and its plotted values are still audited: a value edited on the shape that
     carries it is refused by the gate, named;
  5. a sentence where a headline belongs — a title over the cap, or ending in a full
     stop — is refused by the builder, named;
  6. a fragment where a sentence belongs — a text block with no full stop — is refused
     by the builder, named;
  7. a cover title carrying the period is refused by the builder, named, and refused
     again by the gate on a deck built before the rule; a subtitle repeating the title is
     refused; a title naming the company builds — that judgment is the critic's
     (skills/check-review), a name in any language being no pattern for code;
  8. a zero in a sentence reads `$0` — the en dash is the table's zero, and a sentence
     that trails off in one states nothing;
  9. a schedule condensed past its own arithmetic — a derived line shown over rows that
     do not make it — is refused by the builder, naming the rows it drops;
 10. the deck's own number conventions (REPORT.md § 4): a dollar figure in a sentence or
     a stat tile is scaled and rounded (`$8.4M`), a schedule declared `scale: thousands`
     is titled `(… $ in thousands)` and divided, and the gate still finds the workbook cell
     behind each;
 11. a schedule declared `dense: true` is set at the dense size, and `where:` with
     `through:` keeps a walk's mechanics and its ruled rows and ends it at the closing
     line — the information line under it is not on the deck;
 13. the opening (REPORT.md § 1) is held by the gate: an executive summary page that is
     not first, or carries no figure block, a key-metrics page not headed as the recipe
     declares, and a first schedule two pages after it are each refused, named.
 12. the recipe's schedule (RECIPE_FORMAT.md § Report) is held by the gate: a deck whose
     walk is trimmed to `max_rows`, or carries no table from the walk's tab, is refused
     naming the schedule and the rows it lacks.
 14. money is the book's currency (scripts/style.py): a EUR book reads `€8.4M` and
     `(€ in thousands)` and passes; a typed `€1,234,567` is refused, not skipped;
 16. on a run that declares its periods (a September year end), a schedule's period
     columns are the plan's and a missing one is named; a check whose periods
     scripts/periods.py refuses is named, not dropped.

Run by check-plugin.py (8q) as `uv run --project <plugin> python3 tests/countz-accounting/report-selftest.py`
from the repository root; it lives outside the plugin so it never ships.
Exit 0 when every case holds, 1 otherwise, naming the case.
"""
from __future__ import annotations

import html
import json
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# The plugin under test: <repo>/countz-accounting/scripts, from <repo>/tests/countz-accounting.
SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"

# --- the kit: scripts/wbkit.py (WORKBOOK_STYLE.md § 9 + WORKBOOK.md § 7) ------------
sys.path.insert(0, str(SCRIPTS))
from wbkit import (ACCENT, BAND, FMT_AMOUNT, FMT_TEXT, MIST, S, SLATE, TINT, next_block,  # noqa: E402,F401
                   WIDTH, amount, band, finish, fit_rows, grid, header, ident, section,
                   status, text)


# --- the fixture run ----------------------------------------------------------------
PERIODS = ["FY2023", "FY2024", "LTM Jul 2025"]
SLUGS = ["fy2023", "fy2024", "ltm_2025-07"]
SUB = "Acme Corp · FY2023–LTM Jul 2025 · accrual · USD whole dollars · tolerance 1"
BRIDGE = [
    ("EBIT", [5_610_000, 6_120_000, 6_890_000], "Body"),
    ("+ D&A", [1_040_000, 1_120_000, 1_180_000], "Body"),
    ("= Reported EBITDA", [6_650_000, 7_240_000, 8_070_000], "Subtotal"),
    ("Non-recurring, supported", [420_000, 310_000, 264_000], "Body"),
    ("Normalization, supported", [-180_000, -95_000, 70_000], "Body"),
    ("Owner compensation, no adjustment taken", [0, 0, 0], "Body"),
    ("= Diligence adjusted EBITDA", [6_890_000, 7_455_000, 8_404_000], "Total"),
    ("Management adjusted EBITDA (information)", [7_800_000, 8_500_000, 9_650_000], "Body"),
]
POSITION = ("Diligence-adjusted EBITDA is 8,404,000 for LTM Jul 2025, 1,246,000 below "
            "management's adjusted figure.")


def make_run(rd: pathlib.Path, cur: str = "USD", q6_params: dict | None = None) -> None:
    """The fixture run. `cur` is the reporting currency its basis line states (the deck's
    money symbol and scale headers follow it); `q6_params` extends the bridge check's
    params — `columns` and `fiscal_year_end` declare the plan's periods."""
    basis = "USD whole dollars" if cur == "USD" else f"{cur} whole units"
    sub = SUB.replace("USD whole dollars", basis)
    figures: dict[str, tuple] = {}
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Exec Summary")
    band(ws, "Exec Summary · quality of earnings", f"Acme Corp · FY2023–LTM Jul 2025 · accrual · {basis}",
         summary=POSITION)
    ws.freeze_panes = "B4"
    text(ws.cell(row=4, column=2), "EBITDA bridge · from: q6 EBITDA bridge", "Section")
    header(ws, 6, ["line"] + PERIODS, ["description", "period", "period", "period"], primary=False)
    r = 6
    for label, vals, st in BRIDGE:
        r += 1
        text(ws.cell(row=r, column=2), label, st)
        for j, v in enumerate(vals):
            amount(ws.cell(row=r, column=3 + j), v, style=st if st != "Body" else None)
    grid(ws, 6, r, 2, 5)
    fit_rows(ws, 4)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT

    ws = wb.create_sheet("q6 EBITDA bridge")
    band(ws, "q6 · EBIT walks to diligence adjusted EBITDA", sub,
         summary="Diligence adjusted EBITDA is 8,404,000 LTM Jul 2025; management's figure is 1,246,000 higher.")
    header(ws, 4, ["id", "line", *PERIODS, "verdict", "note"],
           ["id", "description", "period", "period", "period", "status", "note"])
    r = 4
    for label, vals, st in BRIDGE:
        r += 1
        slug = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
        ident(ws.cell(row=r, column=2), f"F.q6.{slug}.ltm_2025-07")
        text(ws.cell(row=r, column=3), label, st)
        for j, v in enumerate(vals):
            amount(ws.cell(row=r, column=4 + j), v, style=st if st != "Body" else None)
            figures[f"F.q6.{slug}.{SLUGS[j]}"] = (label, v, "q6_bridge")
        status(ws.cell(row=r, column=7), "supported")
        text(ws.cell(row=r, column=8), "walk mechanics" if st != "Body" else "the rung's supported subtotal")
    last = r
    r = section(ws, next_block(r), "Exceptions")
    header(ws, r, ["id", "item", "amount", "owner", "what would clear it"],
           ["id", "description", "amount", "status", "note"], primary=False)
    hx = r
    r += 1
    ident(ws.cell(row=r, column=2), "X.q6.mgmt_residual")
    text(ws.cell(row=r, column=3), "Management adjusted EBITDA does not reconcile to its schedule")
    amount(ws.cell(row=r, column=4), 54_000)
    text(ws.cell(row=r, column=5), "CFO")
    text(ws.cell(row=r, column=6), "Management's revised schedule")
    grid(ws, hx, r, 2, 6)
    r = next_block(r)
    section(ws, r, "Notes")
    r += 1
    ident(ws.cell(row=r, column=2), "F.q6.ebit.ltm_2025-07")
    text(ws.cell(row=r, column=3), "Source: the income statement and the general ledger, as tied at q1.", "Note")
    r += 2
    section(ws, r, "To reperform")
    text(ws.cell(row=r + 1, column=2), "1. Read pl.xlsx · IS · rows 5-40, column E; sum to EBIT; add D&A.")
    finish(ws, last)

    ws = wb.create_sheet("Basis of Preparation")
    band(ws, "Basis of Preparation", f"Acme Corp · FY2023–LTM Jul 2025 · accrual · {basis}")
    header(ws, 4, ["source", "what it is", "class", "periods"], ["description", "description", "status", "description"])
    for i, row in enumerate([("gl", "General ledger export, gl.csv", "system-of-record", "Jan 2023 – Jul 2025"),
                             ("pl", "Income statements as presented, pl.xlsx", "management-prepared", "FY2023, FY2024, monthly")], 5):
        for j, v in enumerate(row):
            text(ws.cell(row=i, column=2 + j), v)
    grid(ws, 4, 6, 2, 5)
    section(ws, next_block(6), "Procedures not performed")
    text(ws.cell(row=10, column=2), "The capex bridge was dropped: the room carries no capitalized-cost accounts.")
    hr = section(ws, 12, "How to read this workbook")
    header(ws, hr, ["prefix", "meaning"], ["status", "description"], primary=False)
    text(ws.cell(row=hr + 1, column=2), "F.")
    text(ws.cell(row=hr + 1, column=3), "a figure — resolves on its Sources row")
    grid(ws, hr, hr + 1, 2, 3)
    finish(ws, 6)

    ws = wb.create_sheet("q1 FY2023 statements")
    band(ws, "q1 · The FY2023 income statement agrees with the trial balance", sub,
         summary="Net income and EBIT tie to the TB within tolerance; no exceptions.")
    header(ws, 4, ["id", "what is agreed", "side A", "source A", "amount A", "side B", "source B", "amount B", "difference", "status"],
           ["id", "description", "description", "description", "amount", "description", "description", "amount", "amount", "status"])
    r = 4
    for slug, what, a in [("net_income", "Net income FY2023", 3_912_000), ("ebit", "EBIT FY2023", 5_610_000)]:
        r += 1
        figures[f"F.q1.{slug}.fy2023"] = (what, a, "q1_fy2023")
        ident(ws.cell(row=r, column=2), f"F.q1.{slug}.fy2023")
        text(ws.cell(row=r, column=3), what)
        text(ws.cell(row=r, column=4), "income statement")
        text(ws.cell(row=r, column=5), "pl.xlsx · IS · C12")
        amount(ws.cell(row=r, column=6), a, hard_input=True)
        text(ws.cell(row=r, column=7), "trial balance")
        text(ws.cell(row=r, column=8), "tb.xlsx · TB · P&L close")
        amount(ws.cell(row=r, column=9), a, hard_input=True)
        amount(ws.cell(row=r, column=10), 0)
        status(ws.cell(row=r, column=11), "tied")
    last = r
    r = next_block(r)
    section(ws, r, "Notes")
    text(ws.cell(row=r + 1, column=2), "Population: every line of the income statement, 31 of 31.", "Note")
    r += 3
    section(ws, r, "To reperform")
    text(ws.cell(row=r + 1, column=2), "1. Read pl.xlsx · IS · column C; read tb.xlsx · TB; difference = A − B.")
    finish(ws, last)

    ws = wb.create_sheet("Coverage")
    band(ws, "Coverage", "Acme Corp · one row per rostered check")
    header(ws, 4, ["token · title", "kind", "step", "status", "what was examined", "what was not examined, and why"],
           ["description", "status", "status", "status", "note", "note"])
    for i, row in enumerate([("q1 · FY2023 statements", "tieout", "q1_fy2023", "complete", "31 income-statement lines, FY2023", "—"),
                             ("q6 · EBITDA bridge", "analysis", "q6_bridge", "complete", "the bridge over 3 periods", "the capex bridge")], 5):
        for j, v in enumerate(row):
            text(ws.cell(row=i, column=2 + j), v)
    finish(ws, 6)

    ws = wb.create_sheet("Open Items")
    band(ws, "Open Items", "Acme Corp · review calls, questions for management, data requests")
    header(ws, 4, ["id", "matter", "size", "what closes it", "owner", "raised by"],
           ["id", "description", "amount", "note", "status", "id"])
    ident(ws.cell(row=5, column=2), "Q.q6.mgmt_residual")
    text(ws.cell(row=5, column=3), "Q.q6.mgmt_residual. Which schedule version produced the published figure?")
    amount(ws.cell(row=5, column=4), 54_000)
    text(ws.cell(row=5, column=5), "The schedule version")
    text(ws.cell(row=5, column=6), "CFO")
    text(ws.cell(row=5, column=7), "")
    grid(ws, 4, 5, 2, 7)
    hr = section(ws, next_block(5), "Data requests")
    header(ws, hr, ["id", "matter", "size", "what closes it", "owner", "raised by"],
           ["id", "description", "amount", "note", "status", "id"], primary=False)
    ident(ws.cell(row=hr + 1, column=2), "D.q6.capex")
    text(ws.cell(row=hr + 1, column=3), "D.q6.capex. The capitalized-cost accounts, for the capex bridge.")
    amount(ws.cell(row=hr + 1, column=4), 0)
    text(ws.cell(row=hr + 1, column=5), "The rollforward")
    text(ws.cell(row=hr + 1, column=6), "Controller")
    text(ws.cell(row=hr + 1, column=7), "")
    grid(ws, hr, hr + 1, 2, 7)
    finish(ws, 5)
    ws.sheet_properties.tabColor = "8A5A00"

    ws = wb.create_sheet("Sources")
    band(ws, "Sources", f"Acme Corp · {len(figures)} figure rows")
    header(ws, 4, ["id", "label", "value", "unit", "disposition", "expression", "inputs", "root source", "To reperform", "population", "caveats and carry-forwards", "from check"],
           ["id_ledger", "description", "amount", "status", "status", "note", "id", "id", "note", "note", "note", "status"])
    r = 4
    for fid, (label, value, check) in figures.items():
        r += 1
        ident(ws.cell(row=r, column=2), fid)
        text(ws.cell(row=r, column=3), label)
        amount(ws.cell(row=r, column=4), value)
        text(ws.cell(row=r, column=5), cur)
        text(ws.cell(row=r, column=6), "measured")
        text(ws.cell(row=r, column=7), f"as read ({label})")
        ident(ws.cell(row=r, column=8), "E.q1.pl")
        ident(ws.cell(row=r, column=9), "E.q1.pl")
        text(ws.cell(row=r, column=10), "pl.xlsx · IS · rows 5-40 · col E · sum")
        text(ws.cell(row=r, column=11), "whole")
        text(ws.cell(row=r, column=12), "")
        text(ws.cell(row=r, column=13), check)
    finish(ws, r, ledger=True)

    ws = wb.create_sheet("Evidence")
    band(ws, "Evidence", "Acme Corp · 1 room read")
    header(ws, 4, ["id", "kind", "file", "source", "role", "sheet", "anchor", "rows", "columns", "filter", "row count", "control total", "value basis", "read date", "note"],
           ["id_ledger", "status", "description", "status", "status", "status", "status", "status", "note", "note", "amount", "amount", "status", "period", "note"])
    ident(ws.cell(row=5, column=2), "E.q1.pl")
    for j, v in enumerate(["range", "statements/pl.xlsx", "pl", "management_prepared", "IS", "B4", "5-40", "B:E", ""], start=3):
        text(ws.cell(row=5, column=j), v)
    amount(ws.cell(row=5, column=12), 36, fmt="#,##0")
    amount(ws.cell(row=5, column=13), 41_250_000)
    text(ws.cell(row=5, column=14), "as_stated")
    text(ws.cell(row=5, column=15), "2 Sep 2025")
    text(ws.cell(row=5, column=16), "income statement as presented")
    finish(ws, 5, ledger=True)

    staging = rd / "out" / ".staging"
    staging.mkdir(parents=True)
    (rd / "workpapers").mkdir()
    wb.save(staging / "workbook.xlsx")
    by_check: dict[str, list[str]] = {}
    for fid, (label, value, check) in figures.items():
        by_check.setdefault(check, []).append(
            f"- id: {fid}\n  label: {label}\n  value: {value}\n  unit: {cur.lower()}\n  expression: as read\n"
            f"  inputs:\n    - {{role: pl, source_type: room_file, citation_id: E.q1.pl}}\n")
    for check, lines in by_check.items():
        (rd / "workpapers" / f"figures-{check}.yaml").write_text("".join(lines), encoding="utf-8")
    (rd / "workpapers" / "evidence-plan.yaml").write_text(
        "- id: E.q1.pl\n  kind: range\n  file: \"statements/pl.xlsx\"\n  control_total: 41250000\n  row_count: 36\n", encoding="utf-8")
    recipe = rd / "QOE.md"
    recipe.write_text(RECIPE, encoding="utf-8")
    (rd / "run.json").write_text(json.dumps({
        "schema": "countz-accounting/run@1", "run_id": "qoe-acme-corp.20250903-101500",
        "goal": "quality-of-earnings", "degraded": False,
        "inputs": {"run_dir": str(rd), "skill": "qoe", "company": "Acme Corp", "params": {}},
        "checks": [{"id": "q1_fy2023", "kind": "tieout", "params": {"family": "q1"}},
                   {"id": "q6_bridge", "kind": "analysis", "params": {"family": "q6", **(q6_params or {})}}],
        "plan": {"recipe": str(recipe)}}), encoding="utf-8")


# The recipe the fixture run pins: its `## Report` declares the walk the deck must carry
# (RECIPE_FORMAT.md § Report), which case 12 holds the gate to.
RECIPE = """---
name: quality-of-earnings
objective: o
headline: q6
lead: [q6]
---
# r

## Exec summary

x

## Report

```json
{"metrics": {"title": "Adjusted EBITDA"},
 "schedules": [
  {"title": "EBITDA walk", "from": "q6", "columns": ["line", "verdict"], "periods": "all",
   "where": {"verdict": "supported"}, "through": "= Diligence adjusted EBITDA", "dense": true}
]}
```

## What the plan notes rather than checks

y
"""

# The spec as REPORT.md § 2 has it: a title is a headline — the subject of a page of
# figures (`EBITDA bridge`), the conclusion of a page that argues one (`The bridge foots at
# every rung`) — and the sentence, where a page needs one, is `message`.
GOOD_SPEC = """schema: countz-accounting/report@1
title: Quality of earnings review
sections:
  - title: Executive summary
    pages:
      - title: Executive summary
        message: "{Exec Summary!B3}"
        blocks:
          - stats:
              - {label: Reported EBITDA · LTM Jul 2025, value: "{q6 | = Reported EBITDA | LTM Jul 2025 | $}"}
              - {label: Diligence adjusted EBITDA · LTM Jul 2025, value: "{q6 | = Diligence adjusted EBITDA | LTM Jul 2025 | $}"}
          - heading: What this rests on
          - bullets: ["The FY2023 income statement ties to the trial balance with no exception."]
      - title: Adjusted EBITDA
        blocks:
          - table: {from: q6, rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025], title: "EBITDA bridge, USD"}
          - chart: {type: column, from: q6, rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [FY2023, FY2024, LTM Jul 2025]}
  - title: The bridge
    pages:
      - title: EBITDA bridge by line
        blocks:
          - table: {from: q6, where: {verdict: supported}, through: "= Diligence adjusted EBITDA", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict], dense: true}
      - title: The bridge foots at every rung
        blocks:
          - table: {from: q6, max_rows: 5}
          - table: {from: q6, block: Exceptions}
  - title: Appendix
    pages:
      - title: Coverage
        message: "Two checks ran; the capex bridge was dropped because the room carries no capitalized-cost accounts."
        blocks:
          - text: "Every rostered check is listed with what it examined and what it did not."
          - text: "No owner-compensation adjustment was taken: the amount is {q6 | Owner compensation, no adjustment taken | LTM Jul 2025 | $}."
          - table: {from: Coverage, columns: ["token · title", status, "what was examined", "what was not examined, and why"]}
          - lines: {from: Basis of Preparation, block: Procedures not performed}
"""
SENTENCE_TITLE = ("Reported EBITDA walks to diligence adjusted EBITDA through supported adjustments "
                  "at every rung of the bridge.")


def stored_slides(deck: pathlib.Path) -> list[str]:
    """The slide XML parts of a deck, in presentation order (the gate's own reading)."""
    sys.path.insert(0, str(SCRIPTS))
    from check_report import slide_parts  # noqa: E402
    with zipfile.ZipFile(deck) as z:
        return [z.read(p).decode("utf-8", "replace") for p in slide_parts(z)]


def structure(deck: pathlib.Path) -> list[str]:
    """What case 1 holds on the deck as stored: headline and message are separate
    named shapes, a page written with no message carries none, and every table reads
    its first column left and the rest right."""
    out: list[str] = []
    slides = stored_slides(deck)
    pages = [x for x in slides if 'name="page:' in x]
    with_msg = [x for x in pages if 'name="message"' in x]
    if len(pages) < 4:
        out.append(f"structure: expected four pages, found {len(pages)}")
    if len(with_msg) != 2:
        out.append(f"structure: two pages carry a message shape, found {len(with_msg)}")
    if not any('name="title"' in x and 'name="message"' not in x for x in pages):
        out.append("structure: a page written without a message must carry no message shape")
    titles = [html.unescape(m) for x in pages for m in re.findall(r'name="title".*?<a:t>(.*?)</a:t>', x, re.S)]
    if "Executive summary" not in titles or "Adjusted EBITDA" not in titles:
        out.append(f"structure: titles are the headlines as written, found {titles}")
    tables = 0
    for x in pages:
        for tbl in re.findall(r"<a:tbl>.*?</a:tbl>", x, re.S):
            rows = re.findall(r"<a:tr\b.*?</a:tr>", tbl, re.S)
            for tr in rows:
                cells = re.findall(r"<a:tc\b.*?</a:tc>", tr, re.S)
                if len(cells) < 2:
                    continue
                tables += 1
                algn = [(re.search(r'algn="(\w+)"', c) or [None, None])[1] for c in cells]
                if algn[0] != "l" or any(a != "r" for a in algn[1:]):
                    out.append(f"structure: a table row reads {algn}; the first column is left, the rest right")
                    break
    if not tables:
        out.append("structure: no table row with two or more columns was checked")
    return out


def first_rows(slide_xml: str) -> list[str]:
    """The header cells of every table on a slide, as text."""
    out = []
    for gf in re.findall(r"<p:graphicFrame\b.*?</p:graphicFrame>", slide_xml, re.S):
        tr = re.search(r"<a:tr\b.*?</a:tr>", gf, re.S)
        if tr:
            out += [html.unescape("".join(re.findall(r"<a:t>(.*?)</a:t>", tc, re.S)))
                    for tc in re.findall(r"<a:tc\b.*?</a:tc>", tr.group(0), re.S)]
    return out


def run(cmd: list[str]) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True)


def main() -> int:
    fails: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        rd = pathlib.Path(td) / "run"
        make_run(rd)
        py = sys.executable
        wb = rd / "out" / ".staging" / "workbook.xlsx"
        r = run([py, str(SCRIPTS / "link_workbook.py"), str(wb)])
        if r.returncode != 0:
            fails.append(f"fixture: link_workbook exited {r.returncode}: {(r.stdout + r.stderr)[-200:]}")
        r = run([py, str(SCRIPTS / "check_workbook.py"), str(wb), "--run-dir", str(rd)])
        if r.returncode != 0:
            fails.append(f"fixture: the workbook must pass check_workbook.py (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[-400:]}")
        spec = rd / "out" / ".staging" / "report.yaml"
        deck = rd / "out" / ".staging" / "report.pptx"
        build = [py, str(SCRIPTS / "build_report.py"), str(rd)]
        gate = [py, str(SCRIPTS / "check_report.py"), str(deck), "--run-dir", str(rd)]

        # 1. good spec: builds and passes
        spec.write_text(GOOD_SPEC, encoding="utf-8")
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"1: a good spec must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-400:]}")
        else:
            g = run(gate)
            if g.returncode != 0:
                fails.append(f"1: the built deck must pass check_report.py (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[-400:]}")
            fails.extend(f"1: {s}" for s in structure(deck))
            good = deck.read_bytes()

            # 4. a figure edited after the build
            tampered = deck.with_name("tampered.pptx")
            hits = 0
            with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(tampered, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("ppt/slides/slide") and b'name="page:' in data \
                            and not hits:
                        # edit a figure inside a text run, not the slide's name attribute
                        edited = re.sub(rb"(<a:t>[^<]*?)8,404,000", rb"\g<1>8,404,100", data, count=1)
                        if edited != data:
                            data, hits = edited, hits + 1
                    zout.writestr(item, data)
            g = run([py, str(SCRIPTS / "check_report.py"), str(tampered), "--workbook", str(wb), "--run-dir", str(rd)])
            if not hits or g.returncode != 1 or "8,404,100" not in g.stdout:
                fails.append(f"4: an edited figure must be refused, named (exit {g.returncode}, edited {hits}): "
                             f"{(g.stdout + g.stderr).strip()[:300]}")
            deck.write_bytes(good)

            # 4a. the chart is drawn, and its values are still audited. A chart part
            # rasterises on import in some viewers, so the builder draws the chart as
            # shapes and carries each plotted value on its shape's name; the audit then
            # rests on that name, and this case is what stops it being lost silently.
            with zipfile.ZipFile(deck) as z:
                parts = z.namelist()
            if any(n.startswith("ppt/charts/") for n in parts):
                fails.append("4a: the deck holds a chart part; a chart is drawn as shapes")
            drawn = deck.with_name("drawn.pptx")
            hits = 0
            with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(drawn, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("ppt/slides/slide") and not hits:
                        edited = re.sub(rb'name="chartval:([^:"]+):[-\d.eE+]+"',
                                        rb'name="chartval:\g<1>:4242.5"', data, count=1)
                        if edited != data:
                            data, hits = edited, hits + 1
                    zout.writestr(item, data)
            g = run([py, str(SCRIPTS / "check_report.py"), str(drawn), "--workbook", str(wb), "--run-dir", str(rd)])
            if not hits:
                fails.append("4a: no chartval shape on the deck — a drawn chart carries its values "
                             "on its shape names, and the gate audits them from there")
            elif g.returncode != 1 or "4242.5" not in g.stdout:
                fails.append(f"4a: an edited chart value must be refused, named (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[:300]}")

        # 2b. the run's own vocabulary on a slide — a step token or a ledger id — is refused
        spec.write_text(GOOD_SPEC.replace(
            '"The FY2023 income statement ties to the trial balance with no exception."',
            '"The FY2023 income statement ties to the trial balance with no exception, as q1 found (Q.q6.mgmt_residual)."'),
            encoding="utf-8")
        r = run(build)
        g = run(gate)
        if r.returncode != 0 or g.returncode != 1 or "vocabulary" not in g.stdout \
                or "Q.q6.mgmt_residual" not in g.stdout:
            fails.append(f"2b: a step token or ledger id on a slide must be refused (build {r.returncode}, "
                         f"gate {g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")

        # 2. a typed figure in a sentence: builds, gate refuses it by token
        spec.write_text(GOOD_SPEC.replace(
            '"The FY2023 income statement ties to the trial balance with no exception."',
            '"The FY2023 income statement ties to the trial balance within $1,234,567."'), encoding="utf-8")
        r = run(build)
        g = run(gate)
        if r.returncode != 0 or g.returncode != 1 or "$1,234,567" not in g.stdout:
            fails.append(f"2: a typed figure must build and be refused by the gate, named (build {r.returncode}, "
                         f"gate {g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")

        # 3. a reference nothing resolves
        spec.write_text(GOOD_SPEC.replace("{q6 | = Reported EBITDA | LTM Jul 2025 | $}", "{q6 | = Reported EBITDA | LTM Aug 2025}", 1), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "LTM Aug 2025" not in r.stdout:
            fails.append(f"3: an unresolved reference must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 5. a sentence where a headline belongs: over the cap, and ending in a full stop
        spec.write_text(GOOD_SPEC.replace("title: Adjusted EBITDA", f'title: "{SENTENCE_TITLE}"'), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "pages[1].title" not in r.stdout or "headline" not in r.stdout:
            fails.append(f"5a: a sentence title must be refused as not a headline, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace("title: Adjusted EBITDA", "title: The bridge foots."), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "pages[1].title" not in r.stdout or "full stop" not in r.stdout:
            fails.append(f"5b: a title ending in a full stop must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 7. the cover: the title names the work, not the period (whether it names the
        #    company is the critic's judgment, skills/check-review — no pattern here)
        spec.write_text(GOOD_SPEC.replace("title: Quality of earnings review",
                                          "title: Quality of earnings review, Acme Corp"), encoding="utf-8")
        r = run(build)
        if r.returncode != 0:
            fails.append(f"7a: a cover title naming the company is the critic's call, not the builder's "
                         f"(exit {r.returncode}): {(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace("title: Quality of earnings review",
                                          "title: Quality of earnings review for FY2023"), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "period" not in r.stdout:
            fails.append(f"7b: a cover title carrying the period must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace(
            "title: Quality of earnings review",
            "title: Quality of earnings review\nsubtitle: Quality of earnings review · FY2023"), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "subtitle:" not in r.stdout or "repeats the title" not in r.stdout:
            fails.append(f"7c: a cover subtitle repeating the title must be refused, named "
                         f"(exit {r.returncode}): {(r.stdout + r.stderr).strip()[:300]}")
        # the gate refuses it on a deck the builder never saw: the cover title, edited after the build
        spec.write_text(GOOD_SPEC, encoding="utf-8")
        r = run(build)
        edited = deck.with_name("cover.pptx")
        hits = 0
        with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(edited, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("ppt/slides/slide") and b'name="cover-title"' in data and not hits:
                    e = data.replace(b"<a:t>Quality of earnings review</a:t>",
                                     b"<a:t>Quality of earnings review FY2023</a:t>", 1)
                    if e != data:
                        data, hits = e, hits + 1
                zout.writestr(item, data)
        g = run([py, str(SCRIPTS / "check_report.py"), str(edited), "--workbook", str(wb), "--run-dir", str(rd)])
        if not hits or g.returncode != 1 or "cover title" not in g.stdout:
            fails.append(f"7d: the gate must refuse a cover title carrying the period "
                         f"(exit {g.returncode}, edited {hits}): {(g.stdout + g.stderr).strip()[:300]}")

        # 8. a zero in a sentence, and 9. a schedule condensed past its own arithmetic
        spec.write_text(GOOD_SPEC, encoding="utf-8")
        r = run(build)
        if r.returncode == 0 and deck.is_file():
            text_ = "\n".join(stored_slides(deck))
            if "the amount is $0." not in html.unescape(text_).replace("</a:t><a:t>", ""):
                fails.append("8: a zero in a sentence must read `$0` — the en dash is the table's zero")
        spec.write_text(GOOD_SPEC.replace(
            'rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025]',
            'rows: ["EBIT", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025]'), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "do not foot" not in r.stdout or "Non-recurring" not in r.stdout:
            fails.append(f"9: a schedule condensed past its own arithmetic must be refused, the dropped "
                         f"rows named (exit {r.returncode}): {(r.stdout + r.stderr).strip()[:400]}")

        # 10. the deck's number conventions
        spec.write_text(GOOD_SPEC.replace(
            "columns: [line, FY2023, FY2024, LTM Jul 2025], title:",
            "columns: [line, FY2023, FY2024, LTM Jul 2025], scale: thousands, title:"), encoding="utf-8")
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"10: a scaled schedule must build (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[-300:]}")
        else:
            body = html.unescape("\n".join(stored_slides(deck))).replace("</a:t><a:t>", "")
            for want in ("$8.4M", "($ in thousands)", "8,404"):
                if want.lower() not in body.lower():   # a table title is set in capitals
                    fails.append(f"10: the deck must carry `{want}` — REPORT.md § 4")
            # the scale is stated once, in the table's title — never in a column header,
            # where `LTM JUL 2025 ($ IN THOUSANDS)` wraps a narrow column
            heads = [h for x in stored_slides(deck) for h in first_rows(x)]
            if not any("in thousands" in t.lower() for x in stored_slides(deck)
                       for t in re.findall(r'name="table-title".*?</p:sp>', x, re.S)):
                fails.append("10: a scaled table's title must state `($ in thousands)`")
            if any("thousands" in h.lower() for h in heads):
                fails.append(f"10: no column header states the scale: {[h for h in heads if 'thousands' in h.lower()]}")
            g = run(gate)
            if g.returncode != 0:
                fails.append(f"10: a scaled deck must pass the gate (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[-300:]}")

        # 11. a dense schedule, filtered and ended
        spec.write_text(GOOD_SPEC, encoding="utf-8")
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"11: the good spec must build (exit {r.returncode})")
        else:
            body = html.unescape("\n".join(stored_slides(deck))).replace("</a:t><a:t>", "")
            if 'sz="825"' not in body:
                fails.append("11: a table declared `dense: true` is set at 8.25pt (`sz=\"825\"`)")
            walk = next((x for x in stored_slides(deck) if "EBITDA bridge by line" in x and "table:q6" in x), "")
            if "Management adjusted EBITDA (information)" in walk:
                fails.append("11: `through:` ends the walk at the closing line — the information "
                             "line under it must not be on the walk's page")
            if "Owner compensation, no adjustment taken" not in walk or "+ D&amp;A" not in walk:
                fails.append("11: `where:` keeps the walk's mechanics and every ruled row")

        # 12. the recipe's schedule is held by the gate
        spec.write_text(GOOD_SPEC.replace(
            'where: {verdict: supported}, through: "= Diligence adjusted EBITDA", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict], dense: true',
            'max_rows: 3, columns: [line, FY2023, FY2024, LTM Jul 2025, verdict]'), encoding="utf-8")
        r = run(build)
        if r.returncode != 0:
            fails.append(f"12: a trimmed walk must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            g = run(gate)
            if g.returncode != 1 or "schedule `EBITDA walk`" not in g.stdout or "are not on the deck" not in g.stdout:
                fails.append(f"12: a walk trimmed to max_rows must be refused naming the schedule and the "
                             f"rows it lacks (exit {g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")
        spec.write_text(re.sub(r"      - title: EBITDA bridge by line\n        blocks:\n          - table: .*\n", "", GOOD_SPEC)
                        .replace("          - table: {from: q6, max_rows: 5}\n", "")
                        .replace("          - table: {from: q6, block: Exceptions}\n", "          - text: \"No exception stands.\"\n")
                        .replace("          - table: {from: q6, rows: [\"= Reported EBITDA\", \"= Diligence adjusted EBITDA\"], "
                                 "columns: [line, FY2023, FY2024, LTM Jul 2025], title: \"EBITDA bridge, USD\"}\n", ""), encoding="utf-8")
        r = run(build)
        if r.returncode != 0:
            fails.append(f"12: a deck without the walk must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            g = run(gate)
            if g.returncode != 1 or "no page carries a table from `q6" not in g.stdout:
                fails.append(f"12: a deck with no table from the walk's tab must be refused naming the "
                             f"schedule (exit {g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")

        # 13. the opening is held by the gate (REPORT.md § 1): the executive summary page
        #     first, with a message and a figure block; the key-metrics page headed as the
        #     recipe declares; the first schedule at most one page after it.
        opening = {
            "a": (GOOD_SPEC.replace("      - title: Executive summary\n", "      - title: Diligence-adjusted EBITDA\n", 1),
                  "headed `Executive summary`"),
            "b": (GOOD_SPEC.replace(
                '          - stats:\n'
                '              - {label: Reported EBITDA · LTM Jul 2025, value: "{q6 | = Reported EBITDA | LTM Jul 2025 | $}"}\n'
                '              - {label: Diligence adjusted EBITDA · LTM Jul 2025, value: "{q6 | = Diligence adjusted EBITDA | LTM Jul 2025 | $}"}\n',
                ''), "no stat tile, table or chart"),
            "c": (GOOD_SPEC.replace("      - title: Adjusted EBITDA\n", "      - title: EBITDA\n", 1),
                  "key-metrics page is headed `Adjusted EBITDA`"),
            "d": (GOOD_SPEC.replace(
                "      - title: EBITDA bridge by line\n",
                '      - title: Context first\n        blocks:\n          - text: "We set the scene here."\n'
                '      - title: Context second\n        blocks:\n          - text: "We set more of the scene here."\n'
                "      - title: EBITDA bridge by line\n", 1), "2 pages after the key-metrics page"),
        }
        for case, (text_, want) in opening.items():
            assert text_ != GOOD_SPEC, f"13{case}: the case did not change the spec"
            spec.write_text(text_, encoding="utf-8")
            r = run(build)
            if r.returncode != 0:
                fails.append(f"13{case}: the spec must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
                continue
            g = run(gate)
            if g.returncode != 1 or want not in g.stdout:
                fails.append(f"13{case}: the gate must refuse the opening naming `{want}` (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[:400]}")

        # 14. money is the book's currency, written and read back by scripts/style.py: a
        #     EUR book states `€8.4M` in a tile and `(€ in thousands)` over a scaled
        #     schedule, and passes; a typed `€1,234,567` is refused, not skipped.
        rd_eur = pathlib.Path(td) / "run_eur"
        make_run(rd_eur, cur="EUR")
        wb_eur = rd_eur / "out" / ".staging" / "workbook.xlsx"
        run([py, str(SCRIPTS / "link_workbook.py"), str(wb_eur)])
        spec_eur = rd_eur / "out" / ".staging" / "report.yaml"
        deck_eur = rd_eur / "out" / ".staging" / "report.pptx"
        build_eur = [py, str(SCRIPTS / "build_report.py"), str(rd_eur)]
        gate_eur = [py, str(SCRIPTS / "check_report.py"), str(deck_eur), "--run-dir", str(rd_eur)]
        spec_eur.write_text(GOOD_SPEC.replace(
            "columns: [line, FY2023, FY2024, LTM Jul 2025], title:",
            "columns: [line, FY2023, FY2024, LTM Jul 2025], scale: thousands, title:"), encoding="utf-8")
        r = run(build_eur)
        if r.returncode != 0 or not deck_eur.is_file():
            fails.append(f"14: a EUR book must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            body = html.unescape("\n".join(stored_slides(deck_eur))).replace("</a:t><a:t>", "")
            for want in ("€8.4M", "(€ in thousands)"):
                if want.lower() not in body.lower():
                    fails.append(f"14: a EUR deck must carry `{want}`")
            if "$8.4M" in body:
                fails.append("14: a EUR deck must not state a figure in dollars")
            g = run(gate_eur)
            if g.returncode != 0:
                fails.append(f"14: a EUR deck must pass the gate (exit {g.returncode}): {(g.stdout + g.stderr).strip()[-300:]}")
        spec_eur.write_text(GOOD_SPEC.replace("with no exception.", "within €1,234,567."),
                            encoding="utf-8")
        r = run(build_eur)
        g = run(gate_eur)
        if r.returncode != 0 or g.returncode != 1 or "€1,234,567" not in g.stdout:
            fails.append(f"14: `€1,234,567` typed into a EUR deck must be refused, named (build {r.returncode}, "
                         f"gate {g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")

        # 16. the plan's periods, on a September year end: the schedule's period columns
        #     are the plan's (`periods: all` wants FY2023, FY2024 and LTM Jul 2025 by the
        #     plan, not a header regex).
        rd_p = pathlib.Path(td) / "run_periods"
        make_run(rd_p, q6_params={"columns": ["fy2023", "fy2024", "ltm_2025-07"],
                                  "fiscal_year_end": "09-30"})
        run([py, str(SCRIPTS / "link_workbook.py"), str(rd_p / "out" / ".staging" / "workbook.xlsx")])
        spec_p = rd_p / "out" / ".staging" / "report.yaml"
        build_p = [py, str(SCRIPTS / "build_report.py"), str(rd_p)]
        gate_p = [py, str(SCRIPTS / "check_report.py"), str(rd_p / "out" / ".staging" / "report.pptx"),
                  "--run-dir", str(rd_p)]
        spec_p.write_text(GOOD_SPEC, encoding="utf-8")
        r, g = run(build_p), run(gate_p)
        if r.returncode != 0 or g.returncode != 0:
            fails.append(f"16: the good spec on a run declaring its periods must build and pass (build "
                         f"{r.returncode}, gate {g.returncode}): {(r.stdout + g.stdout).strip()[-300:]}")
        spec_p.write_text(GOOD_SPEC.replace(
            "through: \"= Diligence adjusted EBITDA\", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict]",
            "through: \"= Diligence adjusted EBITDA\", columns: [line, FY2023, LTM Jul 2025, verdict]")
            .replace("{from: q6, max_rows: 5}", "{from: q6, max_rows: 5, columns: [line, FY2023, LTM Jul 2025, verdict]}"),
            encoding="utf-8")
        r, g = run(build_p), run(gate_p)
        if r.returncode != 0 or g.returncode != 1 or "`FY2024`" not in g.stdout:
            fails.append(f"16: a walk without the plan's FY2024 column must be refused, named (build "
                         f"{r.returncode}, gate {g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")
        #     A check whose periods scripts/periods.py refuses (here a non-IANA zone) is
        #     named by the gate.
        rj = json.loads((rd_p / "run.json").read_text(encoding="utf-8"))
        rj["checks"][1]["params"]["timezone"] = "Mars/Olympus_Mons"
        (rd_p / "run.json").write_text(json.dumps(rj), encoding="utf-8")
        spec_p.write_text(GOOD_SPEC, encoding="utf-8")
        r, g = run(build_p), run(gate_p)
        if r.returncode != 0 or g.returncode != 1 or "check `q6_bridge`" not in g.stdout \
                or "Mars/Olympus_Mons" not in g.stdout:
            fails.append(f"16: a check whose periods periods.py refuses must be named by the gate "
                         f"(build {r.returncode}, gate {g.returncode}): "
                         f"{(g.stdout + g.stderr).strip()[:300]}")
        # 17. the recipe decides the structure (RECIPE_FORMAT.md § Report): a schedule
        #     placed `appendix` sits under the `Appendix` kicker, after the narrative. The
        #     good spec (walk under `The bridge`) is refused naming `Appendix`; the walk
        #     moved to the close of the appendix passes; a page after the appendix is
        #     refused; the narrative's sections are held to the declared order.
        recipe_path = rd / "QOE.md"
        appx_recipe = RECIPE.replace('"dense": true}', '"dense": true, "place": "appendix"}')
        assert appx_recipe != RECIPE
        walk_page = ("      - title: EBITDA bridge by line\n        blocks:\n          - table: {from: q6, where: {verdict: supported}, "
                     "through: \"= Diligence adjusted EBITDA\", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict], dense: true}\n")
        assert walk_page in GOOD_SPEC
        in_appendix = GOOD_SPEC.replace(walk_page, "") + walk_page
        structure_cases = {
            "a": (appx_recipe, GOOD_SPEC, "kicker `Appendix`"),
            "b": (appx_recipe, in_appendix, None),
            "c": (appx_recipe, in_appendix + "  - title: Afterword\n    pages:\n      - title: Afterword\n"
                                             "        blocks:\n          - text: \"We close here.\"\n",
                  "follows the appendix"),
            "d": (appx_recipe.replace('"schedules"', '"narrative": ["Findings"], "schedules"'), in_appendix,
                  "not a narrative section"),
            "e": (appx_recipe.replace('"schedules"', '"narrative": ["The bridge"], "schedules"'), in_appendix, None),
        }
        for case, (rtext, stext, want) in structure_cases.items():
            recipe_path.write_text(rtext, encoding="utf-8")
            spec.write_text(stext, encoding="utf-8")
            r = run(build)
            if r.returncode != 0:
                fails.append(f"17{case}: the spec must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
                continue
            g = run(gate)
            if want is None and g.returncode != 0:
                fails.append(f"17{case}: the deck follows the recipe's structure and must pass (exit "
                             f"{g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")
            elif want is not None and (g.returncode != 1 or want not in g.stdout):
                fails.append(f"17{case}: the gate must refuse the structure naming `{want}` (exit "
                             f"{g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")
        recipe_path.write_text(RECIPE, encoding="utf-8")

        # 18. two rows that read the same in every column shown are refused: the reader
        #     cannot tell them apart.
        spec.write_text(GOOD_SPEC.replace(
            'rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025], title: "EBITDA bridge, USD"',
            'rows: ["+ D&A", "+ D&A"], columns: [line, FY2023, FY2024, LTM Jul 2025], title: "EBITDA bridge, USD"'),
            encoding="utf-8")
        r, g = run(build), run(gate)
        if r.returncode != 0 or g.returncode != 1 or "read the same in every column" not in g.stdout:
            fails.append(f"18: a table repeating a row must be refused (build {r.returncode}, gate "
                         f"{g.returncode}): {(r.stdout + g.stdout).strip()[:400]}")

        # 19. a walk as a waterfall: builds, each bar carries its value, and passes the
        #     gate; a waterfall whose steps do not reach a total it shows is refused.
        wf_rows = ('["EBIT", "+ D&A", "= Reported EBITDA", "Non-recurring, supported", '
                   '"Normalization, supported", "= Diligence adjusted EBITDA"]')
        chart_line = ('          - chart: {type: column, from: q6, rows: ["= Reported EBITDA", '
                      '"= Diligence adjusted EBITDA"], columns: [FY2023, FY2024, LTM Jul 2025]}\n')
        assert chart_line in GOOD_SPEC
        spec.write_text(GOOD_SPEC.replace(chart_line, f"          - chart: {{type: waterfall, from: q6, rows: "
                                                      f"{wf_rows}, columns: [LTM Jul 2025]}}\n"), encoding="utf-8")
        r = run(build)
        if r.returncode != 0:
            fails.append(f"19: a waterfall must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            body = html.unescape("\n".join(stored_slides(deck)))
            if 'name="chart-value"' not in body or "8,404,000" not in body:
                fails.append("19: each waterfall bar carries its value beside it")
            if "Source: workbook.xlsx / EBITDA bridge" not in body or "· q6 " in body:
                fails.append("19: a footer names each tab without its roster token "
                             "(`Source: workbook.xlsx / EBITDA bridge`)")
            if "Axis from $" not in body:
                fails.append("19: a walk whose steps are small beside its totals says where its axis starts")
            g = run(gate)
            if g.returncode != 0:
                fails.append(f"19: a waterfall deck must pass (exit {g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")
        spec.write_text(GOOD_SPEC.replace(chart_line, "          - chart: {type: waterfall, from: q6, rows: "
                                          "[\"EBIT\", \"= Reported EBITDA\", \"= Diligence adjusted EBITDA\"], "
                                          "columns: [LTM Jul 2025]}\n"), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "do not foot" not in r.stdout:
            fails.append(f"19: a waterfall skipping its steps must be refused (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 21. a check tab stating a status as the run's code (`not_supported`) is refused by
        #     the workbook gate (WORKBOOK.md § 3 Language); a check id naming a check is not.
        from openpyxl import load_workbook  # noqa: PLC0415
        coded = rd / "out" / ".staging" / "coded.xlsx"
        wbc = load_workbook(wb)
        wbc["q6 EBITDA bridge"]["H6"].value = "not_supported"
        wbc["q6 EBITDA bridge"]["H7"].value = "q1_fy2023"
        wbc.save(coded)
        g = run([py, str(SCRIPTS / "check_workbook.py"), str(coded)])
        if g.returncode != 1 or "machine vocabulary" not in g.stdout or "`not_supported`" not in g.stdout:
            fails.append(f"21: a status written as the run's code must be refused, named (exit "
                         f"{g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")
        elif "q1_fy2023" in g.stdout:
            fails.append("21: a check id naming its check is navigation, not machine vocabulary")

        # 20. days on the deck: never `-0.0`; a table's negative in parentheses.
        import build_report  # noqa: PLC0415 — the builder's own formatter
        for v, want in ((-0.004, "–"), (0.0, "–"), (-5.66, "(5.7)"), (5.66, "5.7")):
            got = build_report.fmt_value(v, "0.0")
            if got != want:
                fails.append(f"20: {v} formatted `0.0` in a table must read `{want}`, not `{got}`")
        if build_report.fmt_value(-0.004, "0.0", prose=True) != "0":
            fails.append("20: a zero in a sentence reads `0`, never `-0.0`")

        # 6. a fragment where a sentence belongs: a text block with no full stop
        spec.write_text(GOOD_SPEC.replace(
            '- text: "Every rostered check is listed with what it examined and what it did not."',
            '- text: "Every rostered check, what it examined and what it did not"'), encoding="utf-8")
        r = run(build)
        if r.returncode != 1 or "blocks[0]" not in r.stdout or "full stop" not in r.stdout:
            fails.append(f"6: a fragment text block must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 22. working-paper terms (REPORT.md § 3): a sentence the author writes with one is
        #     refused by the gate, named; the same term opening its definition in a status
        #     note passes, as does a copied table's cell.
        bullet = '- bullets: ["The FY2023 income statement ties to the trial balance with no exception."]'
        assert bullet in GOOD_SPEC
        defined = GOOD_SPEC.replace(bullet, bullet + '\n          - note: "Candidate: an adjustment the '
                                    'records do not yet support, outside the adjusted figure."')
        spec.write_text(defined, encoding="utf-8")
        r, g = run(build), None
        if r.returncode == 0:
            g = run(gate)
        if r.returncode != 0 or g.returncode != 0:
            fails.append(f"22: a status note defining `Candidate:` must build and pass (build {r.returncode}, "
                         f"gate {g.returncode if g else '-'}): {((g or r).stdout + (g or r).stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace(bullet, '- bullets: ["We walked reported EBITDA to the adjusted '
                                          'figure, and two candidates remain."]'), encoding="utf-8")
        r = run(build)
        g = run(gate) if r.returncode == 0 else r
        if g.returncode != 1 or "working-paper terms" not in g.stdout or "walked" not in g.stdout \
                or "candidates" not in g.stdout:
            fails.append(f"22: a working-paper term in a bullet must be refused, named (exit {g.returncode}): "
                         f"{(g.stdout + g.stderr).strip()[:300]}")

        # 23. a walk's axis (REPORT.md § 2): steps small beside the totals start the axis
        #     above zero; a walk crossing zero keeps zero; a money axis carries its currency.
        lo, hi, _, cut = build_report.waterfall_axis(14_600_000, 18_700_000)
        if not (cut and 0 < lo <= 14_600_000 and hi >= 18_700_000):
            fails.append(f"23: a walk from $18.7M to $14.6M starts its axis above zero ({lo}, {hi}, cut {cut})")
        if build_report.waterfall_axis(-1_000_000, 5_000_000)[3]:
            fails.append("23: a walk crossing zero keeps zero on its axis")
        if build_report.waterfall_axis(100_000, 5_000_000)[3]:
            fails.append("23: a walk whose steps are large beside its totals keeps zero on its axis")
        for v, step, want in ((12_500_000, 2_500_000, "$12.5M"), (-500_000, 100_000, "($0.5M)"),
                              (0.0, 100_000, "0")):
            got = build_report.axis_label(v, step, "usd")
            if got != want:
                fails.append(f"23: a money axis tick {v} reads `{want}`, not `{got}`")

        # 24. continuation pages (REPORT.md § 2 Fit): a table split over pages leaves no
        #     piece of fewer than three rows, and a note that does not fit after a table
        #     takes the table's last rows with it.
        def table_of(n: int) -> dict:
            rows = [[build_report.Cell(f"Item {i}", "", False, "", f"A{i}"),
                     build_report.Cell(1000.0 * i, "#,##0", False, "", f"B{i}")] for i in range(n)]
            return {"t": "table", "table": build_report.Table("x1 Items", None, ["item", "amount"], rows,
                                                               kinds=["body"] * n, numeric=[False, True])}
        widows, lone = [], []
        for n in range(5, 70):
            pages = build_report.flow(build_report.Page("s", "s", "Items", [table_of(n)]))
            sizes = [len(b["table"].rows) for pg in pages for b in pg.blocks if b["t"] == "table"]
            if len(sizes) > 1 and min(sizes) < build_report.WIDOW_ROWS:
                widows.append((n, sizes))
            pages = build_report.flow(build_report.Page("s", "s", "Items", [
                table_of(n), {"t": "note", "text": "Supported: the records carry the amount."}]))
            if any(all(b["t"] == "note" for b in pg.blocks) for pg in pages):
                lone.append(n)
        if widows:
            fails.append(f"24: a table continued over pages leaves a piece under three rows: {widows[:3]}")
        if lone:
            fails.append(f"24: a note stands alone on a continuation page, for tables of {lone[:5]} rows")

    if fails:
        print(f"report-selftest: {len(fails)} failure(s)")
        for f in fails:
            print(f"  {f}")
        return 1
    print("report-selftest: ok — a spec-conformant workbook builds a deck that passes, headlines and "
          "messages as separate shapes and tables first-column-left, rest-right; a typed figure, an "
          "unresolved reference, an edited figure, a sentence title, a fragment, a cover title "
          "carrying the period and a schedule condensed past its own arithmetic are refused; a "
          "zero in a sentence reads $0, a money figure reads $8.4M (€8.4M in a EUR book) and a "
          "declared schedule reads ($ in thousands); a typed non-dollar figure is refused; "
          "periods are the plan's; a dense, "
          "filtered walk builds, and the recipe's "
          "schedule trimmed or absent is refused; the opening — the executive summary with its message "
          "and a figure block, the key-metrics page headed as the recipe declares, the first schedule "
          "at most one page after it — is held; the recipe's structure — an appendix schedule under "
          "the `Appendix` kicker, nothing after the appendix, the narrative in its declared order — is "
          "held; twin rows are refused; a waterfall builds and foots; days never read `-0.0`; a status "
          "written as the run's code is refused on the workbook.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
