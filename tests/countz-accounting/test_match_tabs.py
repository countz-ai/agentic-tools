#!/usr/bin/env python3
"""Self-test for scripts/match_tabs.py.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_match_tabs.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import datetime as dt
import pathlib
import sys

import polars as pl

# The plugin under test: <repo>/countz-accounting/scripts, from <repo>/tests/countz-accounting.
SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from match_tabs import match_tabs  # noqa: E402
from wbkit import save as wbkit_save  # noqa: E402  every formula's result cached, as a tab script saves


def main() -> int:
    import json
    import tempfile
    import zipfile

    from openpyxl import Workbook, load_workbook

    import check_workbook
    import link_workbook
    from figures import Ledger
    from resolve import Pass, from_assignment, resolve, rules
    bad = []
    D = dt.date(2024, 3, 1)
    left = pl.DataFrame([("i1", "c1", D, 100.37), ("i2", "c1", D, 250.13), ("i3", "c2", D, 75.25),
                         ("i4", "c3", D + dt.timedelta(days=9), 81.10),
                         ("i5", "c4", D + dt.timedelta(days=40), 12.34)],
                        orient="row", schema=["id", "entity", "date", "value"])
    right = pl.DataFrame([("d1", "A", D, 350.50), ("d2", "A", D, 75.25),
                          ("d3", "A", D + dt.timedelta(days=9), 81.10),
                          ("d4", "B", D + dt.timedelta(days=20), 9.99)],
                         orient="row", schema=["id", "entity", "date", "value"])
    res = resolve(left, right, rules(window=(0, 2)))
    with tempfile.TemporaryDirectory() as tmp:
        run = pathlib.Path(tmp)
        (run / "workpapers").mkdir()
        L = Ledger(run, "m1", fresh=True)
        L.cite({"id": "E.m1.left", "kind": "span", "file": "l.xlsx", "source": "l",
                "file_role": "system_export", "sheet": "S", "header_at": "A1", "rows": "2:6",
                "columns": [], "filter": "none - full sheet consumed", "row_count": 5,
                "control_total": {"column": "v", "value": 519.19}})
        L.population("P.m1.left", "left items", 6, 6, citations=["E.m1.left"])
        L.population("P.m1.right", "right items", 4, 4, citations=["E.m1.left"])
        wb = Workbook()
        wb.active.title = "m1 Reconciliation"
        out = match_tabs(wb, res, left, right, check="m1", token="m1", ledger=L,
                         subtitle="Fixture · March 2024 · USD", inputs=[("left", "E.m1.left")],
                         population="P.m1.left", right_population="P.m1.right",
                         others=pl.DataFrame({"id": ["i9"], "value": [0.0], "reason": ["open"]}),
                         left_label=pl.DataFrame({"id": ["i1"], "label": ["Customer One"]}))
        L.write()
        path = run / "tab.xlsx"
        wbkit_save(wb, path)
        wbb = load_workbook(path)
        if wbb.sheetnames[:4] != [out["summary"], out["schedule"], out["reconciling"], out["rules"]]:
            bad.append(f"the four tabs lead the file: {wbb.sheetnames}")
        sch = wbb[out["schedule"]]
        states = [sch.cell(r, 6).value for r in range(5, 11)]
        if states != ["Matched"] * 4 + ["In transit", "No cash"]:
            bad.append(f"the schedule is sorted by status: {states}")
        with zipfile.ZipFile(path) as z:
            if fails := check_workbook.audit_match(z, assembled=False):
                bad.append(f"the match gate refuses the builder's own tabs: {fails[:3]}")
            if design := [f for f in check_workbook.audit_design(z) if out["summary"] in f or out["schedule"] in f]:
                bad.append(f"the design gate refuses the tabs: {design[:3]}")
        rec = load_workbook(path)
        rec[out["reconciling"]]["F5"].value = rec[out["reconciling"]]["F5"].value + 1   # and a gross
        rec.save(run / "gross.xlsx")                                                   # that is not the net
        with zipfile.ZipFile(run / "gross.xlsx") as z:
            if not any("positive and negative" in f for f in check_workbook.audit_match(z, assembled=False)):
                bad.append("a line whose gross is not its net passes the gate")
        ws = wbb[out["summary"]]
        ws["D6"].value = ws["D6"].value + 1                    # a tampered line is refused
        wbb.save(path)
        with zipfile.ZipFile(path) as z:
            if not check_workbook.audit_match(z, assembled=False):
                bad.append("a summary line that is not the filter passes the gate")
        wbl = load_workbook(path)
        targets = link_workbook.match_targets(wbl)
        if targets.get((out["summary"], "C6")) != (out["schedule"], "B5:K8") or \
                targets.get((out["summary"], "C7")) != (out["schedule"], "B9:K9"):
            bad.append(f"each line links to its rows: {targets}")

        # a match built with joins gets the same tabs, and they say where it came from
        asg = res.items.select("side", "id", "status", "group", "pass", reason=pl.lit(None, pl.Utf8))
        fa = from_assignment(left, right, asg, [Pass(r, "as resolve() ran it") for r in res.by_rule["rule"]],
                             transit=2)
        L2 = Ledger(run, "m2", fresh=True)
        L2.population("P.m2.left", "left items", 5, 5, citations=["E.m1.left"])
        wb2 = Workbook()
        wb2.active.title = "m2 Reconciliation"
        out2 = match_tabs(wb2, fa, left, right, check="m2", token="m2", ledger=L2,
                          subtitle="Fixture · March 2024 · USD", inputs=[("left", "E.m1.left")],
                          population="P.m2.left", right_population="P.m2.right")
        wbkit_save(wb2, run / "src.xlsx")
        with zipfile.ZipFile(run / "src.xlsx") as z:
            if fails := check_workbook.audit_match(z, assembled=False):
                bad.append(f"the match gate refuses tabs from from_assignment(): {fails[:3]}")
        said = " ".join(str(c.value) for row in load_workbook(run / "src.xlsx")[out2["summary"]].iter_rows()
                        for c in row if c.value)
        if "the passes on m2 Match rules" not in said or "resolve.py run on" in said:
            bad.append("the summary's To reperform names the passes of a join-built match")

        # GATE 7: every reconciliation carries its match tabs
        (run / "run.json").write_text(json.dumps({"checks": [
            {"id": "m1", "kind": "recon"}, {"id": "m2", "kind": "recon"}, {"id": "t1", "kind": "tieout"}]}))

        def gate7(name, sheets, notes=None, drop=()):
            w = load_workbook(run / "src.xlsx")
            for t in [t for t in w.sheetnames if t not in sheets] + list(drop):
                del w[t]
            if notes:
                w[w.sheetnames[-1]]["B20"] = notes
            w.save(run / f"{name}.xlsx")
            return check_workbook.audit(run / f"{name}.xlsx", run_dir=run)["recon"]
        every = load_workbook(run / "src.xlsx").sheetnames
        for name, sheets, notes, drop, refused, what in (
                ("m2", every, None, (), False, "a reconciliation with its match tabs"),
                ("m2", ["m2 Reconciliation"], None, (), True, "a reconciliation with no match tabs"),
                ("m2", every, None, (out2["reconciling"], out2["rules"]), True, "a reconciliation with half a set"),
                ("m2", ["m2 Reconciliation"], "No item grain: the statement holds its closing balance alone.",
                 (), False, "a reconciliation that states a side with no item grain"),
                ("t1", ["m2 Reconciliation"], None, (), False, "a tie-out, which matches nothing")):
            if bool(gate7(name, sheets, notes, drop)) != refused:
                bad.append(f"GATE 7 {'passes' if refused else 'refuses'} {what}")
        with zipfile.ZipFile(run / "src.xlsx") as z:
            if check_workbook.audit_recon(z, "m2", None):
                bad.append("GATE 7 holds a tab gated without the run's roster")
        seal = load_workbook(run / "src.xlsx")
        seal.create_sheet("Exec Summary", 0)
        seal.save(run / "workbook.xlsx")
        with zipfile.ZipFile(run / "workbook.xlsx") as z:
            if check_workbook.audit_recon(z, "workbook", run):
                bad.append("GATE 7 refuses a sealed reconciliation with its match tabs")
        for t in (out2["summary"], out2["schedule"], out2["reconciling"], out2["rules"]):
            del seal[t]
        seal.save(run / "workbook.xlsx")
        with zipfile.ZipFile(run / "workbook.xlsx") as z:
            if [f.split(":")[0] for f in check_workbook.audit_recon(z, "workbook", run)] != ["m2"]:
                bad.append("GATE 7 passes a sealed reconciliation whose match tabs were dropped")
        bad += regressions(run)
    for b in bad:
        print("FAIL", b)
    print("match_tabs.py self-check:", "FAIL" if bad else "ok")
    return 1 if bad else 0


def regressions(run: pathlib.Path) -> list[str]:
    """Each input the tabs once showed wrongly, or crashed on, beside its valid neighbour:
    shown faithfully, or refused before anything is written."""
    import copy
    import io
    import re
    import zipfile

    from openpyxl import Workbook, load_workbook

    import check_workbook
    from figures import Ledger
    from resolve import Pass, Rule, from_assignment, resolve
    from wbkit import STATUS, register_status
    bad = []
    D = dt.date(2024, 3, 1)
    day = lambda n: D + dt.timedelta(days=n)  # noqa: E731

    def tabs(res, left, right, wb=None, ledger=None, **kw):
        wb = Workbook() if wb is None else wb
        ledger = Ledger(run, "r1", fresh=True) if ledger is None else ledger
        kw = {"check": "r1", "token": "r1", "subtitle": "Fixture · USD", "inputs": [("left", "E.m1.left")],
              "population": "P.m1.left", "right_population": "P.m1.right", **kw}
        return wb, match_tabs(wb, res, left, right, ledger=ledger, **kw), ledger

    def saved(wb):
        """The workbook as stored, every formula's result cached (wbkit.save)."""
        p = run / "saved.xlsx"
        wbkit_save(wb, p)
        return io.BytesIO(p.read_bytes())

    def gate(wb):
        with zipfile.ZipFile(saved(wb)) as z:
            return check_workbook.audit_match(z, assembled=False)

    def refused(words, fn) -> bool:
        try:
            fn()
        except ValueError as e:
            return words in str(e)
        return False

    def listed(ws) -> dict:
        """The Reconciling items' list of items not matched: {id: row values}."""
        first = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 3).value == "Side") + 1
        return {ws.cell(r, 2).value: [ws.cell(r, c).value for c in range(3, 9)] for r in range(first, ws.max_row + 1)}

    left = pl.DataFrame({"id": ["a", "b", "c"], "date": [D, D, day(-30)], "value": [10.0, 20.0, 5.0]})
    right = pl.DataFrame({"id": ["x", "y"], "date": [D, day(5)], "value": [10.0, 7.0]})
    res = resolve(left, right)

    # only the `others` are kept out; a reason reading `kept out` does not make an item one
    asg = pl.DataFrame({"side": ["left", "left", "left", "right", "right"], "id": ["a", "b", "c", "x", "y"],
                        "status": ["matched", "unmatched", "unmatched", "matched", "unmatched"],
                        "group": [1, None, None, 1, None], "pass": ["p1", None, None, "p1", None],
                        "reason": [None, "kept out of p1: two candidates", None, None, None]})
    wb, out, _ = tabs(from_assignment(left, right, asg, [Pass("p1", "same amount")], transit=0), left, right)
    rec = wb[out["reconciling"]]
    if (rec["D5"].value, rec["E5"].value, rec["E7"].value) != (3, 35.0, -25.0):
        bad.append(f"an item whose reason opens `kept out` is counted as kept out: {rec['D5'].value} "
                   f"items, {rec['E5'].value} per books, {rec['E7'].value} not matched")
    # the `others` count in the total; those with an amount are also listed
    others = pl.DataFrame({"id": ["k1", "k2", "k3"], "value": [0.0, 0.0, 50.0],
                           "reason": ["void", "credited", "disputed"]})
    wb, out, _ = tabs(res, left, right, others=others)
    summ, rec = wb[out["summary"]], wb[out["reconciling"]]
    got = listed(rec)
    if (rec["D5"].value, rec["E5"].value) != (6, 85.0) or rec["D7"].value != summ["D8"].value or \
            rec["E7"].value != -summ["E8"].value or "k3" not in got or {"k1", "k2"} & set(got) or gate(wb):
        bad.append(f"the reconciling items and the summary disagree on the items kept out: per books "
                   f"{rec['D5'].value}/{rec['E5'].value}, not matched {rec['D7'].value} (summary "
                   f"{summ['D8'].value}), listed {sorted(got)}")

    # a label that is not text is written as its digits; a label table's words win
    l2 = pl.DataFrame({"id": ["a", "b"], "date": [D, D], "value": [10.0, 20.0], "entity": [12345, 678]})
    r2 = pl.DataFrame({"id": ["x"], "date": [D], "value": [10.0]})
    wb, out, _ = tabs(resolve(l2, r2), l2, r2, left_label=pl.DataFrame({"id": ["b"], "label": ["Acme"]}))
    sch = load_workbook(saved(wb))[out["schedule"]]
    if {sch["B5"].value: sch["C5"].value, sch["B6"].value: sch["C6"].value} != {"a": "12345", "b": "Acme"}:
        bad.append(f"a customer number as the label is lost: {sch['C5'].value!r}, {sch['C6'].value!r}")

    # a deposit of 2,000 lines: as many as a cell holds, then the rest counted and summed
    n = 2000
    l3 = pl.DataFrame({"id": ["a"], "date": [D], "value": [float(n)]})
    r3 = pl.DataFrame({"id": [f"line-{i:05d}" for i in range(n)], "date": [D] * n, "value": [1.0] * n,
                       "batch": ["B1"] * n})
    wb, out, _ = tabs(resolve(l3, r3), l3, r3)
    to = load_workbook(saved(wb))[out["schedule"]]["I5"].value
    m = re.search(r"and ([\d,]+) more bank lines in match group 1, together ([\d,.]+)$", to or "")
    shown = (to or "").count(" · 1.00")
    if len(to or "") > 32_767 or not m or shown + int(m.group(1).replace(",", "")) != n or \
            float(m.group(2).replace(",", "")) != n - shown:
        bad.append(f"a match to {n:,} lines loses some from Matched to ({len(to or ''):,} characters)")
    wb, out, _ = tabs(resolve(l3.with_columns(value=pl.lit(3.0)), r3.head(3)), l3.with_columns(value=pl.lit(3.0)),
                      r3.head(3))
    if wb[out["schedule"]]["I5"].value.count(" · 1.00") != 3 or "more" in wb[out["schedule"]]["I5"].value:
        bad.append("a match to three lines does not list the three")

    # a group's difference once, on its first row: the column sums to the differences
    l4 = pl.DataFrame({"id": ["a", "b", "c", "d"], "date": [D] * 4, "value": [100.0, 50.0, 25.0, 40.0],
                       "entity": ["e", "e", "e", "f"]})
    r4 = pl.DataFrame({"id": ["x", "y"], "date": [D, D], "value": [172.0, 39.0]})
    fee = [Rule("day's items, fee", (), (0, 0), group_left=("entity", "date"), tolerance=(-5.0, 0.0),
                difference="fee")]
    wb, out, _ = tabs(resolve(l4, r4, fee), l4, r4)
    sch, rec = wb[out["schedule"]], wb[out["reconciling"]]
    diffs = {sch.cell(r, 2).value: sch.cell(r, 10).value for r in range(5, 9)}
    line_ = next(rec.cell(r, 5).value for r in range(5, 12) if rec.cell(r, 3).value == "Difference: fee")
    if diffs != {"a": -3.0, "b": None, "c": None, "d": -1.0} or sum(v or 0 for v in diffs.values()) != line_:
        bad.append(f"the schedule's Difference repeats a group's difference on each row: {diffs}, the line {line_}")

    # frames other than the Resolution's are refused, whatever differs; the same frame in
    # another order, its ids numbers, is not
    for what, l_ in (("a row more", pl.concat([pl.DataFrame({"id": ["0"], "date": [day(9)], "value": [9.0]}), left])),
                     ("a row less", left.head(2)),
                     ("an unmatched amount", left.with_columns(value=pl.Series([10.0, 21.0, 5.0]))),
                     ("a matched amount", left.with_columns(value=pl.Series([11.0, 20.0, 5.0]))),
                     ("an unmatched date", left.with_columns(date=pl.Series([D, day(1), day(-30)])))):
        wb = Workbook()
        if not refused("the stream the Resolution was built from", lambda: tabs(res, l_, right, wb=wb)) or \
                wb.sheetnames != ["Sheet"]:
            bad.append(f"a left frame with {what} than the Resolution's is shown")
    l5 = pl.DataFrame({"id": list(range(11, 0, -1)), "date": [D] * 11, "value": [float(i) for i in range(11, 0, -1)]})
    wb, out, _ = tabs(resolve(l5, r2), l5, r2)
    sch = wb[out["schedule"]]
    if any(sch.cell(r, 2).value != str(int(sch.cell(r, 5).value)) for r in range(5, 16)) or gate(wb):
        bad.append("a stream whose ids are numbers, in another order, is shown with the wrong amounts")

    # every amount to the currency's decimals, which are the Resolution's
    l6 = pl.DataFrame({"id": ["a", "b"], "date": [D, D], "value": [1.234, 2.345]})
    r6 = pl.DataFrame({"id": ["x"], "date": [D], "value": [1.234]})
    kwd = resolve(l6, r6, decimals=3)
    wb, out, _ = tabs(kwd, l6, r6, currency="kwd")
    if load_workbook(saved(wb), data_only=True)[out["summary"]]["E5"].value != 3.579 or \
            wb[out["schedule"]]["E5"].number_format != '#,##0.000;(#,##0.000);"–"':
        bad.append(f"a 3-decimal currency is shown to 2: {wb[out['schedule']]['E5'].number_format}")
    if not refused("matched at 3", lambda: tabs(kwd, l6, r6)) or \
            not refused("decimals=2", lambda: tabs(kwd, l6, r6, currency="kwd", decimals=2)):
        bad.append("a currency or decimals other than the Resolution's is taken")
    l7, r7 = l6.with_columns(value=pl.Series([1234.0, 99.0])), r6.with_columns(value=pl.lit(1234.0))
    wb, out, _ = tabs(resolve(l7, r7, decimals=0), l7, r7, currency="jpy", decimals=0)
    if wb[out["schedule"]]["E5"].number_format != '#,##0;(#,##0);"–"':
        bad.append("a currency with no minor unit is shown with cents")

    # a tab name the workbook holds, in any case, is refused; two sets of one check, under
    # two tokens, share a workbook and a ledger
    wb = Workbook()
    wb.create_sheet("R1 MATCH SUMMARY")
    if not refused("already holds 'R1 MATCH SUMMARY'", lambda: tabs(res, left, right, wb=wb)) or \
            wb.sheetnames != ["Sheet", "R1 MATCH SUMMARY"]:
        bad.append("a set whose tab name the workbook holds is written beside it, renamed")
    wb, L = Workbook(), Ledger(run, "r2", fresh=True)
    try:
        _, o1, _ = tabs(res, left, right, wb=wb, ledger=L, check="r2", token="r2 receipts")
        _, o2, _ = tabs(res, left, right, wb=wb, ledger=L, check="r2", token="r2 payments")
        if wb.sheetnames[:8] != [o2[k] for k in ("summary", "schedule", "reconciling", "rules")] + \
                [o1[k] for k in ("summary", "schedule", "reconciling", "rules")] or gate(wb) or \
                set(o1["figures"].values()) & set(o2["figures"].values()):
            bad.append(f"two sets of one check: {wb.sheetnames}")
    except ValueError as e:
        bad.append(f"two sets of one check are refused: {e}")
    if o1["figures"][("left", "Matched", "count")] != "F.r2.match.receipts.left.matched.count":
        bad.append(f"a second set keys by the token past the check id: {o1['figures']}")

    # the check id is written once, so a long one keeps every id one the design gate reads
    # as an id, not prose in a narrow column
    try:
        wb, out, _ = tabs(res, left, right, check="ar_subledger_recon", token="ar_subledger_recon")
        with zipfile.ZipFile(saved(wb)) as z:
            prose = [f for f in check_workbook.audit_design(z) if "prose" in f]
        longest = max(out["figures"].values(), key=len)
        if not longest.startswith("F.ar_subledger_recon.match.right.") or prose:
            bad.append(f"an 18-character check id: {longest}, {prose[:2]}")
    except ValueError as e:
        bad.append(f"an 18-character check id is refused: {e}")

    # an id opening with `=` is text, never a formula
    l8 = pl.DataFrame({"id": ["=1+2", "=HYPERLINK(\"http://x\")"], "date": [D, D], "value": [1.0, 2.0]})
    r8 = pl.DataFrame({"id": ["=SUM(A1:A2)"], "date": [D], "value": [1.0]})
    wb, out, _ = tabs(resolve(l8, r8), l8, r8)
    back = load_workbook(saved(wb))
    ids = [(back[out["schedule"]].cell(r, 2).value, back[out["schedule"]].cell(r, 2).data_type) for r in (5, 6)]
    if sorted(ids) != [("=1+2", "s"), ('=HYPERLINK("http://x")', "s")] or gate(wb):
        bad.append(f"an id opening with `=` is stored as a formula: {ids}")

    # right_population is required; a zero count is measured_zero when every bank line
    # matched, not_measured when there are none
    wb = Workbook()
    if not refused("right_population", lambda: tabs(res, left, right, wb=wb, right_population=None)) or \
            wb.sheetnames != ["Sheet"]:
        bad.append("a set with no population for its right items is written")
    wb, out, L = tabs(resolve(l2.head(1), r2), l2.head(1), r2)
    none = L.entries[out["figures"][("right", "Not in the book", "count")]]
    if (none["value"], none["zero_basis"]) != (0, "measured_zero"):
        bad.append(f"every bank line matched: {none['value']}, {none['zero_basis']}")
    empty = r2.head(0)
    wb, out, L = tabs(resolve(l2.head(1), empty), l2.head(1), empty)
    none = L.entries[out["figures"][("right", "Matched", "amount")]]
    if (none["value"], none["zero_basis"]) != (0.0, "not_measured"):
        bad.append(f"no bank lines at all: {none['value']}, {none['zero_basis']}")

    # a difference or an after_word in any words keys its figure; two that key one are refused
    l9 = pl.DataFrame({"id": ["a", "b"], "date": [D, D], "value": [100.0, 200.0]})
    r9 = pl.DataFrame({"id": ["x", "y"], "date": [D, day(1)], "value": [97.1, 195.0]})

    def fees(n1, n2):
        return resolve(l9, r9, [Rule("same day", (), (0, 0), percent=(-0.03, 0.0), difference=n1),
                                Rule("next day", (), (1, 1), percent=(-0.03, 0.0), difference=n2)])
    try:
        wb, out, L = tabs(fees("card fee (2.9%)", "bank's charge"), l9, r9)
        if not {"F.r1.match.recon.difference.card_fee_2_9",
                "F.r1.match.recon.difference.bank_s_charge"} <= set(L.entries):
            bad.append(f"the differences' figures: {sorted(k for k in L.entries if 'difference' in k)}")
    except ValueError as e:
        bad.append(f"a difference named in words with punctuation is refused: {e}")
    if not refused("name them apart", lambda: tabs(fees("fee", "fee."), l9, r9)):
        bad.append("two differences that key one figure are taken")
    # a difference name that runs an id past PROSE_MIN is refused before anything is written
    wb = Workbook()
    if not refused("Shorten the check id", lambda: tabs(
            fees("fee the card processor withheld at settlement", "bank's charge"), l9, r9, wb=wb)) or \
            wb.sheetnames != ["Sheet"]:
        bad.append("a figure id past PROSE_MIN is written")
    for word in ("En tránsito", "Outstanding (cheques)"):
        try:
            wb, out, L = tabs(res, left, right, after_word=word)
            if out["figures"][("left", word, "count")] != "F.r1.match.left.in_transit.count":
                bad.append(f"after_word {word!r} keys {out['figures'][('left', word, 'count')]}")
        except ValueError as e:
            bad.append(f"after_word {word!r} is refused: {e}")

    # dates: a Date, a Datetime or ISO text; anything else refused, never a crash
    l10 = left.with_columns(date=pl.Series(["2024-03-01 09:30", "2024-03-01", "2024-01-31"]))
    wb, out, _ = tabs(resolve(l10, right), l10, right)
    if wb[out["schedule"]]["D5"].value != D:
        bad.append(f"a date as ISO text with a time: {wb[out['schedule']]['D5'].value}")
    l11 = left.with_columns(date=pl.Series(["2024-3-1", "2024-03-01", "2024-01-31"]))
    if not refused("not ISO YYYY-MM-DD", lambda: tabs(res, l11, right)):
        bad.append("a date that is not ISO is taken, or crashes")

    def kept(ids, values, reasons):
        return pl.DataFrame({"id": ids, "value": values, "reason": reasons}, schema_overrides={"value": pl.Float64})

    # a token Excel cannot name a tab with, text a cell cannot store, a status word another
    # tab styled otherwise, others that cannot be shown, a Resolution naming a rule twice,
    # a check id outside the grammar: refused
    for words, kw in (("never holds '/'", {"token": "recon a/r"}), ("apostrophe", {"token": "'r1"}),
                      ("tab name limit", {"token": "r1 " + "x" * 20}), ("a check id", {"check": "R-1"}),
                      ("control character", {"subtitle": "Fixture\x0b"}),
                      ("are also in the left stream", {"others": kept(["a"], [1.0], ["x"])}),
                      ("present and unique", {"others": kept(["k", "k"], [1.0, 1.0], ["x", "y"])}),
                      ("no reason", {"others": kept(["k"], [1.0], [" "])}),
                      ("finite number", {"others": kept(["k"], [None], ["x"])}),
                      ("another status reads so", {"after_word": "unmatched"})):
        wb = Workbook()
        if not refused(words, lambda: tabs(res, left, right, wb=wb, **kw)) or wb.sheetnames != ["Sheet"]:
            bad.append(f"refused with `{words}`, before writing: {kw}")
    ctl = pl.DataFrame({"id": ["a\x0bb"], "date": [D], "value": [10.0]})
    if not refused("control character", lambda: tabs(resolve(ctl, r2), ctl, r2)):
        bad.append("an id holding a control character is not refused")
    register_status("Pending review", "break")
    held = dict(STATUS)
    if not refused("one word, one style", lambda: tabs(res, left, right, after_word="Pending review")) or \
            STATUS != held:
        bad.append("a status word another tab styled otherwise is restyled")
    twice = copy.copy(res)
    twice.by_rule = pl.concat([res.by_rule, res.by_rule.head(1)])
    if not refused("names a rule twice", lambda: tabs(twice, left, right)):
        bad.append("a Resolution naming a rule twice is shown, each line with both's matches")

    # a refusal the ledger makes, after every other check, leaves the ledger and the
    # workbook as they were
    wb, L = Workbook(), Ledger(run, "r3", fresh=True)
    before = (set(L.minted), dict(L.entries), list(wb.sheetnames), dict(STATUS))
    if not refused("inputs", lambda: tabs(res, left, right, wb=wb, ledger=L, inputs=[])) or \
            (set(L.minted), dict(L.entries), list(wb.sheetnames), dict(STATUS)) != before:
        bad.append("a refused call leaves tabs or figures behind")

    # the figures returned: every status's count and amount, each in the ledger
    wb, out, L = tabs(res, left, right)
    if len(out["figures"]) != 12 or not all(f in L.entries for f in out["figures"].values()):
        bad.append(f"the figures returned: {sorted(out['figures'])}")
    return bad


if __name__ == "__main__":
    raise SystemExit(main())
