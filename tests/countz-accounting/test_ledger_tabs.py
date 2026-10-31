#!/usr/bin/env python3
"""Self-test for scripts/ledger_tabs.py: Sources, Tables, Constants and Evidence written from a run's
records, and the trail down them once link_workbook.py has wired it.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_ledger_tabs.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import json
import pathlib
import subprocess
import sys
import tempfile

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import cache  # noqa: E402
import ledger_tabs  # noqa: E402
import sql_shell  # noqa: E402
from figures import Ledger  # noqa: E402
from sql import db, run_sql  # noqa: E402
from wbkit import band  # noqa: E402


def rows_by(ws, col: int = 2) -> dict:
    """{value in column B: row} below the header."""
    return {ws.cell(r, col).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, col).value}


def header(ws) -> dict:
    return {ws.cell(4, c).value: c for c in range(2, ws.max_column + 1) if ws.cell(4, c).value}


def main() -> int:
    from openpyxl import Workbook, load_workbook
    pl = cache._pl()
    bad: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        tdp = pathlib.Path(td).resolve()
        room, rd = tdp / "room", tdp / "run"
        room.mkdir()
        rd.mkdir()
        ids = [f"I-{i:02d}" for i in range(1, 13)]
        amounts = [100.0 * i for i in range(1, 13)]
        (room / "inv.csv").write_text("id,amount\n" + "".join(f"{i},{a}\n" for i, a in zip(ids, amounts)))
        (room / "terms.md").write_text("Invoices above a threshold of $500 are large.\n")
        (rd / "run.json").write_text(json.dumps({
            "checks": [{"id": "t1", "params": {}}],
            "sources": [{"id": "dataroom", "name": "room", "path": str(room), "kind": "folder"}]}))
        cache.write(rd, "inv", pl.DataFrame({"id": ids, "amount": amounts}), file="inv.csv",
                    source="dataroom", file_role="system_export", header_at="A1", rows="2:13",
                    control="amount", what="the invoices", grain="one invoice",
                    columns={"id": {"at": "A", "parse": "as written", "description": "invoice number"},
                             "amount": {"at": "B", "parse": "as written", "description": "amount",
                                        "unit": "usd"}})
        L = Ledger(rd, "t1")
        L.constants([{"name": "large_invoice", "value": 500, "unit": "usd", "file": "terms.md",
                      "source": "dataroom", "file_role": "correspondence", "locator": "line 1",
                      "quote": "a threshold of $500"},
                     {"name": "review_floor", "value": 100, "unit": "usd",
                      "basis": "the smallest invoice in E.t1.*"}])
        run_sql("CREATE TABLE t1.large AS SELECT id, amount FROM room.inv "
                "WHERE amount > const('large_invoice');"
                "COMMENT ON TABLE t1.large IS 'One row per invoice above the threshold.'", rd)
        run_sql("CREATE TABLE t1.tagged AS SELECT 'F.t1.' || id AS tag, amount FROM room.inv;"
                "COMMENT ON TABLE t1.tagged IS 'One row per invoice, tagged.'", rd)
        L.figures("SELECT 'F.t1.large' AS id, 'Large invoices' AS label, 'usd' AS unit, "
                  "sum(amount) AS value FROM t1.large")
        L.figures("SELECT 'F.t1.all' AS id, 'All invoices' AS label, 'usd' AS unit, "
                  "sum(amount) AS value FROM room.inv")
        L.figures("SELECT 'F.t1.again' AS id, 'All invoices again' AS label, 'usd' AS unit, "
                  "fig('F.t1.all') AS value")
        L.tie("T.t1.again", "All invoices, twice", "F.t1.all", "F.t1.again")
        L.write()

        wb = Workbook()
        band(wb.active, "t1 · Billing", "Test Co · 2025")
        wb.active.title = "t1 Billing"
        got = ledger_tabs.write(wb, rd, "Test Co · 2025")
        if wb.sheetnames != ["t1 Billing", "Sources", "Tables", "Constants", "Evidence"]:
            bad.append(f"tab strip: {wb.sheetnames}")
        if got["tables"] != 3 or got["constants"] != 2:
            bad.append(f"counts: {got}")
        # a tab holding one table freezes its header with the band
        for name in ("Sources", "Tables", "Constants", "Evidence"):
            want = "B5" if len(wb[name].tables) == 1 else "B4"
            if wb[name].freeze_panes != want:
                bad.append(f"{name}: frozen at {wb[name].freeze_panes}, want {want}")
        if wb["Tables"].freeze_panes != "B5":
            bad.append("Tables holds one table and freezes its header")

        # Sources: the statement, its reads, and how to reperform it
        src = wb["Sources"]
        h, at = header(src), rows_by(src)
        reads = lambda fid: [src.cell(at[fid], c).value for n, c in h.items()  # noqa: E731
                             if n.startswith("reads") and src.cell(at[fid], c).value]
        if reads("F.t1.large") != ["t1.large"] or reads("F.t1.all") != ["room.inv"] \
                or reads("F.t1.again") != ["F.t1.all"]:
            bad.append(f"reads: {reads('F.t1.large')} {reads('F.t1.all')} {reads('F.t1.again')}")
        rep = src.cell(at["F.t1.large"], h["To reperform"]).value
        if not rep.startswith("Run the expression over the tables in reads") or "gives $6,300.00." not in rep:
            bad.append(f"To reperform, a SQL figure: {rep!r}")
        diff = next(i for i in at if str(i).startswith("F.t1") and i not in
                    ("F.t1.large", "F.t1.all", "F.t1.again"))
        if "Expected:" not in src.cell(at[diff], h["To reperform"]).value:
            bad.append(f"To reperform, a tie's difference: {src.cell(at[diff], h['To reperform']).value!r}")
        if src.cell(at["F.t1.all"], h["root source"]).value != "E.t1.inv":
            bad.append(f"root source: {src.cell(at['F.t1.all'], h['root source']).value!r}")

        # Tables: what a row is, how it was made, its counts, what it reads
        tab = wb["Tables"]
        h, at = header(tab), rows_by(tab)
        big, inv = at.get("t1.large"), at.get("room.inv")
        cell = lambda r, n: tab.cell(r, h[n]).value  # noqa: E731
        if cell(big, "what one row is") != "One row per invoice above the threshold." \
                or cell(big, "made by") != "t1 Billing" or cell(big, "rows") != 7 \
                or cell(big, "totals") != "amount 6,300.00" or cell(big, "result hash") \
                or "const('large_invoice')" not in cell(big, "SQL or read") \
                or [cell(big, "reads"), cell(big, "reads 2")] != ["large_invoice", "room.inv"]:
            bad.append(f"a check's table: {[tab.cell(big, c).value for c in h.values()]}")
        con = db(rd).con
        cur = con.execute("SELECT * FROM room.inv")
        foot = sql_shell.Shell().render(cur.description, cur.fetchall()).strip().splitlines()[-1]
        if cell(inv, "result hash") not in foot or cell(inv, "rows") != 12 \
                or cell(inv, "what one row is") != "One invoice · the invoices" \
                or cell(inv, "made by") != "read from the data room" \
                or "control total amount 7,800.00" not in cell(inv, "SQL or read") \
                or cell(inv, "evidence") != "E.t1.inv":
            bad.append(f"a data-room table: {[tab.cell(inv, c).value for c in h.values()]} ({foot})")
        if "'F.' || 't1.'" not in cell(at["t1.tagged"], "SQL or read"):
            bad.append(f"a table's SQL shows its id fragment whole: {cell(at['t1.tagged'], 'SQL or read')!r}")
        if "large_invoice" in at or "review_floor" in at:
            bad.append("a constant is on the Tables tab")
        # Constants: one row per constant
        con_tab = wb["Constants"]
        at = rows_by(con_tab)
        if con_tab.cell(at["review_floor"], 7).value != "the smallest invoice in E.t1.const.large_invoice, E.t1.inv":
            bad.append(f"a constant's wildcard id: {con_tab.cell(at['review_floor'], 7).value!r}")
        if "large_invoice" not in at or con_tab.cell(at["large_invoice"], 3).value != "500" \
                or con_tab.cell(at["large_invoice"], 7).value != "a threshold of $500":
            bad.append("the constant's row")

        ev = wb["Evidence"]
        if "E.t1.inv" not in rows_by(ev):
            bad.append("the data-room read's Evidence row")

        # the trail, once linked
        path = rd / "workbook.xlsx"
        wb.save(path)
        r = subprocess.run([sys.executable, str(SCRIPTS / "link_workbook.py"), str(path), "--run-dir",
                            str(rd), "--json"], capture_output=True, text=True)
        if r.returncode != 0:
            print(r.stdout, r.stderr)
            return 1
        wb = load_workbook(path)
        src, tab = wb["Sources"], wb["Tables"]
        loc = lambda c: c.hyperlink.location if c.hyperlink else None  # noqa: E731
        s_at, t_at, e_at = rows_by(src), rows_by(tab), rows_by(wb["Evidence"])
        c_at = rows_by(wb["Constants"])
        hop = {
            "a figure's read to its table": (loc(src.cell(s_at["F.t1.large"], header(src)["reads"])),
                                             f"'Tables'!B{t_at['t1.large']}"),
            "a figure's read to its Sources row": (loc(src.cell(s_at["F.t1.again"], header(src)["reads"])),
                                                    f"'Sources'!B{s_at['F.t1.all']}"),
            "a table's read to the table it reads": (loc(tab.cell(t_at["t1.large"], header(tab)["reads 2"])),
                                                     f"'Tables'!B{t_at['room.inv']}"),
            "a constant to its row": (loc(tab.cell(t_at["t1.large"], header(tab)["reads"])),
                                      f"'Constants'!B{c_at['large_invoice']}"),
            "a data-room table to its Evidence row": (loc(tab.cell(t_at["room.inv"], header(tab)["evidence"])),
                                                      f"'Evidence'!B{e_at['E.t1.inv']}"),
        }
        for what, (got_, want) in hop.items():
            if got_ != want:
                bad.append(f"{what}: {got_}, want {want}")
        gate = subprocess.run([sys.executable, str(SCRIPTS / "check_workbook.py"), str(path), "--json"],
                              capture_output=True, text=True)
        g = json.loads(gate.stdout[gate.stdout.index("{"):])
        design = [f for f in g.get("design") or [] if f.split(":")[0] in ("Sources", "Tables", "Constants",
                                                                            "Evidence")]
        if g["links"]["link_failures"] or design:
            bad.append(f"the gate refuses the ledger tabs: {g['links']['link_failures']} {design}")

        # dead ends at the check's own gate: ids its ledger rows cite that no tab states
        # (accounts_receivable, 2026-10-22: `'(X.a7.'` in a table's statement; arr_analysis,
        # 2026-10-20: a tie id in a figure's expression; both refused only at the seal)
        run_sql("CREATE TABLE t1.flagged AS SELECT '(X.t1.' || id || ')' AS note FROM room.inv;"
                "COMMENT ON TABLE t1.flagged IS 'One row per invoice, flagged.'", rd)
        L.figures("SELECT 'F.t1.flagged' AS id, 'Flagged invoices' AS label, 'count' AS unit, "
                  "count(*) AS value, ['S.t1.election', 'T.t1.nowhere'] AS caveats "
                  "FROM t1.flagged")
        L.write()
        if got := ledger_tabs.dead_ends(rd, "t1"):
            bad.append(f"dead ends before the check has a tab: {got}")
        staged = rd / "out" / ".staging" / "t1.xlsx"
        staged.parent.mkdir(parents=True)

        def stage(*cells: str) -> list[str]:
            wb = Workbook()
            wb.active.title = "t1 Billing"
            for i, v in enumerate(cells, 5):
                wb.active.cell(i, 2, v)
            wb.save(staged)
            return ledger_tabs.dead_ends(rd, "t1")

        got = stage("T.t1.again")
        if sorted(x.split("`")[1] for x in got) != ["T.t1.nowhere", "X.t1"] \
                or not any("statement of table t1.flagged" in x for x in got):
            bad.append(f"dead ends on a staged tab stating neither: {got}")
        if got := stage("T.t1.again", "X.t1", "T.t1.nowhere: the tie that holds nowhere"):
            bad.append(f"dead ends once the tab states both: {got}")
        stage("T.t1.again")
        from sql import verify
        if not any("X.t1" in f for f in verify(rd, "t1")["failures"]):
            bad.append("sql.py verify --check does not hold the dead ends")
    if bad:
        for b in bad:
            print(f"FAIL {b}")
        print("ledger_tabs: FAIL")
        return 1
    print("ledger_tabs: ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
