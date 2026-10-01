#!/usr/bin/env python3
"""Self-test for scripts/wbkit.py.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_wbkit.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import pathlib
import sys

# The plugin under test: <repo>/countz-accounting/scripts, from <repo>/tests/countz-accounting.
SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from wbkit import S, BAND, BREAK_T, FMT_AMOUNT, FMT_COUNT, FMT_DATE, FMT_FX, FONT, TIED_T, amount, band, count, finish, header, ident, register_status, section, status, table, table_blocks, text, next_block  # noqa: E402


def main() -> int:
    """One tab through every helper; the styles register on the workbook and the sheet
    carries the § 9 settings."""
    import io
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "k1 Kit"
    band(ws, "k1 · the kit builds a tab", "Fixture · FY2026 · USD", "Every helper ran once.")
    header(ws, 4, ["id", "description", "amount", "status", "invoices", "euro"],
           ["id", "description", "amount", "status", "count", "amount"],
           currency={"euro": "eur"})
    ident(ws.cell(row=5, column=2), "F.k1.total")
    text(ws.cell(row=5, column=3), "A wrapped description long enough to need a second line "
                                   "inside a forty-two wide column.")
    amount(ws.cell(row=5, column=4), 1234567.89)
    status(ws.cell(row=5, column=5), "pass")
    count(ws.cell(row=5, column=6), 1204)
    amount(ws.cell(row=5, column=7), 5000.0)
    ident(ws.cell(row=6, column=2), "=1+2")             # an id opening with `=`: text, not a formula
    text(ws.cell(row=6, column=3), 12345)               # a number as a label: its digits, not blank
    register_status("matched", "tied")
    status(ws.cell(row=6, column=5), "matched")
    try:
        register_status("matched", "break")
        clash = False
    except ValueError:
        clash = True
    section(ws, 7, "Notes")
    import datetime as _dt
    last = table(ws, 9, [("id", "id"), ("item", "text"), ("balance", "amount", "eur"),
                         ("lines", "count"), ("rate", "fx_rate"), ("as of", "date"),
                         ("result", "status")],
                 [["X.k1.a", "First item", 1204.4, 3, 1.0679, _dt.date(2025, 9, 30), "pass"],
                  {"id": "X.k1.b", "item": "Second", "balance": -50.0, "lines": 1,
                   "rate": None, "as of": None, "result": "fail"}], primary=False)
    try:
        table(ws, 20, [("x", "money")], [])
        bad_kind = False
    except ValueError:
        bad_kind = True
    finish(ws, 5)
    # A tab of several tables: each its own Excel table, the Total row outside it, a
    # subtotal never read as a header, repeated header labels refused at authoring time.
    w2 = wb.create_sheet("k2 Tables")
    band(w2, "k2 · two tables", "Fixture · FY2026 · USD", "Each table filters alone.")
    table(w2, 4, [("id", "id"), ("item", "text"), ("amount", "amount")],
          [["A.1", "one", 10.0], ["A.2", "two", 20.0]])
    # a subtotal of labels and a formula: its figure is computed, so it is no header
    for c, v in ((2, "S.1"), (3, "Subtotal"), (4, "=SUM(D5:D6)")):
        w2.cell(row=7, column=c, value=v).style = S["Subtotal"]
    w2.cell(row=8, column=2, value="A.3")
    w2.cell(row=8, column=4, value=5.0)
    for c, v in ((2, "Total"), (4, 35.0)):
        w2.cell(row=9, column=c, value=v).style = S["Total"]
    head2 = section(w2, next_block(9), "Second table")     # two blank rows, heading, one blank
    table(w2, head2, [("id", "id"), ("reason", "text")], [["B.1", "why"]], primary=False)
    try:
        header(w2, 20, ["Amount", "amount "], ["amount", "amount"])
        dup_refused = False
    except ValueError:
        dup_refused = True
    finish(w2, 9)
    blocks = table_blocks(w2)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    back = load_workbook(buf)
    sheet = back["k1 Kit"]
    ok = (sheet.freeze_panes == "B4" and sheet.print_title_rows == "$1:$3"
          and sheet["B1"].font.name == FONT
          and sheet["B4"].fill.fgColor.rgb.endswith(BAND)
          and sheet["D5"].number_format == FMT_AMOUNT
          and sheet["F5"].number_format == FMT_COUNT and FMT_COUNT != FMT_AMOUNT
          and sheet["G4"].value == "euro (€)" and sheet["D4"].value == "amount"
          and sheet["E6"].font.color.rgb.endswith(TIED_T) and clash
          and FMT_DATE == "mmm d, yyyy"
          and sheet.row_dimensions[5].height is not None
          and last == 11 and bad_kind and sheet["D9"].value == "balance (€)"
          and sheet["D10"].number_format == FMT_AMOUNT and sheet["E10"].number_format == FMT_COUNT
          and sheet["F10"].number_format == FMT_FX and sheet["G10"].number_format == FMT_DATE
          and sheet["H11"].font.color.rgb.endswith(BREAK_T)
          and sheet["B11"].border.bottom.style is not None
          and (sheet["B6"].value, sheet["B6"].data_type) == ("=1+2", "s")
          and (sheet["C6"].value, sheet["C6"].data_type) == ("12345", "s")
          and sheet["B5"].value == "F.k1.total" and sheet["C5"].value.startswith("A wrapped"))
    t2 = back["k2 Tables"]
    refs = sorted(t.ref for t in t2.tables.values())
    ok = ok and dup_refused and (next_block(9), head2) == (12, 14) \
        and blocks == [(4, 2, 4, 8), (14, 2, 3, 15)] \
        and refs == ["B14:C15", "B4:D8"] and t2.auto_filter.ref is None \
        and [c.name for c in t2.tables[sorted(t2.tables)[0]].tableColumns] == ["id", "item", "amount"] \
        and len(sheet.tables) == 2 and len({n.casefold() for n in (*sheet.tables, *t2.tables)}) == 4 \
        and all(t.tableStyleInfo is not None and t.tableStyleInfo.name == "TableStyleLight1"
                and not t.tableStyleInfo.showRowStripes for t in t2.tables.values()) \
        and t2["B14"].fill.fgColor.rgb.endswith(BAND)
    fails = live()
    for f in fails:
        print(f"FAIL {f}")
    ok = ok and not fails
    print("wbkit: ok" if ok else "wbkit: self-check FAILED")
    return 0 if ok else 1


def live() -> list[str]:
    """The arithmetic is live (WORKBOOK.md § 7): `total` writes the formula and refuses
    one that does not reach the script's figure, `save` caches every result and writes the
    cells map, and check_workbook.py GATE 8 holds a mapped sheet's totals to formulas."""
    import json
    import subprocess
    import tempfile
    from openpyxl import Workbook, load_workbook
    from wbkit import S, amount, finish, header, save, stated, text, total
    bad: list[str] = []
    wb = Workbook()
    ws = wb.active
    ws.title = "k3 Walk"
    band(ws, "k3 · the walk foots", "Fixture · FY2026 · USD", "Opening plus sales less cash is closing.")
    header(ws, 4, ["id", "line", "FY2026", "verdict"], ["id", "description", "amount", "status"])
    lines = [("F.k3.open", "Opening receivable", 100.0, None), ("F.k3.sales", "Add: billed", 50.25, None),
             ("F.k3.cash", "Less: cash received", 40.10, None)]
    for r, (fid, lab, v, _) in enumerate(lines, start=5):
        ident(ws.cell(r, 2), fid)
        text(ws.cell(r, 3), lab)
        amount(ws.cell(r, 4), v, fid=fid)
    text(ws.cell(8, 3), "= Closing receivable")
    got = total(ws.cell(8, 4), 110.15, rows=[5, 6], less=[7], style="Subtotal", fid="F.k3.close")
    if ws["D8"].value != "=D5+D6-D7" or abs(got - 110.15) > 1e-9:
        bad.append(f"a walk's derived line: {ws['D8'].value} = {got}")
    try:
        total(ws.cell(9, 4), 999.0, rows=[5, 6])
        bad.append("a total that does not reach the script's figure must be refused")
    except ValueError as e:
        if "missing a row" not in str(e):
            bad.append(f"the refusal names the cause: {e}")
    # a conditional subtotal over the verdict column, and a copy of another check's figure
    for r, (v, verdict) in enumerate(((7.0, "supported"), (3.0, "candidate")), start=10):
        amount(ws.cell(r, 4), v, src="F.k1.total" if r == 10 else None)
        text(ws.cell(r, 5), verdict)
    total(ws.cell(12, 4), 7.0, rows=[10, 11], when=("E", "supported"), style="Subtotal")
    if not ws["D12"].value.startswith('=SUMIFS(D10:D11,$E$10:$E$11,"supported")'):
        bad.append(f"a conditional subtotal: {ws['D12'].value}")
    text(ws.cell(13, 3), "Total")
    amount(ws.cell(13, 4), 4.0, style="Total")      # an item count: stated, not a sum
    stated(ws.cell(13, 4))
    finish(ws, 13)
    try:
        amount(ws.cell(20, 4), 1.0, fid="k3 total")
        bad.append("a figure id outside the grammar must be refused")
    except ValueError:
        pass
    ws["D20"].value = None
    with tempfile.TemporaryDirectory() as td:
        path = pathlib.Path(td) / "k3_walk.xlsx"
        side = save(wb, path)
        back = load_workbook(path, data_only=True)["k3 Walk"]
        if (back["D8"].value, back["D12"].value) != (110.15, 7.0):
            bad.append(f"save caches every result: D8 {back['D8'].value}, D12 {back['D12'].value}")
        m = json.loads(side.read_text())["sheets"]["k3 Walk"]
        if m.get("D5", {}).get("fid") != "F.k3.open" or m.get("D10", {}).get("src") != "F.k1.total" \
                or m.get("D8", {}).get("formula") != "=D5+D6-D7" or not m.get("D13", {}).get("stated"):
            bad.append(f"the cells map: {m}")
        gate = [sys.executable, str(SCRIPTS / "check_workbook.py"), str(path), "--json"]
        rep = json.loads(subprocess.run(gate, capture_output=True, text=True).stdout)
        if rep.get("formulas"):
            bad.append(f"a mapped sheet with live totals passes GATE 8: {rep['formulas']}")
        # the same sheet with the closing line typed as a value is refused
        typed = load_workbook(path)
        typed["k3 Walk"]["D8"].value = 110.15
        typed.save(path)
        rep = json.loads(subprocess.run(gate, capture_output=True, text=True).stdout)
        if not any("typed as a value" in f and "D8" in f for f in rep.get("formulas", [])):
            bad.append(f"a total typed as a value on a mapped sheet must be refused: {rep.get('formulas')}")
        # a stale result: the formula computes one figure, the file stores another
        save(wb, path)
        import formula as fx
        fx.cache_results(path, {"k3 Walk": {"D8": 999.0}})
        rep = json.loads(subprocess.run(gate, capture_output=True, text=True).stdout)
        if not any("stored result" in f for f in rep.get("formulas", [])):
            bad.append(f"a stale result must be refused: {rep.get('formulas')}")
    return bad


if __name__ == "__main__":
    raise SystemExit(main())
