#!/usr/bin/env python3
"""Self-test for scripts/link_workbook.py § 8, the live arithmetic.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_link_workbook.py`
from the repository root; it lives outside the plugin so it never ships.

A run of four check tabs: two written before the cells maps existed (retrofitted on
evidence alone) and two with maps (wired by figure id). After the pass every total is a
formula over its rows, every copy a reference to its producer's cell, a copy that
disagrees with its producer is reported and left, every formula's result is cached, and
check_workbook.py GATE 8 has nothing to refuse.
"""
from __future__ import annotations

import json
import pathlib
import subprocess
import sys
import tempfile

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from wbkit import amount, band, header, ident, text  # noqa: E402

PERIOD = "September 2025"


def tab(wb, name, rows):
    """A check tab: header on row 4, then (id, label, value, style) rows."""
    ws = wb.create_sheet(name)
    band(ws, f"{name.split()[0]} · fixture", f"Fixture · {PERIOD} · USD", "A fixture tab.")
    header(ws, 4, ["id", "line", PERIOD], ["id", "description", "amount"])
    for r, (fid, label, v, style, kw) in enumerate(rows, start=5):
        ident(ws.cell(r, 2), fid)
        text(ws.cell(r, 3), label, style or "Body")
        amount(ws.cell(r, 4), v, style=style, **(kw or {}))
    return ws


def main() -> int:
    from openpyxl import Workbook, load_workbook
    bad: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        rd = pathlib.Path(td) / "run"
        (rd / "out" / "tabs").mkdir(parents=True)
        (rd / "workpapers").mkdir()
        checks = ["a1_base", "b1_walk", "m1_producer", "c1_mapped"]
        (rd / "run.json").write_text(json.dumps({"checks": [{"id": c, "params": {}} for c in checks]}))

        # the assembled workbook: two tabs written before the maps (values only), two
        # with maps, and a Sources tab
        wb = Workbook()
        wb.remove(wb.active)
        tab(wb, "a1 Base", [("F.a1.x", "Measured figure", 1234.56, None, None),
                            ("F.a1.y", "Another figure", 77.7, None, None),
                            ("F.a1.z", "A third figure", 0.04, None, None),
                            (None, "Total", 1312.30, "Total", None)])
        tab(wb, "b1 Walk", [("F.b1.open", "Opening, as a1 measured it", 1234.56, None, None),
                            ("F.b1.add", "Add: billed", 10.0, None, None),
                            ("F.b1.less", "Less: cash received", 4.56, None, None),
                            (None, "= Closing", 1240.0, None, None)])
        tab(wb, "m1 Producer", [("F.m1.x", "Produced figure", 500.25, None, None)])
        tab(wb, "c1 Mapped", [("F.c1.copy", "Copied figure", 500.25, None, None),
                              ("F.c1.wrong", "Copied, but wrong", 501.0, None, None),
                              ("F.c1.lost", "Copied from nowhere", 9.0, None, None)])
        maps = {"m1_producer": {"m1 Producer": {"D5": {"fid": "F.m1.x"}}},
                "c1_mapped": {"c1 Mapped": {"D5": {"src": "F.m1.x"}, "D6": {"src": "F.m1.x"},
                                            "D7": {"src": "F.zz.none"}}}}
        for name, sheets in maps.items():           # one map per placed tab, as place_tab leaves them
            (rd / "out" / "tabs" / f"{name}.cells.json").write_text(
                json.dumps({"schema": "countz-accounting/cells@1", "sheets": sheets}))
        src_tab = wb.create_sheet("Sources")
        band(src_tab, "Sources", "Fixture", None)
        path = rd / "out" / ".staging" / "workbook.xlsx"
        path.parent.mkdir(parents=True)
        wb.save(path)

        r = subprocess.run([sys.executable, str(SCRIPTS / "link_workbook.py"), str(path), "--run-dir", str(rd),
                            "--json"], capture_output=True, text=True)
        if r.returncode != 0:
            print(r.stdout, r.stderr)
            return 1
        rep = json.loads(r.stdout)
        f = load_workbook(path)
        v = load_workbook(path, data_only=True)
        want = {("a1 Base", "D8"): "=SUM(D5:D7)",                    # a total: its rows
                ("b1 Walk", "D5"): "='a1 Base'!D5",                  # a copy, retrofitted by value
                ("b1 Walk", "D8"): "=D5+D6-D7",                      # a walk, its Less line deducted
                ("c1 Mapped", "D5"): "='m1 Producer'!D5"}            # a copy, wired by its map
        for (sh, co), formula in want.items():
            if f[sh][co].value != formula:
                bad.append(f"{sh}!{co}: {f[sh][co].value!r}, want {formula}")
        for (sh, co), value in {("a1 Base", "D8"): 1312.30, ("b1 Walk", "D8"): 1240.0,
                                ("b1 Walk", "D5"): 1234.56, ("c1 Mapped", "D5"): 500.25}.items():
            if v[sh][co].value is None or abs(v[sh][co].value - value) > 1e-9:
                bad.append(f"{sh}!{co} caches {v[sh][co].value!r}, want {value}")
        if f["c1 Mapped"]["D6"].value != 501.0 or not any("c1 Mapped!D6" in p for p in rep["copy_problems"]):
            bad.append(f"a copy that disagrees with its producer is left and reported: {rep['copy_problems']}")
        if not any("F.zz.none" in p for p in rep["copy_problems"]):
            bad.append("a copy of a figure no map places is reported")
        if f["a1 Base"]["D6"].data_type == "f":
            bad.append("a figure with no earlier match stays a value")
        if (rep["copies_wired"], rep["copies_retrofitted"], rep["totals_made_live"]) != (1, 1, 2):
            bad.append(f"counts: {rep['copies_wired']} wired, {rep['copies_retrofitted']} retrofitted, "
                       f"{rep['totals_made_live']} totals")
        gate = subprocess.run([sys.executable, str(SCRIPTS / "check_workbook.py"), str(path), "--run-dir",
                               str(rd), "--json"], capture_output=True, text=True)
        g = json.loads(gate.stdout[gate.stdout.index("{"):])
        if g["uncached"] or g.get("formulas"):
            bad.append(f"GATE 1 and 8 pass the linked workbook: {g['uncached']} uncached, {g.get('formulas')}")
    if bad:
        for b in bad:
            print(f"FAIL {b}")
        print("link_workbook: FAIL")
        return 1
    print("link_workbook: ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
