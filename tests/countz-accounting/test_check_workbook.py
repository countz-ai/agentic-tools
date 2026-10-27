#!/usr/bin/env python3
"""Self-test for scripts/check_workbook.py's cut-text, frozen-pane and tab-name rules.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_check_workbook.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import io
import json
import pathlib
import sys
import tempfile
import zipfile

from openpyxl import Workbook

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import check_workbook  # noqa: E402


def cut_fails(texts: list[str]) -> list[str]:
    """The cut-text failures audit_design reports for a tab holding `texts` in column B."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sources"
    for i, t in enumerate(texts, 5):
        ws.cell(i, 2, t)
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        return [f for f in check_workbook.audit_design(z) if "text cut mid-word" in f]


def pane_fails(tables: int, pane: str) -> list[str]:
    """The frozen-pane failures both gates report for a ledger tab of `tables` tables frozen
    at `pane`."""
    from wbkit import band, bind_styles, excel_tables, header, ident, next_block, section
    wb = Workbook()
    bind_styles(wb)
    ws = wb.active
    ws.title = "Sources"
    band(ws, "Sources", "Test Co · 2025")
    r = 4
    for n in range(tables):
        if n:
            r = section(ws, next_block(r), f"Part {n + 1}")
        header(ws, r, ["id", "label"], ["id", "description"])
        for k in range(1, 3):
            ident(ws.cell(r + k, 2), f"F.t{n}.x{k}")
            ws.cell(r + k, 3, "a figure")
        r += 2
    excel_tables(ws)
    ws.freeze_panes = pane
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        design = [f for f in check_workbook.audit_design(z) if "freeze panes" in f]
        links = [f for f in check_workbook.audit_links(z)["link_failures"] if "frozen pane" in f]
    return design + links


def main() -> int:
    bad: list[str] = []
    # a whole id that opens a longer id is a key, not a cut claim (arr_analysis, 2026-10-09:
    # 28 cohort-count ids ending `.n` beside the `.nrr` ids blocked the seal)
    ids = ["F.g5.rbd.industry.financial_services.fy2022.n",
           "F.g5.rbd.industry.financial_services.fy2022.nrr"]
    if got := cut_fails(ids):
        bad.append(f"an id that opens a longer id is read as cut: {got}")
    # a sentence cut to a length beside its whole statement still fails
    whole = "The customer was billed after termination and the balance has since removed the credit."
    if not cut_fails([whole, whole[:60]]):
        bad.append("a sentence cut mid-word beside its whole statement is taken")
    # a title that repeats the check id the token spells (arr_analysis, 2026-10-16:
    # `g6 stale arr Stale ARR`); the id alone, or a one-check family's token and title, pass
    checks = [{"id": "g6_stale_arr", "params": {"family": "g6"}},
              {"id": "g6_customer_concentration", "params": {"family": "g6"}},
              {"id": "g4_cube", "params": {"family": "g4"}}]
    _, fails = check_workbook.tab_owners(checks, ["g6 stale arr Stale ARR"])
    if not any("repeats the check id" in f for f in fails):
        bad.append(f"a title repeating the check id is taken: {fails}")
    _, fails = check_workbook.tab_owners(
        checks, ["g6 Stale ARR", "g6 Customer concentration", "g4 Customer revenue cube"])
    if fails:
        bad.append(f"a tab named by its id, or by its family token and title, is refused: {fails}")
    # a check's file staged before assembly is held to names that pair with that check
    # (revenue_leak, 2026-10-23: ten r6 tabs refused at the seal), and `&` in a name reads
    # as `&` (arr_analysis, 2026-10-20: `g6 Contract term & renewals` read as `&amp;`)
    with tempfile.TemporaryDirectory() as td:
        rd = pathlib.Path(td)
        (rd / "run.json").write_text(json.dumps({"checks": [
            {"id": "r6_deductions", "params": {"family": "r6"}},
            {"id": "r6_discounts", "params": {"family": "r6"}},
            {"id": "r6_contract_term_renewals", "params": {"family": "r6"}},
            {"id": "r0_population", "params": {"family": "r0"}}]}))

        def staged(check: str, *names: str) -> list[str]:
            wb = Workbook()
            wb.active.title = names[0]
            for n in names[1:]:
                wb.create_sheet(n)
            buf = io.BytesIO()
            wb.save(buf)
            with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
                return check_workbook.audit_order(z, rd, check)

        for check, names, want in (
                ("r6_deductions", ["r6 Customer deductions"], "opens with `r6 deductions`"),
                ("r6_deductions", ["r6 discounts"], "pairs with r6_discounts"),
                ("r6_deductions", ["r6 deductions"], None),
                ("r6_contract_term_renewals", ["r6 Contract term & renewals"], None),
                ("r0_population", ["r0 Population and scope"], None),
                ("workbook", ["Any name"], None)):
            got = staged(check, *names)
            if (want is None and got) or (want and not any(want in f for f in got)):
                bad.append(f"{check} staged as {names}: {got or 'taken'}")
    # a tab of one table freezes its header with the band; a tab of several freezes the band
    for tables, pane, ok in ((1, "B5", True), (1, "B4", False), (2, "B4", True), (2, "B5", False)):
        got = pane_fails(tables, pane)
        if bool(got) == ok:
            bad.append(f"{tables} table(s) frozen at {pane}: {'refused' if got else 'taken'} {got}")
    for b in bad:
        print("FAIL", b)
    print("check_workbook: " + ("FAIL" if bad else "ok"))
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
