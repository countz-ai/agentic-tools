#!/usr/bin/env python3
"""Self-test for scripts/render.py: a check's template rendered to its md, its summary and
its tab, and the template's refusals.

Run by check-plugin.py (8z) as `uv run --project <plugin> python3 tests/countz-accounting/test_render.py`
from the repository root; it lives outside the plugin so it never ships.
"""
from __future__ import annotations

import json
import pathlib
import subprocess
import sys

SCRIPTS = pathlib.Path(__file__).resolve().parents[2] / "countz-accounting" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import cache  # noqa: E402
import render  # noqa: E402
from figures import Ledger  # noqa: E402
from items import write_items  # noqa: E402
from sql import run_sql  # noqa: E402

TEMPLATE = """---
tab: t1 Billing
title: t1 · Billing agreed to the ledger
subtitle: Test Co · May and June 2025 · $ · tolerance: display rounding
summary: Billing is {F.t1.billed.total} over two months.
statuses:
  Billed: tied
---
# t1_billing

## Billing by month

Billing is {F.t1.billed.total}.

```sql primary
SELECT month AS "Month", 'Billed in ' || month AS "What", cell('F.t1.billed.' || month) AS "Billed",
       'Billed' AS "Status", 'From the invoice register.' AS "Basis", NULL AS _row, month AS _seq
FROM (SELECT DISTINCT month FROM room.inv)
UNION ALL
SELECT NULL, 'Billed, both months', cell('F.t1.billed.total'), NULL, NULL, 'total', 'z'
ORDER BY _seq
```

## Invoices

```sql
-- heading: Invoices by kind
SELECT id AS "Invoice", kind AS "Kind", amount AS "Amount", start AS "Service start"
FROM room.inv ORDER BY id
```

## Walk

```sql
SELECT 'June' AS "Line", cell('F.t1.billed.2025-06') AS "Amount", NULL AS _row, 1 AS _seq
UNION ALL SELECT 'Less: May', cell('F.t1.billed.2025-05'), 'less', 2
UNION ALL SELECT 'June over May', cell('F.t1.net'), 'subtotal', 3
ORDER BY _seq
```

```sql
-- kinds: Count=count
SELECT id AS "Invoice", CASE kind WHEN 'INV' THEN 'Billed' ELSE 'Credit' END AS "Status",
       amount AS "Amount", NULL::DOUBLE AS "Count", NULL AS _row, NULL AS _when, NULL AS _stated, id AS _seq
FROM room.inv
UNION ALL
SELECT NULL, 'Billed', fig('F.t1.inv'), fig('F.t1.inv_count'), 'total', 'Billed', 'Count', 'z'
ORDER BY _seq
```

## Ties

{ties}

```sql
SELECT 'T.t1.ledger' AS "Tie", tie("Tie").status AS "Result", cell(tie("Tie").difference) AS "Difference",
       cell('F.t0.opening') AS "Opening", 'The ledger reads the same.' AS "Basis"
```

{manifests}

## Notes

- `F.t1.billed.total`: Both months, {F.t1.billed.total}.

## To reperform

1. Sum the invoice amounts by month.

## Summary

We agreed billing of {F.t1.billed.total} to the ledger.
"""


def main() -> int:
    import tempfile
    pl = cache._pl()
    bad: list[str] = []

    def refused(label, text, want, run_dir):
        (run_dir / "workpapers" / "t1.tpl.md").write_text(text)
        try:
            render.render(run_dir, "t1")
        except (ValueError, KeyError, FileNotFoundError) as exc:
            if want not in str(exc):
                bad.append(f"{label}: refused, but the message lacks {want!r}: {exc}")
            return
        bad.append(f"{label}: not refused")

    with tempfile.TemporaryDirectory() as td:
        tdp = pathlib.Path(td).resolve()
        room, rd = tdp / "room", tdp / "run"
        room.mkdir()
        rd.mkdir()
        (room / "inv.csv").write_text("id,month,amount,kind,start\nI-1,2025-05,100,INV,2025-01-01\n"
                                      "I-2,2025-05,50,CM,2025-01-01\nI-3,2025-06,200,INV,2025-03-01\n")
        (rd / "run.json").write_text(json.dumps({"sources": [
            {"id": "dataroom", "name": "room", "path": str(room), "kind": "folder"}]}))
        import datetime as dt
        inv = pl.DataFrame({"id": ["I-1", "I-2", "I-3"], "month": ["2025-05", "2025-05", "2025-06"],
                            "amount": [100.0, 50.0, 200.0], "kind": ["INV", "CM", "INV"],
                            "start": [dt.date(2025, 1, 1), dt.date(2025, 1, 1), dt.date(2025, 3, 1)]})
        col = lambda at, d, u=None: {"at": at, "parse": "as written", "description": d,  # noqa: E731
                                     **({"unit": u} if u else {})}
        cache.write(rd, "inv", inv, file="inv.csv", source="dataroom", file_role="system_export",
                    header_at="A1", rows="2:4", control="amount", what="the invoices",
                    grain="one invoice",
                    columns={"id": col("A", "invoice number"), "month": col("B", "billing month"),
                             "amount": col("C", "invoice amount", "usd"), "kind": col("D", "INV or CM"),
                             "start": col("E", "service start")})
        write_items(rd, "t1", "invoices", about="One row per invoice.", sql="SELECT id AS item_id, 'book' AS side, month AS period, "
                    "amount, 'usd' AS unit FROM room.inv", sides=["book"], keys=["item_id"],
                    closes="F.t1.billed.total")
        L0 = Ledger(rd, "t0")
        L0.figures("SELECT 'F.t0.opening' AS id, 'Opening' AS label, 'usd' AS unit, "
                   "sum(amount) AS value FROM room.inv WHERE kind = 'CM'")
        L0.write()
        L = Ledger(rd, "t1")
        L.figures("SELECT 'F.t1.billed.' || month AS id, 'Billed in ' || month AS label, 'usd' AS unit, "
                  "sum(amount) AS value FROM room.inv GROUP BY month")
        L.figures("SELECT 'F.t1.billed.total' AS id, 'Billed' AS label, 'usd' AS unit, "
                  "sum(amount) AS value FROM room.inv")
        L.figures("SELECT 'F.t1.ledger' AS id, 'Ledger' AS label, 'usd' AS unit, "
                  "fig('F.t1.billed.2025-05') + fig('F.t1.billed.2025-06') AS value")
        L.figures("SELECT 'F.t1.net' AS id, 'June over May' AS label, 'usd' AS unit, "
                  "fig('F.t1.billed.2025-06') - fig('F.t1.billed.2025-05') AS value")
        L.figures("SELECT 'F.t1.inv' AS id, 'Invoices billed' AS label, 'usd' AS unit, "
                  "sum(amount) AS value FROM room.inv WHERE kind = 'INV'")
        L.figures("SELECT 'F.t1.inv_count' AS id, 'Invoices billed' AS label, 'count' AS unit, "
                  "count(*) AS value FROM room.inv WHERE kind = 'INV'")
        L.tie("T.t1.ledger", "Billing to the ledger", "F.t1.billed.total", "F.t1.ledger")
        L.write()
        (rd / "workpapers" / "t1.tpl.md").write_text(TEMPLATE)

        # the macros the template reads with
        got = run_sql("SELECT tie('T.t1.ledger').status AS s, cell('F.t1.ledger').value AS v", rd)
        if got.row(0) != ("pass", 350.0):
            bad.append(f"tie() and cell(): {got.row(0)}")

        out = render.render(rd, "t1")
        md = out["md"].read_text()
        for want in ("Billing is $350.", "| 2025-05 | Billed in 2025-05 | $150 | Billed |",
                     "| **Billed, both months** | **$350** |", "Statement:\n\n```sql", "| I-1 | INV | $100 | Jan 1, 2025 |",
                     "T.t1.ledger Billing to the ledger", "| T.t1.ledger | pass | $0 |",
                     "**Item table** `checks/t1-invoices.csv`", "Closes to: F.t1.billed.total",
                     "- `F.t1.billed.total`: Both months, $350.", "1. Sum the invoice amounts"):
            if want not in md:
                bad.append(f"md lacks {want!r}")
        if "agreed billing" in md or "## Summary" in md:
            bad.append("the summary section is in the check's md")
        if out["summary"].read_text() != "We agreed billing of $350 to the ledger.\n":
            bad.append(f"summary: {out['summary'].read_text()!r}")

        from openpyxl import load_workbook
        ws = load_workbook(out["tab"]).active
        cells = json.loads(out["tab"].with_name("t1.cells.json").read_text())["sheets"]["t1 Billing"]
        if (ws.title, ws["B1"].value, ws["B3"].value) != ("t1 Billing", "t1 · Billing agreed to the ledger",
                                                          "Billing is $350 over two months."):
            bad.append(f"band: {ws.title!r} {ws['B1'].value!r} {ws['B3'].value!r}")
        if [ws.cell(row=r, column=2).value for r in (4, 5, 6)] != ["Month", "2025-05", "2025-06"]:
            bad.append("the primary table is not under the band on row 4")
        if cells.get("D5", {}).get("fid") != "F.t1.billed.2025-05" or ws["D5"].value != 150:
            bad.append(f"a figure cell: {ws['D5'].value} {cells.get('D5')}")
        if ws["D7"].value != "=SUM(D5:D6)" or cells.get("D7", {}).get("fid") != "F.t1.billed.total":
            bad.append(f"the total row: {ws['D7'].value} {cells.get('D7')}")
        formulas = {c.coordinate: c.value for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")}
        if not any("-" in f and "SUMIFS" not in f for f in formulas.values()):
            bad.append(f"a less line is not deducted: {formulas}")
        if not any("SUMIFS" in f and '"Billed"' in f for f in formulas.values()):
            bad.append(f"a _when total is not a conditional sum: {formulas}")
        if not any(e.get("stated") for e in cells.values()):
            bad.append("a _stated figure is not declared stated")
        copied = [e for e in cells.values() if e.get("src") == "F.t0.opening"]
        if len(copied) != 1 or any(e.get("fid") == "F.t0.opening" for e in cells.values()):
            bad.append(f"another check's figure is not written as a copy: {copied}")
        if ws["E5"].style != "cz_tied":
            bad.append(f"a registered status word: {ws['E5'].style}")
        heads = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)
                 if ws.cell(row=r, column=2).style == "cz_section"]
        if heads != ["Invoices by kind", "Walk", "Ties", "Notes", "To reperform"]:
            bad.append(f"sections: {heads}")
        if len(ws.tables) != 6:
            bad.append(f"Excel tables: {len(ws.tables)}")
        g = subprocess.run([sys.executable, str(SCRIPTS / "check_workbook.py"), str(out["tab"]),
                            "--run-dir", str(rd)], capture_output=True, text=True)
        if g.returncode != 0:
            bad.append(f"check_workbook refuses the rendered tab: {(g.stdout + g.stderr)[-400:]}")

        # a recon check passes the workbook holding its match tabs
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "t1 Match summary"
        render.render(rd, "t1", wb=wb)
        if load_workbook(out["tab"]).sheetnames != ["t1 Match summary", "t1 Billing"]:
            bad.append("render(wb=): the check tab does not follow the match tabs")

        primary = TEMPLATE.index("```sql primary")
        refused("no primary table", TEMPLATE.replace("```sql primary", "```sql"), "primary", rd)
        refused("two primary tables", TEMPLATE.replace("```sql\n-- heading", "```sql primary\n-- heading"),
                "2 primary", rd)
        refused("an unknown option", TEMPLATE.replace("-- heading:", "-- title:"), "option", rd)
        refused("an unknown kind", TEMPLATE.replace("-- heading: Invoices by kind",
                                                    "-- kinds: Kind=word"), "kind 'word'", rd)
        refused("a _row word", TEMPLATE.replace("'total', 'z'", "'grand', 'z'"), "_row", rd)
        refused("a figure no ledger holds", TEMPLATE.replace("cell('F.t1.billed.total')",
                                                             "cell('F.t1.nothing')"), "F.t1.nothing", rd)
        refused("a reference no ledger holds", TEMPLATE.replace("Billing is {F.t1.billed.total}.",
                                                                "Billing is {F.t1.none}."), "F.t1.none", rd)
        refused("a statement that writes", TEMPLATE[:primary] + "```sql primary\nCREATE TABLE t1.x AS "
                "SELECT 1 AS a\n```\n" + TEMPLATE[TEMPLATE.index("## Invoices"):], "SELECT", rd)
        refused("a note with no id", TEMPLATE.replace("- `F.t1.billed.total`: Both", "- Both"),
                "a note is", rd)
        refused("a total the rows do not reach", TEMPLATE.replace("FROM (SELECT DISTINCT month FROM room.inv)",
                "FROM (SELECT DISTINCT month FROM room.inv WHERE month = '2025-05')"), "not the", rd)
        refused("no tab name", TEMPLATE.replace("tab: t1 Billing\n", ""), "`tab`", rd)
    if bad:
        print("render.py self-check FAILED:")
        for b in bad:
            print(f"  - {b}")
        return 1
    print("render.py self-check: ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
