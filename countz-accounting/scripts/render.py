#!/usr/bin/env python3
"""A check's page, written once: `workpapers/<check>.tpl.md` rendered to `checks/<check>.md`,
`checks/<check>-summary.md` and the staged tab `out/.staging/<check>.xlsx` (WORKBOOK.md § 7).

    uv run --project ${CLAUDE_PLUGIN_ROOT} python3 ${CLAUDE_PLUGIN_ROOT}/scripts/render.py <run_dir> <check>

renders the three files, then gates and places the tab (`step_record.place_tab`); exit 1
names what a gate refused. `--no-place` stops after staging. A recon check writes its match
tabs first and passes the workbook: `render(run_dir, check, wb=wb)`, then `place_tab`.

The template is Markdown:

- Front matter between `---` lines, one `key: value` per line: `tab` (the tab name,
  WORKBOOK.md § 2), `title`, `subtitle` and `summary` (the tab's rows 1 to 3); optionally
  `currency` (the unit of a money column that is not a figure, default `usd`) and
  `statuses:` with one indented `Word: tied | review | break | note` line per status word.
- `{F.id}` and `{F.id:deck}` anywhere, in prose and in a query's text values.
- A fenced `sql` block: one SELECT, a table in the md and on the tab, the column aliases
  its headers. `cell('F.id')` is a figure: the cell shows its value in its unit and
  carries its id. Otherwise the type sets the kind: a number with decimals is money in
  `currency`, a whole number a count, a date a date. The first column is the id column; a
  text column headed Status, Verdict, Result or Disposition is the status column; a last
  text column is the prose column. `sql primary` marks the primary table, which the tab
  puts under the band. The md shows each statement under its table.
- Hidden columns, named with a leading `_`, are not shown:
  - `_row` marks a row `subtotal` (the body rows since the last subtotal), `total` (every
    body row), `derived` (a walk's line: the line above it and the body rows since) or
    `less` (a body row the sums after it deduct). Each number on a summing row is a
    formula over those rows, refused when they do not add to it. A `cell()` there keeps
    its id; a `fig('F.x')` in a column of plain amounts does not.
  - `_when` on a summing row adds only the rows whose status reads that word.
  - `_stated` on a summing row names the columns it states rather than sums (a count).
  - any other (`_seq`) sorts.
- Options, as comment lines opening a block: `-- kinds: Header=kind, ...` (`wbkit.KINDS`)
  and `-- heading: Text`, the tab's heading when the md heading above does not fit.
- `{ties}` on its own line: the check's tie schedule. A table of some ties, or with more
  columns, is a query over `tie('T.x')` (`tie(t).status`, `cell(tie(t).a)`).
- `{manifests}` on its own line: each item table's manifest block (EVIDENCE.md § 5), md
  only.
- A list under a heading `Notes` (`` - `F.x`: the note ``) and a numbered list under
  `To reperform` are those blocks on the tab, last.
- A `## Summary` section is `checks/<check>-summary.md`, and is not in the check's md.

Other prose is in the md only. Refused, writing nothing: a reference no ledger holds, a
block that is not one SELECT, no primary table or two, an unknown option or kind, a `_row`
word other than those four, a `_when` on a table with no status column, a note with no id.

Example, a template's primary table with its total:

    ```sql primary
    SELECT contract_id AS "Contract", cell('F.m5.writeoff.' || contract_id) AS "Written off",
           verdict AS "Verdict", basis AS "Basis", NULL AS _row
    FROM m5_amortization.writeoffs
    UNION ALL
    SELECT 'Written off in FY2025', cell('F.m5.writeoff.total'), NULL, NULL, 'total'
    ORDER BY _row NULLS FIRST, "Contract"
    ```
"""
from __future__ import annotations

import argparse
import datetime as dt
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import figures  # noqa: E402
import items  # noqa: E402
import style  # noqa: E402
import wbkit  # noqa: E402

STATUS_HEADERS = {"status", "verdict", "result", "disposition"}
FENCE = re.compile(r"^```\s*(\w*)\s*(.*?)\s*$")
HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
NOTE = re.compile(r"^\s*[-*]\s+`?([A-Z]{1,2}\.[A-Za-z0-9_.-]*[A-Za-z0-9])`?\s*[:—-]\s*(.+)$")
STEP = re.compile(r"^\s*\d+\.\s+\S")
FIGURE_FORMAT = {"pct": wbkit.FMT_PCT, "rate": wbkit.FMT_RATE, "fx_rate": wbkit.FMT_FX,
                 "days": wbkit.FMT_DAYS, "ratio": wbkit.FMT_DAYS, "quantity": wbkit.FMT_DAYS}
KIND_FORMAT = {"amount": wbkit.FMT_AMOUNT, "cents": wbkit.FMT_CENTS, "pct": wbkit.FMT_PCT,
               "rate": wbkit.FMT_RATE, "fx_rate": wbkit.FMT_FX, "days": wbkit.FMT_DAYS}
RIGHT = wbkit.Alignment(horizontal="right", vertical="top")
NUMERIC = ("figure", "amount", "cents", "count", "pct", "rate", "fx_rate", "days")


class TemplateError(ValueError):
    pass


# --- reading the template ------------------------------------------------------------------
def _front(text: str, where: str) -> tuple[dict, list[str]]:
    """`key: value` lines, the value the rest of the line as written; `statuses:` takes
    indented `Word: kind` lines. Not YAML: a subtitle holds colons and a summary opens
    with `{F.`."""
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        raise TemplateError(f"{where}: opens with front matter between `---` lines")
    try:
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
    except StopIteration:
        raise TemplateError(f"{where}: the front matter has no closing `---`") from None
    meta: dict = {}
    for n, line in enumerate(lines[1:end], start=2):
        if not line.strip():
            continue
        key, sep, value = line.strip().partition(":")
        if not sep:
            raise TemplateError(f"{where} line {n}: `key: value`")
        if line[:1].isspace():
            if not isinstance(meta.get("statuses"), dict):
                raise TemplateError(f"{where} line {n}: an indented line is a status word "
                                    f"under `statuses:`")
            meta["statuses"][key.strip()] = value.strip()
        elif key.strip() == "statuses" and not value.strip():
            meta["statuses"] = {}
        else:
            meta[key.strip()] = value.strip()
    unknown = set(meta) - {"tab", "title", "subtitle", "summary", "currency", "statuses"}
    if unknown:
        raise TemplateError(f"{where}: front matter {sorted(unknown)}: tab, title, subtitle, "
                            f"summary, currency, statuses")
    for k in ("tab", "title", "subtitle"):
        if not isinstance(meta.get(k), str) or not meta[k].strip():
            raise TemplateError(f"{where}: front matter `{k}` is required")
    return meta, lines[end + 1:]


def _options(query: str, where: str) -> tuple[dict, str]:
    opts, lines = {}, query.splitlines()
    i = 0
    while i < len(lines) and (not lines[i].strip() or lines[i].lstrip().startswith("--")):
        m = re.match(r"\s*--\s*(\w+)\s*:\s*(.*)$", lines[i])
        if m:
            k, v = m.group(1), m.group(2).strip()
            if k == "kinds":
                kinds = {}
                for part in filter(None, (p.strip() for p in v.split(","))):
                    h, _, kind = part.partition("=")
                    if kind.strip() not in wbkit.KINDS:
                        raise TemplateError(f"{where}: kind {kind.strip()!r} for {h.strip()!r}: "
                                            f"one of {', '.join(wbkit.KINDS)}")
                    kinds[h.strip()] = kind.strip()
                opts["kinds"] = kinds
            elif k == "heading":
                opts["heading"] = v
            else:
                raise TemplateError(f"{where}: option `{k}`: kinds or heading")
        i += 1
    body = lines[i:]
    return opts, "\n".join(body).strip()


def parse(text: str, where: str = "template") -> tuple[dict, list[tuple]]:
    """The front matter and the template's parts in order: `("line", text, section)`,
    `("heading", level, text)`, `("sql", primary, options, query, heading)`,
    `("ties", heading)`, `("manifests",)`, `("note", id, text)`, `("step", text)` and
    `("summary", text)`. `section` is the heading the line sits under."""
    meta, lines = _front(text, where)
    parts, heading, h2, i = [], None, None, 0
    while i < len(lines):
        line = lines[i]
        m = FENCE.match(line)
        if m and m.group(1).lower() == "sql":
            flags = m.group(2).split()
            if set(flags) - {"primary"}:
                raise TemplateError(f"{where} line {i + 1}: ```sql takes `primary` only")
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            if j == len(lines):
                raise TemplateError(f"{where} line {i + 1}: the sql block is not closed")
            opts, query = _options("\n".join(lines[i + 1:j]), f"{where} line {i + 1}")
            if h2 == "summary":
                raise TemplateError(f"{where} line {i + 1}: the summary holds no table")
            parts.append(("sql", "primary" in flags, opts, query, opts.get("heading", heading)))
            i = j + 1
            continue
        h = HEADING.match(line)
        if h:
            level, heading = len(h.group(1)), h.group(2)
            if level <= 2:
                h2 = heading.strip().lower()
            if not (level == 2 and h2 == "summary"):
                parts.append(("heading", level, heading))
        elif line.strip() == "{ties}":
            parts.append(("ties", heading))
        elif line.strip() == "{manifests}":
            parts.append(("manifests",))
        elif h2 == "summary":
            parts.append(("summary", line))
        else:
            sec = (heading or "").strip().lower()
            if sec == "notes" and line.strip():
                n = NOTE.match(line)
                if not n:
                    raise TemplateError(f"{where} line {i + 1}: a note is `- \\`<id>\\`: <the "
                                        f"note>`, the id it qualifies first")
                parts.append(("note", n.group(1), n.group(2)))
            elif sec == "to reperform" and STEP.match(line):
                parts.append(("step", line.strip()))
            else:
                parts.append(("line", line, heading))
        i += 1
    primaries = [p for p in parts if p[0] == "sql" and p[1]]
    if len(primaries) != 1:
        raise TemplateError(f"{where}: {len(primaries)} primary tables; mark exactly one "
                            f"```sql primary")
    return meta, parts


# --- the tables ---------------------------------------------------------------------------
def _is_cell(dtype) -> bool:
    import polars as pl
    return isinstance(dtype, pl.Struct) and {f.name for f in dtype.fields} >= {"id", "value"}


def _kinds(df, opts: dict, where: str) -> list[str]:
    import polars as pl
    names, over = list(df.columns), opts.get("kinds", {})
    if missing := sorted(set(over) - set(names)):
        raise TemplateError(f"{where}: kinds name {missing}, not columns of {names}")
    out = []
    texts = [n for n, t in zip(names, df.dtypes) if t == pl.String]
    for i, (n, t) in enumerate(zip(names, df.dtypes)):
        if _is_cell(t):
            k = "figure"
        elif n in over:
            k = over[n]
        elif i == 0:
            k = "id"
        elif t == pl.String or t == pl.Boolean:
            k = "status" if n.strip().lower() in STATUS_HEADERS else \
                "note" if (texts and n == texts[-1] and i == len(names) - 1) else "text"
        elif t in (pl.Date, pl.Datetime) or isinstance(t, pl.Datetime):
            k = "date"
        elif t.is_integer():
            k = "count"
        elif t.is_numeric():
            k = "amount"
        else:
            raise TemplateError(f"{where}: column {n!r} is {t}; a number, a date, text or "
                                f"cell('F.id')")
        out.append(k)
    return out


ROLES = {"subtotal": "Subtotal", "total": "Total", "derived": "Subtotal", "less": None}


class Table:
    """One `sql` block or the tie schedule, as rows of (value, kind, unit, fid); each
    row's role (None for a body row, else a key of ROLES), and on a summing row the status
    word it adds (`when`) and the headers it states rather than sums (`stated`)."""

    def __init__(self, heading, headers, kinds, rows, roles=None, sql=None, primary=False,
                 when=None, stated=None):
        self.heading, self.headers, self.kinds, self.rows = heading, headers, kinds, rows
        self.roles = roles or [None] * len(rows)
        self.when = when or [None] * len(rows)
        self.stated = stated or [set() for _ in rows]
        self.sql, self.primary = sql, primary

    def width(self, i: int) -> str:
        k = self.kinds[i]
        if k == "figure":
            units = {r[i][2] for r in self.rows if r[i][0] is not None}
            u = units.pop() if len(units) == 1 else None
            return {"count": "count", "pct": "percent", "rate": "percent"}.get(u, "amount")
        return wbkit.KINDS[k]


def _figure(L, fid: str, where: str) -> dict:
    f = L.get(fid)
    if f is None:
        raise TemplateError(f"{where}: cell('{fid}'): no figure with this id in the run's ledgers")
    return f


def _query_table(run_dir, L, part, meta, where) -> Table:
    import sql as sqlmod
    _, primary, opts, query, heading = part
    try:
        df = sqlmod.db(run_dir).frame(query)
    except Exception as e:
        raise TemplateError(f"{where}: {e}") from None
    def hidden(name):
        if name not in df.columns:
            return [None] * df.height
        return [None if v in (None, "") else str(v) for v in df[name].to_list()]
    roles, when = hidden("_row"), hidden("_when")
    if bad := sorted({r for r in roles if r is not None and r not in ROLES}):
        raise TemplateError(f"{where}: _row {bad}: subtotal, total, derived or less")
    shown = [c for c in df.columns if not c.startswith("_")]
    stated = [set() if s is None else {h.strip() for h in s.split(",")} for s in hidden("_stated")]
    for i, (r, w, s) in enumerate(zip(roles, when, stated)):
        if (w or s) and r not in ("subtotal", "total", "derived"):
            raise TemplateError(f"{where}: row {i + 1}: _when and _stated belong on a subtotal, "
                                f"total or derived row")
        if missing := sorted(s - set(shown)):
            raise TemplateError(f"{where}: row {i + 1}: _stated names {missing}, not columns")
    df = df.drop([c for c in df.columns if c.startswith("_")])
    kinds = _kinds(df, opts, where)
    cur = str(meta.get("currency") or "usd").lower()
    rows = []
    for rec in df.iter_rows():
        row = []
        for v, k in zip(rec, kinds):
            if k == "figure":
                if v is None or v.get("id") is None:
                    row.append((None, k, None, None))
                    continue
                f = _figure(L, v["id"], where)
                row.append((f.get("value"), k, f.get("unit"), v["id"]))
            elif isinstance(v, str):
                row.append((L.sub(v), k, None, None))
            elif isinstance(v, bool):
                row.append(("yes" if v else "no", k, None, None))
            else:
                unit = cur if k in ("amount", "cents") else \
                    {"count": "count", "pct": "pct", "rate": "rate", "fx_rate": "fx_rate",
                     "days": "days"}.get(k)
                row.append((v, k, unit, None))
        rows.append(row)
    if any(when) and "status" not in kinds:
        raise TemplateError(f"{where}: _when adds the rows whose status reads a word; the table "
                            f"has no status column")
    return Table(heading, list(df.columns), kinds, rows, roles, query, primary, when,
                 [{shown.index(h) for h in s} for s in stated])


def _tie_table(ties, heading) -> Table:
    rows = []
    for t in ties:
        u = t["unit"]
        rows.append([(t["id"], "id", None, None), (t["label"], "text", None, None),
                     (t["side_a"]["value"], "figure", u, t["side_a"]["figure_id"]),
                     (t["side_b"]["value"], "figure", u, t["side_b"]["figure_id"]),
                     (t["difference"], "figure", u, t["difference_figure"]),
                     (t["status"], "status", None, None)])
    return Table(heading, ["Tie", "What is agreed", "Side A", "Side B", "Difference", "Status"],
                 ["id", "text", "figure", "figure", "figure", "status"], rows)


# --- the md -------------------------------------------------------------------------------
def _md_value(v, kind, unit):
    if v is None:
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return style.date_short(v)
    if unit is not None and not isinstance(v, str):
        return figures.fmt(v, unit, "cell" if kind == "cents" else "prose")
    return str(v)


def _md_table(t: Table) -> str:
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")  # noqa: E731
    right = [k in ("figure", "amount", "cents", "count", "pct", "rate", "fx_rate", "days")
             for k in t.kinds]
    out = ["| " + " | ".join(esc(h) for h in t.headers) + " |",
           "|" + "|".join("---:" if r else "---" for r in right) + "|"]
    for row, role in zip(t.rows, t.roles):
        cells = ["" if (x := _md_value(*c[:3])) is None else esc(x) for c in row]
        if role in ("subtotal", "total", "derived"):
            cells = [f"**{c}**" if c else c for c in cells]
        out.append("| " + " | ".join(cells) + " |")
    md = "\n".join(out) + "\n"
    if t.sql:
        md += "\nStatement:\n\n```sql\n" + t.sql + "\n```\n"
    return md


def _manifests(run_dir: pathlib.Path, check: str) -> str:
    import sql as sqlmod
    blocks = []
    for p in sorted((run_dir / "checks").glob(f"{check}-*.parquet")):
        m = sqlmod.table_metadata(p) or {}
        if m.get("kind") == "item" and m.get("manifest"):
            blocks.append(items.manifest_md(m["manifest"]))
    return "\n".join(blocks)


# --- the tab ------------------------------------------------------------------------------
def _write_cell(ws, r, c, value, kind, unit, fid, own):
    """One body cell. A figure of another check is written as a copy (`src`), which
    assembly links to that check's cell (WORKBOOK.md § 7)."""
    cell = ws.cell(row=r, column=c)
    ids = {"fid": fid} if fid in own else {"src": fid} if fid else {}
    if value is None:
        cell.style = wbkit.S["Body"]
    elif kind == "id":
        wbkit.ident(cell, value)
    elif kind in ("text", "note"):
        wbkit.text(cell, value)
    elif kind == "status":
        wbkit.status(cell, value)
    elif kind == "period":
        wbkit.text(cell, value if isinstance(value, str) else style.date_short(value))
        cell.alignment = RIGHT
    elif kind == "date":
        wbkit.amount(cell, value, fmt=wbkit.FMT_DATE)
        cell.alignment = RIGHT
    elif kind == "count" or (kind == "figure" and unit == "count"):
        wbkit.count(cell, value, **ids)
        cell.alignment = RIGHT
    else:
        fmt = FIGURE_FORMAT.get(unit, wbkit.FMT_AMOUNT) if kind == "figure" else KIND_FORMAT[kind]
        wbkit.amount(cell, value, fmt=fmt, **ids)
        cell.alignment = RIGHT


def _widths(ws, need: dict) -> None:
    for col, kind in need.items():
        ws.column_dimensions[wbkit.get_column_letter(col)].width = wbkit.WIDTH[kind]


def _place(ws, t: Table, row: int, need: dict, own: set) -> int:
    """Write `t` with its header on `row`; return its last row, the Total included."""
    wbkit.header(ws, row, t.headers, [t.width(i) for i in range(len(t.headers))])
    _widths(ws, need)                          # one column plan for every table on the tab
    status_col = (wbkit.get_column_letter(2 + t.kinds.index("status"))
                  if "status" in t.kinds else None)
    r, body, since, above = row, [], [], None   # body and since: (row, deducted)
    for rec, role, when, stated in zip(t.rows, t.roles, t.when, t.stated):
        r += 1
        if role in (None, "less"):
            for c, (v, k, u, fid) in enumerate(rec, start=2):
                _write_cell(ws, r, c, v, k, u, fid, own)
            body.append((r, role == "less"))
            since.append((r, role == "less"))
            continue
        adds = {"subtotal": since, "total": body,
                "derived": ([(above, False)] if above else []) + since}[role]
        plus, minus = [x for x, d in adds if not d], [x for x, d in adds if d]
        style_ = ROLES[role]
        for i, (v, k, u, fid) in enumerate(rec):
            cell = ws.cell(row=r, column=2 + i)
            number = v is not None and k in NUMERIC and not isinstance(v, str)
            if number and i not in stated:
                fmt = (FIGURE_FORMAT.get(u, wbkit.FMT_COUNT if u == "count" else wbkit.FMT_AMOUNT)
                       if k == "figure" else wbkit.FMT_COUNT if k == "count" else KIND_FORMAT[k])
                wbkit.total(cell, v, rows=plus + minus if when else plus,
                            less=None if when else minus, when=(status_col, when) if when else None,
                            fid=fid if fid in own else None, fmt=fmt, style=style_)
                cell.alignment = RIGHT
            elif number:
                _write_cell(ws, r, 2 + i, v, k, u, fid, own)
                fmt = cell.number_format
                cell.style = wbkit.S[style_]
                cell.number_format = fmt
                wbkit.stated(cell)
            elif v is None:
                cell.style = wbkit.S[style_]
            else:
                wbkit.text(cell, v if isinstance(v, str) else str(v), style_)
        since, above = [], r
    wbkit.grid(ws, row, r, 2, 1 + len(t.headers))
    return r


def _tab(wb, meta, tables, notes, steps, L):
    from openpyxl import Workbook
    for word, kind in (meta.get("statuses") or {}).items():
        wbkit.register_status(str(word), str(kind))
    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = meta["tab"]
    else:
        ws = wb.create_sheet(meta["tab"])
    need: dict[int, str] = {}
    order = ["id", "id_ledger", "count", "percent", "period", "status", "amount", "description", "note"]
    for t in tables:
        for i in range(len(t.headers)):
            k, have = t.width(i), need.get(2 + i)
            if have is None or order.index(k) > order.index(have):
                need[2 + i] = k
    if notes:                                   # the id in B, the note in a wrapped C
        need.setdefault(2, "id")
        if order.index(need.get(3, "id")) < order.index("description"):
            need[3] = "description"
    wbkit.band(ws, L.sub(meta["title"]), L.sub(meta["subtitle"]),
               L.sub(meta["summary"]) if meta.get("summary") else None)
    primary = next(t for t in tables if t.primary)
    own = set(L.entries)
    last = primary_last = _place(ws, primary, 4, need, own)
    shown = primary.heading
    for t in tables:
        if t is primary:
            continue
        r = wbkit.next_block(last)
        if t.heading and t.heading != shown:  # a second table under one heading follows the first
            r = wbkit.section(ws, r, L.sub(t.heading))
        shown = t.heading
        last = _place(ws, t, r, need, own)
    for title, lines in (("Notes", notes), ("To reperform", steps)):
        if not lines:
            continue
        r = wbkit.next_block(last)
        wbkit.section(ws, r, title)
        for j, line in enumerate(lines, start=1):
            if title == "Notes":
                wbkit.ident(ws.cell(row=r + j, column=2), line[0])
                wbkit.text(ws.cell(row=r + j, column=3), L.sub(line[1]), "Note")
            else:
                wbkit.text(ws.cell(row=r + j, column=2), L.sub(line))
        last = r + len(lines)
    _widths(ws, need)
    wbkit.finish(ws, primary_last)
    return wb


# --- rendering ----------------------------------------------------------------------------
def render(run_dir, check: str, wb=None) -> dict:
    """Render `workpapers/<check>.tpl.md`: the check's md, its summary and its staged tab.
    `wb` is a workbook already holding the check's match tabs; the check tab is added to it.
    Returns the paths written."""
    run_dir = pathlib.Path(run_dir).resolve()
    src = run_dir / "workpapers" / f"{check}.tpl.md"
    if not src.is_file():
        raise FileNotFoundError(f"{src}: no template")
    where = src.name
    meta, parts = parse(src.read_text(encoding="utf-8"), where)
    L = figures.Ledger(run_dir, check)
    sub = L.sub
    md, summary, tables, notes, steps = [], [], [], [], []
    for n, part in enumerate(parts):
        kind = part[0]
        if kind == "line":
            md.append(sub(part[1]))
        elif kind == "heading":
            md.append(f"{'#' * part[1]} {sub(part[2])}")
        elif kind == "summary":
            summary.append(sub(part[1]))
        elif kind == "note":
            sub(part[2])
            notes.append((part[1], part[2]))
            md.append(f"- `{part[1]}`: {sub(part[2])}")
        elif kind == "step":
            steps.append(part[1])
            md.append(sub(part[1]))
        elif kind == "ties":
            if L.ties:
                tables.append(_tie_table(L.ties, part[1]))
            md.append(figures.tie_table(L.ties))
        elif kind == "manifests":
            md.append(_manifests(run_dir, check))
        else:
            t = _query_table(run_dir, L, part, meta, f"{where} table {n + 1}")
            tables.append(t)
            md.append(_md_table(t))
    wb = _tab(wb, meta, tables, notes, steps, L)
    out = {"md": run_dir / "checks" / f"{check}.md",
           "tab": run_dir / "out" / ".staging" / f"{check}.xlsx"}
    out["md"].parent.mkdir(parents=True, exist_ok=True)
    out["md"].write_text("\n".join(md).strip() + "\n", encoding="utf-8")
    text_ = "\n".join(summary).strip()
    if text_:
        out["summary"] = run_dir / "checks" / f"{check}-summary.md"
        out["summary"].write_text(text_ + "\n", encoding="utf-8")
    out["tab"].parent.mkdir(parents=True, exist_ok=True)
    wbkit.save(wb, out["tab"])
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("run_dir")
    ap.add_argument("check")
    ap.add_argument("--no-place", action="store_true", help="stage the tab; do not gate it")
    a = ap.parse_args(argv)
    try:
        out = render(a.run_dir, a.check)
    except (TemplateError, KeyError, ValueError, FileNotFoundError) as e:
        print(f"render: {e}", file=sys.stderr)
        return 1
    for k, p in out.items():
        print(f"{k}: {p}")
    if a.no_place:
        return 0
    import step_record
    try:
        print(f"placed: {step_record.place_tab(a.run_dir, a.check)}")
    except step_record.GateRefused as e:
        print(str(e), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
