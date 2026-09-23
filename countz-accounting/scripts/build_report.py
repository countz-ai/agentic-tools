#!/usr/bin/env python3
"""Render the report deck — `out/report.pptx` — from `report.yaml` and the sealed workbook.

The deck is what the reader who decides is handed; the workbook is the trail. The two
must never disagree, so the deck MINTS NOTHING: every table on it is copied from a
workbook tab by this script, and every number inside a sentence is a reference —
`{tab | row label | column header}` or `{tab!B3}` — that this script resolves from the
workbook cell and formats. An author never types a figure into the deck; a reference
that resolves to nothing is a refusal, never a blank.

What the author writes is the story (reference/REPORT.md): `report.yaml` names the
sections and, in each, the pages — a title that is the page's headline (the subject of
a page of figures, the conclusion of a page that argues one), an optional `message`
stating that in one complete sentence under it, and the blocks that carry it: text,
bullets, stats, a table or chart copied from a tab, a check's result strip. This script
lays the cover and the footers (a section is the kicker its pages carry; there is no
contents page and no divider), refuses a sentence where a headline belongs and a
fragment where a sentence does, measures every block so a page that does not fit is
refused rather than shrunk, and continues a long table onto the next page.

Usage:
    build_report.py <run_dir> [--spec out/.staging/report.yaml] [--workbook out/.staging/workbook.xlsx]
                    [--out out/.staging/report.pptx] [--json]

Exit 0 on a rendered deck, 1 when the spec refuses to render — an unresolved reference,
an unknown tab or block, a page whose blocks do not fit — each named, and 2 on a usage
error. A non-zero exit writes no deck. The gate on the result is scripts/check_report.py.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import pathlib
import re
import sys
from copy import deepcopy
from dataclasses import dataclass, field, replace

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from link_workbook import LABEL_COLS, RUN_TABS, normalize  # noqa: E402

try:
    import openpyxl
    import yaml
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError as exc:  # pragma: no cover
    print(f"build_report.py needs openpyxl, pyyaml and python-pptx ({exc}); run it as\n"
          f"  uv run --project ${{CLAUDE_PLUGIN_ROOT}} python3 scripts/build_report.py <run_dir>",
          file=sys.stderr)
    raise SystemExit(2)

# --- the design -----------------------------------------------------------------
# The Countz deck system, as read out of the Paper artboard "Find Fast, Fix Faster — v2"
# and carried by the c4 decks: scripts/design/slide_layout.py in the monorepo holds it in
# px on a 1280x720 artboard (STYLES, RULE, BODY_TOP, FOOTER_TOP). The plugin ships alone,
# so the same values are mirrored here in inches and points (px / 96, px * 0.75). A change
# there is a change here.
BONE = "F5F3EE"                          # the ground
INK, BODY, MUTED = "1C2A2A", "2E3D3C", "566665"
TEAL, BAND, DOT, MARKER = "0F756D", "005C53", "2A9D90", "2A9D90"
RULE, WHITE = "D9D4C8", "FFFFFF"
ACCENT, SLATE = TEAL, MUTED              # the workbook's names for the same two colours
BREAK_T, REVIEW_T, TIED_T = "B42318", "9A5B00", "1E7B3C"   # WORKBOOK_STYLE.md § 1c, text only
FONT = "Inter"
SERIES = [TEAL, MUTED, MARKER, REVIEW_T]

SLIDE_W, SLIDE_H = 13.333, 7.5          # inches, 16:9 (1280 x 720 px)
MARGIN = 0.667                           # 64 px
KICKER_Y, KICKER_H = 0.54, 0.2           # 52 px
TITLE_Y = 0.81                           # 78 px
BODY_Y, BODY_BOTTOM = 1.55, 6.55         # body top, tighter than c4's 174 px; the band starts at 648 px
HEADER_GAP = 0.1                         # between the header (headline, message) and the body
FOOTER_Y, FOOTER_H = 6.75, 0.75          # the band
GAP = 0.13                               # between stacked blocks
COL_GAP = 0.667                          # 64 px, the c4 column gap
BODY_W = SLIDE_W - 2 * MARGIN
RULE_H, RULE_HAIR = 0.03, 0.01           # 3 px accent rule above a block; 1 px hairline
# The chart's box. A chart is DRAWN — rectangles, connectors, text boxes — never placed as
# a chart part: a chart part is redrawn by some viewers and rasterised on import by others
# (Google Slides), and the deck went soft there while Keynote and PowerPoint held. Shapes
# are vector in all three. The values are audited through the shape names (`chartval:`),
# so check_report.py reads a drawn chart exactly as it read a chart part.
CHART_H = 3.1                            # the plot, its axis strip and its legend
CHART_AXIS_H = 0.26                      # the strip under the plot: categories, or ticks on a bar
CHART_LEGEND_H = 0.30                    # the legend strip, only where there is more than one series
CHART_TICKS = 5                          # gridlines the value axis aims for
CHART_BAR_GAP = 0.22                     # of a category slot, left empty between clusters
CHART_DOT = 0.09                         # a line chart's marker
GRID_PT, BASE_PT = 0.75, 1.0             # the gridline and the zero line, in points

# Type scale in points (c4 px * 0.75). Names follow slide_layout.STYLES where one exists.
# `title` is the c4 display title (44 px); `message` the caption style (18 px), the
# cover-subtitle treatment carried onto a page: a muted sentence under an ink headline.
# `stat_value` sits below the display size: a stat tile carries a whole dollar figure in
# a quarter of the body width, and at the display size it wraps mid-figure. `table_xs` and
# `heading_xs` are a schedule the author declares `dense: true` — a walk or a roster shown
# at full population (REPORT.md § 2) — never a size the builder picks to make a page fit.
PT = {"kicker": 9, "title": 33, "message": 13.5, "lead": 16.5, "body": 12, "small": 11.25,
      "heading": 9, "note": 9.75, "table": 10.5, "table_sm": 9.75, "table_xs": 8.25,
      "heading_xs": 7.5, "stat_value": 24,
      "stat_label": 9, "footer": 8.25, "tagline": 15, "cover_title": 42, "cover_sub": 18,
      "result_word": 18, "caption": 13.5}
TRACK = {"kicker": 0.12, "heading": 0.08, "stat_label": 0.08}    # em, on caps labels
MESSAGE_GAP = 0.08                       # between the headline's line box and the message
# A title is a headline — `EBITDA`, `Acquired-intangible amortisation rejected` — never
# the sentence; two lines at the display size is the most a headline runs.
TITLE_MAX, TITLE_LINES = 80, 2
MESSAGE_MAX = 160                        # characters — the workbook's result-sentence cap (WORKBOOK.md § 3)
# The cover is four strings — the company as the kicker, the title, the subtitle, the
# prepared line — and each fact appears once: the title is the work, the subtitle the
# entity detail and the period, the prepared line the company and the date. The Exec
# Summary's subtitle carries the whole basis line — entity, period, basis, unit, tolerance
# — so the default is its leading company mention dropped and the rest cut to whole `·`
# segments inside this budget. An author's own `subtitle` is refused, never cut.
COVER_SUB_MAX = 72                       # characters
COVER_TITLE_MAX, COVER_TITLE_LINES = 60, 2
# A period, wherever it is written: `FY2023`, `2023`, `2023-09-30`, `30 September 2023`.
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December")
PERIOD_TOKEN = re.compile(
    rf"\b(?:FY|CY)\s?(?:19|20)?\d\d\b|\b(?:19|20)\d\d\b|\b\d{{4}}-\d{{2}}-\d{{2}}\b"
    rf"|\b(?:{'|'.join(m[:3] for m in MONTHS)})[a-z]*\.?\s+\d{{1,4}}\b", re.I)
LEGAL_FORM = re.compile(r"[\s,]*\b(?:inc|llc|l\.l\.c|ltd|limited|corp|corporation|co|company|plc|"
                        r"gmbh|s\.a|sa|sas|bv|nv|ag|pty|llp|lp|holdings?|group)\b\.?", re.I)
CONTINUED = " (continued)"

# Result words and their colour (WORKBOOK_STYLE.md § 5). Colour repeats the word; never alone.
STATUS_COLOR = {"pass": TIED_T, "supported": TIED_T, "tied": TIED_T, "clean": TIED_T,
                "warn": REVIEW_T, "candidate": REVIEW_T, "review": REVIEW_T, "open": REVIEW_T,
                "fail": BREAK_T, "unexplained": BREAK_T, "break": BREAK_T, "exception": BREAK_T,
                "withheld": MUTED, "not_run": MUTED, "not run": MUTED, "blocked": MUTED,
                "degraded": MUTED}
REF = re.compile(r"\{([^{}]+)\}")
EXEC = "Exec Summary"
TABLE_STYLE_NONE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # "No Style, No Grid"
ID_HEADERS = {"id"}


class SpecError(Exception):
    """A defect in report.yaml or in what it asks of the workbook. Collected, then fatal."""


# --- measurement --------------------------------------------------------------
# Arial's average advance is ~0.5 em; bold ~0.55. Widths are estimates for layout, never
# for truth: a block that measures over the page is refused, so an underestimate surfaces
# as an overflow in PowerPoint and the constants below lean generous.
def text_w(s: str, pt: float, bold: bool = False) -> float:
    # Inter runs wider than Arial; measured on a live deck, 0.52 em under-wrapped a
    # 15-row text table into the footer band
    return len(s) * pt * (0.6 if bold else 0.56) / 72


def lines_for(s: str, width: float, pt: float, bold: bool = False) -> int:
    n = 0
    for para in (s or "").split("\n"):
        words = para.split(" ")
        cur, ln = 0.0, 1
        for w in words:
            ww = text_w(w + " ", pt, bold)
            if cur + ww > width and cur > 0:
                ln += 1
                cur = ww
            else:
                cur += ww
        n += ln
    return max(n, 1)


def text_h(s: str, width: float, pt: float, bold: bool = False, spacing: float = 1.2) -> float:
    return lines_for(s, width, pt, bold) * pt * spacing / 72 + 0.06


def nice_axis(lo: float, hi: float, target: int = CHART_TICKS) -> tuple[float, float, list[float]]:
    """The value axis a reader can read: round ticks spanning the data, zero on the scale.
    The step is 1, 2, 2.5 or 5 times a power of ten, never the raw range over the tick
    count, so the labels are numbers a reader holds in their head."""
    lo, hi = min(0.0, lo), max(0.0, hi)
    if hi - lo <= 0:
        hi = lo + 1.0
    raw = (hi - lo) / max(target, 1)
    mag = 10.0 ** math.floor(math.log10(raw)) if raw > 0 else 1.0
    step = next((m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw - 1e-12), 10 * mag)
    lo, hi = math.floor(lo / step) * step, math.ceil(hi / step) * step
    ticks, v = [], lo
    while v <= hi + step * 1e-6:
        ticks.append(0.0 if abs(v) < step * 1e-6 else v)
        v += step
    return lo, hi, ticks


def axis_text(v: float, step: float) -> str:
    """A tick, in the deck's number conventions: thousands separated, negatives in
    parentheses, and only as many decimals as the step actually distinguishes."""
    if step >= 1e5:                          # a money axis is read at its scale (§ 4)
        a, unit = abs(v) / 1e6, "m"
    elif step >= 100:
        a, unit = abs(v) / 1e3, "k"
    else:
        a, unit = abs(v), ""
    if unit:
        s = f"{a:,.1f}".rstrip("0").rstrip(".") + unit if a else "0"
        return f"({s})" if v < 0 else s
    d = 0 if step >= 1 else min(4, int(math.ceil(-math.log10(step))))
    return (f"({abs(v):,.{d}f})" if v < 0 else f"{v:,.{d}f}")


# --- prose ----------------------------------------------------------------------
# A sentence on a slide is complete: it ends with a full stop, a question mark or an
# exclamation (a closing bracket or quote may follow), opens with a capital, and is never
# clipped. A check token (`q6`) or a file name (`gl.csv`) may open it lowercase.
SENTENCE_END = re.compile(r"[.?!][)\]\"'”’]*$")
LOWER_OPENER = re.compile(r"^[a-z][\w-]*[\d.]")


def incomplete(s: str) -> str | None:
    """Why `s` is not a complete sentence, or None when it is."""
    s = (s or "").strip()
    if not s:
        return "is empty"
    if "…" in s or "..." in s:
        return "is clipped (…) — write the whole sentence"
    if not SENTENCE_END.search(s):
        return "does not end with a full stop — a complete sentence, not a fragment"
    if s[0].islower() and not LOWER_OPENER.match(s):
        return "opens lowercase — a complete sentence opens with a capital"
    return None


def sentence_case(s: str) -> str:
    """Capitalise the first letter unless the first word is a token or a file name."""
    s = s.strip()
    if s and s[0].islower() and not LOWER_OPENER.match(s):
        return s[0].upper() + s[1:]
    return s


def company_phrase(company: str) -> str:
    """The company without its legal form — `Demo DGII Corp` reads `demo dgii` — for
    matching a mention of it in another cover string."""
    return " ".join(LEGAL_FORM.sub(" ", company or "").split()).strip(" ,.-").lower()


def drops_company(s: str, company: str) -> str:
    """`s` with a leading mention of the company dropped: the Exec Summary's basis line
    opens with the entity (`Demo DGII Corp, seventeen operating legal entities`), and on
    the cover the company is the kicker."""
    name = company_phrase(company)
    if not name:
        return s
    m = re.match(re.escape(name), s.strip(), re.I)
    if not m:
        return s
    rest = s.strip()[m.end():]
    form = LEGAL_FORM.match(rest)
    if form:
        rest = rest[form.end():]
    rest = rest.lstrip(" ,·-–—")
    return sentence_case(rest) if rest else s


def cover_subtitle(s: str, company: str = "") -> str:
    """The Exec Summary's basis line, its leading company mention dropped and the rest cut
    to the whole `·` segments that fit the cover."""
    parts = [p.strip() for p in drops_company(s, company).split("·")]
    out: list[str] = []
    for part in parts:
        if not out or len(" · ".join(out + [part])) <= COVER_SUB_MAX:
            out.append(part)
            continue
        # The period segment carries a parenthetical and an as-of date and is usually the
        # one that does not fit whole; its head (`FY2023`) still states the period, so take
        # that where it fits and stop.
        head = re.split(r"[(,;]", part, 1)[0].strip()
        if head and len(" · ".join(out + [head])) <= COVER_SUB_MAX:
            out.append(head)
        break
    return " · ".join(out)


def not_cover_title(title: str, company: str) -> str | None:
    """Why `title` is not a cover title, or None. It names the work — `Quality of earnings
    review`, `Revenue leak: billed to collected` — and carries neither the company (the
    kicker and the prepared line do) nor the period (the subtitle does)."""
    if len(title) > COVER_TITLE_MAX:
        return f"{len(title)} characters; the title names the work, at most {COVER_TITLE_MAX}"
    if lines_for(title, BODY_W, PT["cover_title"], True) > COVER_TITLE_LINES:
        return f"runs past {COVER_TITLE_LINES} lines at {PT['cover_title']}pt — shorter"
    name = company_phrase(company)
    if name and name in company_phrase(title):
        return (f"names the company (`{company}`) — the cover carries it as the kicker and in "
                f"`Prepared for`, so the title names the work alone")
    m = PERIOD_TOKEN.search(title)
    if m:
        return f"carries the period (`{m.group(0)}`) — the period is the subtitle's, the date the prepared line's"
    return None


def not_cover_subtitle(sub: str, title: str, company: str) -> str | None:
    """Why `sub` is not a cover subtitle, or None: it carries the entity detail and the
    period on one line, and repeats neither the company nor the title."""
    if len(sub) > COVER_SUB_MAX:
        return f"{len(sub)} characters; the entity and the period, at most {COVER_SUB_MAX}"
    if lines_for(sub, BODY_W, PT["cover_sub"]) > 1:
        return f"runs past one line at {PT['cover_sub']}pt — the entity and the period, shorter"
    name = company_phrase(company)
    if name and name in company_phrase(sub):
        return f"names the company (`{company}`) — the cover's kicker and prepared line carry it"
    words = [w for w in re.split(r"[^\w]+", title.lower()) if len(w) > 3]
    shared = [w for w in words if re.search(rf"\b{re.escape(w)}\b", sub.lower())]
    if len(shared) >= 2:
        return (f"repeats the title ({', '.join(shared[:3])}) — the title names the work, "
                f"the subtitle the entity and the period")
    return None


def not_headline(title: str) -> str | None:
    """Why `title` is not a headline, or None: at most TITLE_MAX characters and
    TITLE_LINES lines at the display size, and no full stop — the sentence is `message`."""
    if len(title) > TITLE_MAX:
        return f"{len(title)} characters; a headline — the subject or the conclusion — at most {TITLE_MAX}"
    if title.rstrip().endswith("."):
        return "ends with a full stop — a headline, not a sentence; the sentence goes in `message`"
    if lines_for(title, BODY_W, PT["title"], True) > TITLE_LINES:
        return f"runs past {TITLE_LINES} lines at {PT['title']}pt — a headline, shorter"
    return None


# --- the workbook ---------------------------------------------------------------
@dataclass
class Cell:
    value: object
    fmt: str
    bold: bool
    fill: str
    coord: str


@dataclass
class Table:
    """A copied schedule: `headers` (column, text), `rows` of Cells aligned to headers,
    `source` the tab, `title` the block's heading, `more` rows left on the tab."""
    source: str
    title: str | None
    headers: list[str]
    rows: list[list[Cell]]
    kinds: list[str] = field(default_factory=list)   # body | subtotal | total per row
    more: int = 0
    numeric: list[bool] = field(default_factory=list)
    dense: bool = False                              # `dense: true` on the block (REPORT.md § 2)


@dataclass
class Lines:
    source: str
    title: str | None
    lines: list[str]


class Book:
    """The sealed workbook, read once with cached values (check_workbook.py has already
    refused a formula with no cached result)."""

    def __init__(self, path: pathlib.Path):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.tabs = self.wb.sheetnames
        sub = str(self.cell(EXEC, "B2").value or "") if EXEC in self.tabs else ""
        self.money = "$" if re.search(r"\bUSD\b|\$|dollar", sub, re.I) else ""
        self._grid: dict[str, tuple] = {}

    # tab resolution: the exact name, or the roster token that opens it (`q6`).
    def tab(self, name: str) -> str:
        name = str(name).strip()
        if name in self.tabs:
            return name
        tok = normalize(name)
        hits = [t for t in self.tabs if normalize(t.split(" ", 1)[0]) == tok or normalize(t) == tok]
        if not hits:                       # `r4 concentration` opens `r4 concentration Customer…`
            hits = [t for t in self.tabs if normalize(t).startswith(tok + "_")]
        if len(hits) == 1:
            return hits[0]
        if not hits:
            raise SpecError(f"no workbook tab named `{name}` (tabs: {', '.join(self.tabs)})")
        raise SpecError(f"`{name}` names {len(hits)} tabs ({', '.join(hits)}); use the full name")

    def check_tabs(self) -> list[str]:
        return [t for t in self.tabs if normalize(t) not in RUN_TABS]

    def cell(self, tab: str, coord: str):
        return self.wb[tab][coord]

    def as_cell(self, c) -> Cell:
        fill = ""
        try:
            if c.fill is not None and c.fill.fill_type == "solid":
                fill = (c.fill.fgColor.rgb or "")[-6:].upper()
        except Exception:
            fill = ""
        return Cell(c.value, c.number_format or "", bool(c.font and c.font.b), fill, c.coordinate)

    def grid(self, tab: str):
        """(headers_at[row][col], row_texts[row] -> {texts in the first LABEL_COLS columns},
        row_cells[row] -> [(col, cell)] non-empty) — the same matching the linker uses for
        the Exec Summary's copied amounts (link_workbook.py rule 6)."""
        if tab in self._grid:
            return self._grid[tab]
        ws = self.wb[tab]
        headers_at, row_texts, row_cells = {}, {}, {}
        hdr: dict[int, str] = {}
        for row in ws.iter_rows():
            if not row:
                continue
            r = row[0].row
            # a column's header is the label the last HEADER ROW above gave it — a row of
            # short bold strings and no number (the kit's Header / HeaderPlain styles) —
            # so a text cell in a data row is looked up under its column's header and
            # never under the text of the row before it
            headers_at[r] = dict(hdr)
            cells = []
            for c in row:
                v = c.value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                cells.append((c.column, c))
                if isinstance(v, str) and c.column <= LABEL_COLS:
                    row_texts.setdefault(r, set()).add(v.strip())
            is_header = (len(cells) >= 2 and all(isinstance(c.value, str) and len(c.value.strip()) <= 48
                                                  and bool(c.font and c.font.b) for _, c in cells))
            if is_header:
                hdr = {col: c.value.strip() for col, c in cells}
            elif len(cells) == 1 and isinstance(cells[0][1].value, str) and cells[0][1].font and cells[0][1].font.b \
                    and len(cells[0][1].value.strip()) <= 120:
                hdr = {}                      # a section heading: the next header row starts afresh
            if cells:
                row_cells[r] = cells
        self._grid[tab] = (headers_at, row_texts, row_cells)
        return self._grid[tab]

    def lookup(self, tab: str, label: str, header: str) -> Cell:
        """The cell at (row whose leading text is `label`, column whose header is
        `header`) — both copied verbatim from the tab, as the linker demands."""
        headers_at, row_texts, row_cells = self.grid(tab)
        label, header = label.strip(), header.strip()
        rows = [r for r in sorted(row_texts) if label in row_texts[r]]
        if not rows:
            fold = normalize(label)
            rows = [r for r in sorted(row_texts) if any(normalize(t) == fold for t in row_texts[r])]
        if not rows:
            raise SpecError(f"`{tab}` has no row whose leading label is `{label}`")
        for r in rows:
            cols = [c for c, h in headers_at.get(r, {}).items() if h == header]
            if not cols:
                fold = normalize(header)
                cols = [c for c, h in headers_at.get(r, {}).items() if normalize(h) == fold]
            for col in cols:
                c = self.wb[tab].cell(r, col)
                if c.value is not None:
                    return self.as_cell(c)
        raise SpecError(f"`{tab}`: row `{label}` has no value under a column headed `{header}`")

    def band(self, tab: str) -> dict:
        """The frozen band of a check tab (WORKBOOK.md § 3): title, subtitle, and the
        summary sentence on B3."""
        ws = self.wb[tab]
        return {"title": str(ws["B1"].value or ""), "subtitle": str(ws["B2"].value or ""),
                "summary": str(ws["B3"].value or "")}

    def blocks(self, tab: str, start: int = 4) -> list[Table | Lines]:
        """Every block on a tab, in order (WORKBOOK.md § 4, § 6): the primary table
        (header on row 4, no title), then each titled block — a table (a header row of
        short labels over body rows) or lines (one statement per row) — each closed by a
        blank row."""
        ws = self.wb[tab]
        _, _, row_cells = self.grid(tab)
        max_row = ws.max_row
        out: list[Table | Lines] = []
        r = start
        first = tab != EXEC

        def is_short_labels(cells) -> bool:
            return len(cells) >= 2 and all(isinstance(c.value, str) and len(c.value.strip()) <= 48
                                            for _, c in cells)

        def is_statement(cells) -> bool:
            """One long sentence in a row — a Notes or How-to-read line, never a heading
            and never a table row."""
            if len(cells) != 1 or not isinstance(cells[0][1].value, str):
                return False
            s = cells[0][1].value.strip()
            c = cells[0][1]
            if c.font and c.font.b and len(s) <= 120:
                return False                  # a Section heading, however long
            return len(s) > 60 or s.endswith((".", "!", "?"))

        def read_table(title, header_row):
            hcells = row_cells[header_row]
            cols = [col for col, _ in hcells]
            headers = [str(c.value).strip() for _, c in hcells]
            rows, kinds = [], []
            rr = header_row + 1
            while rr <= max_row and rr in row_cells and not is_statement(row_cells[rr]):
                cells = {col: c for col, c in row_cells[rr]}
                line = [self.as_cell(cells[col]) if col in cells
                        else Cell(None, "", False, "", f"{openpyxl.utils.get_column_letter(col)}{rr}")
                        for col in cols]
                lead = next((c for c in line if c.value is not None), None)
                label = str(lead.value) if lead else ""
                kind = "body"
                if lead and lead.bold:
                    kind = "total" if re.match(r"^(=\s*)?total\b", label, re.I) or \
                        (cells and any(getattr(c.border.bottom, "style", None) == "double"
                                       for c in cells.values())) else "subtotal"
                rows.append(line)
                kinds.append(kind)
                rr += 1
            numeric = [any(isinstance(row[i].value, (int, float)) and not isinstance(row[i].value, bool)
                           for row in rows) for i in range(len(cols))]
            return Table(tab, title, headers, rows, kinds, 0, numeric), rr

        while r <= max_row:
            if r not in row_cells:
                r += 1
                continue
            cells = row_cells[r]
            if first and r == start:
                first = False
                if len(cells) >= 2:
                    t, r = read_table(None, r)
                    out.append(t)
                    continue
            if len(cells) == 1 and isinstance(cells[0][1].value, str) and not is_statement(cells):
                title = str(cells[0][1].value).strip()
                nr = r + 1
                while nr <= max_row and nr not in row_cells and nr - r <= 2:
                    nr += 1
                if nr <= max_row and nr in row_cells and is_short_labels(row_cells[nr]) \
                        and (nr + 1) in row_cells:
                    t, r = read_table(title, nr)
                    out.append(t)
                    continue
                lines = []
                rr = nr
                while rr <= max_row and rr in row_cells:
                    parts = [str(c.value).strip() for _, c in row_cells[rr] if c.value is not None]
                    lines.append("  ".join(parts))
                    rr += 1
                if lines:
                    out.append(Lines(tab, title, lines))
                    r = rr
                else:
                    out.append(Lines(tab, None, [title]))
                    r = nr
                continue
            if len(cells) >= 2 and is_short_labels(cells) and (r + 1) in row_cells:
                t, r = read_table(None, r)
                out.append(t)
                continue
            # a row of statements with no title (e.g. the source-id lines under Evidence)
            lines = []
            rr = r
            while rr <= max_row and rr in row_cells:
                lines.append("  ".join(str(c.value).strip() for _, c in row_cells[rr]
                                       if c.value is not None))
                rr += 1
            out.append(Lines(tab, None, lines))
            r = rr
        return out

    def block(self, tab: str, title: str) -> Table | Lines:
        fold = normalize(title)
        for b in self.blocks(tab):
            if b.title and (b.title.strip() == title.strip() or normalize(b.title) == fold
                            or normalize(b.title).startswith(fold)):
                return b
        for b in self.blocks(tab):            # an untitled statement block, by its first line
            if isinstance(b, Lines) and not b.title and b.lines and normalize(b.lines[0]).startswith(fold):
                return b
        raise SpecError(f"`{tab}` has no block titled `{title}`")

    def primary(self, tab: str) -> Table:
        bl = self.blocks(tab)
        if bl and isinstance(bl[0], Table) and bl[0].title is None:
            return bl[0]
        raise SpecError(f"`{tab}` has no primary table on row 4")


# --- formatting (REPORT.md § 4; the workbook's own conventions are DOCTRINE.md) ------
def compact_money(v: float, money: str = "$") -> str:
    """A dollar figure as the deck states it (REPORT.md § 4): scaled and rounded —
    `$50.5m`, `$81k` — never to the dollar. The workbook keeps the exact figure."""
    a = abs(v)
    if a >= 999_500_000:
        s = f"{money}{a / 1e9:,.1f}bn"
    elif a >= 999_500:
        s = f"{money}{a / 1e6:,.1f}m"
    elif a >= 999.5:
        s = f"{money}{a / 1e3:,.0f}k"
    else:
        s = f"{money}{a:,.0f}"
    return f"({s})" if v < 0 else s


def fmt_value(v, fmt: str = "", prose: bool = False, money: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%-d %b %Y")
    if isinstance(v, str):
        return v.strip()
    if not isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return ""
    fmt = fmt or ""
    if "%" in fmt:
        pct = v * 100
        s = f"{abs(pct):.0f}%" if float(abs(pct)).is_integer() else f"{abs(pct):.1f}%"
        return f"({s})" if pct < 0 else s
    if "@" in fmt:
        return str(v)
    m = re.fullmatch(r"0(?:\.(0+))?", fmt)
    if m:                                    # days, ratios, counts: `0.0` / `0`
        d = len(m.group(1)) if m.group(1) else 0
        if d > 1 and abs(v) >= 1:            # one decimal on the deck (REPORT.md § 4)
            d = 1
        return f"{v:,.{d}f}"
    if fmt in ("#,##0",):
        return f"{v:,.0f}"
    if v == 0:                               # an en dash is the table's zero (§ 4 Style)
        return (money + "0") if prose else "–"
    if prose and money:                      # a sentence or a figure tile: scaled (§ 4)
        return compact_money(v, money)
    # any other figure keeps the decimals its cell shows, capped at the one the deck shows
    d = 2 if re.search(r"0\.00(?![0-9])", fmt) else 0
    if d > 1 and abs(v) >= 1:
        d = 1
    if prose and round(v, d) == 0:
        return "0"
    s = f"{abs(v):,.{d}f}"
    return f"({s})" if v < 0 else s


# --- the spec --------------------------------------------------------------------
class Resolver:
    """Every `{…}` reference in the author's text, resolved from the workbook and
    formatted for prose. Unresolved references are collected, not blanked."""

    def __init__(self, book: Book):
        self.book = book
        self.errors: list[str] = []
        self.count = 0
        self.sources: set[str] = set()      # every tab the deck draws on
        self.touched: set[str] = set()      # the tabs the page being built draws on

    def use(self, tab: str) -> None:
        self.sources.add(tab)
        self.touched.add(tab)

    def resolve(self, text: str, where: str) -> str:
        if not isinstance(text, str):
            return "" if text is None else str(text)

        def sub(m: re.Match) -> str:
            ref = m.group(1).strip()
            money = ""
            parts = [x.strip() for x in ref.split("|")]
            if parts and parts[-1] in ("$", "plain"):
                money = self.book.money if parts[-1] == "$" else ""
                ref = " | ".join(parts[:-1])
            try:
                cell, tab = self._cell(ref)
            except SpecError as exc:
                self.errors.append(f"{where}: {{{ref}}} — {exc}")
                return m.group(0)
            self.count += 1
            self.use(tab)
            return fmt_value(cell.value, cell.fmt, prose=True, money=money)
        return REF.sub(sub, text)

    def _cell(self, ref: str) -> tuple[Cell, str]:
        parts = [p.strip() for p in ref.split("|")]
        if len(parts) == 1 and "!" in parts[0]:
            tab, coord = parts[0].rsplit("!", 1)
            tab = self.book.tab(tab.strip().strip("'"))
            coord = coord.strip().replace("$", "").upper()
            if not re.fullmatch(r"[A-Z]{1,3}\d+", coord):
                raise SpecError(f"`{coord}` is not a cell reference")
            c = self.book.cell(tab, coord)
            if c.value is None:
                raise SpecError(f"`{tab}`!{coord} is empty")
            return self.book.as_cell(c), tab
        if len(parts) == 3:
            tab = self.book.tab(parts[0])
            return self.book.lookup(tab, parts[1], parts[2]), tab
        raise SpecError("a reference is `{tab | row label | column header}` or `{tab!B3}`, "
                        "with an optional trailing `| $` for a dollar figure")


# --- page model ---------------------------------------------------------------------
@dataclass
class Page:
    section: str
    kicker: str
    title: str              # the headline
    blocks: list            # block dicts, resolved
    sources: list[str] = field(default_factory=list)
    kind: str = "page"      # cover | page
    number: int = 0
    tagline: str = ""       # the page's takeaway in the footer band, optional
    message: str = ""       # the page's message as one sentence under the headline, optional


def as_blocks(raw, res: Resolver, where: str, book: Book) -> list[dict]:
    """Normalise the author's block list: resolve references in text, copy tables and
    chart data from the workbook, validate shapes."""
    out: list[dict] = []
    if raw is None:
        return out
    if not isinstance(raw, list):
        raise SpecError(f"{where}: blocks must be a list")
    for i, b in enumerate(raw):
        at = f"{where}[{i}]"
        if isinstance(b, str):
            out.append({"t": "text", "text": prose(res.resolve(b, at), at)})
            continue
        if not isinstance(b, dict) or len(b) == 0:
            raise SpecError(f"{at}: a block is a mapping with one key")
        keys = [k for k in b if k not in ("lead", "widths", "size")]
        if len(keys) != 1:
            raise SpecError(f"{at}: a block has one kind key, got {', '.join(b)}")
        kind = keys[0]
        v = b[kind]
        if kind == "text":
            out.append({"t": "text", "text": prose(res.resolve(v, at), at), "lead": bool(b.get("lead"))})
        elif kind == "heading":
            out.append({"t": "heading", "text": res.resolve(v, at)})
        elif kind == "note":
            out.append({"t": "note", "text": res.resolve(v, at)})
        elif kind == "bullets":
            if not isinstance(v, list):
                raise SpecError(f"{at}: bullets is a list of strings")
            out.append({"t": "bullets", "items": [res.resolve(s, f"{at}.bullets[{j}]") for j, s in enumerate(v)]})
        elif kind == "stats":
            if not isinstance(v, list) or not v:
                raise SpecError(f"{at}: stats is a list of {{label, value}}")
            if len(v) > 4:
                raise SpecError(f"{at}: at most four stats on a row")
            tiles = []
            for j, s in enumerate(v):
                if not isinstance(s, dict) or "value" not in s or "label" not in s:
                    raise SpecError(f"{at}.stats[{j}]: needs label and value")
                known(s, STAT_KEYS, f"{at}.stats[{j}]", "stat")
                tiles.append({"label": res.resolve(s["label"], at), "value": res.resolve(str(s["value"]), at),
                              "note": res.resolve(s.get("note", ""), at)})
            out.append({"t": "stats", "tiles": tiles})
        elif kind == "kv":
            if not isinstance(v, list):
                raise SpecError(f"{at}: kv is a list of {{label, value}}")
            out.append({"t": "kv", "items": [(res.resolve(str(p.get("label", "")), at),
                                              res.resolve(str(p.get("value", "")), at)) for p in v]})
        elif kind == "table":
            out.append(table_block(v, at, res, book))
        elif kind == "lines":
            # a tab's statement block — Notes, To reperform, a How-to-read line list — as bullets
            if not isinstance(v, dict) or "from" not in v or "block" not in v:
                raise SpecError(f"{at}: lines needs `from: <tab>` and `block: <title>`")
            tab = book.tab(v["from"])
            blk = book.block(tab, v["block"])
            if not isinstance(blk, Lines):
                raise SpecError(f"{at}: `{v['block']}` on `{tab}` is a table, not lines — use table:")
            res.use(tab)
            if v.get("title", blk.title):
                out.append({"t": "heading", "text": res.resolve(str(v.get("title", blk.title)), at)})
            out.append({"t": "bullets", "items": list(blk.lines), "small": True})
        elif kind == "chart":
            out.append(chart_block(v, at, res, book))
        elif kind == "columns":
            items = v.get("items") if isinstance(v, dict) else v
            widths = v.get("widths") if isinstance(v, dict) else b.get("widths")
            if not isinstance(items, list) or not (2 <= len(items) <= 3):
                raise SpecError(f"{at}: columns holds two or three lists of blocks")
            cols = [as_blocks(col, res, f"{at}.columns[{j}]", book) for j, col in enumerate(items)]
            if widths and (len(widths) != len(cols) or abs(sum(widths) - 1) > 0.01):
                raise SpecError(f"{at}: widths are one share per column and sum to 1")
            out.append({"t": "columns", "cols": cols, "widths": widths or [1 / len(cols)] * len(cols)})
        else:
            raise SpecError(f"{at}: unknown block kind `{kind}` (text, heading, note, bullets, "
                            f"stats, kv, table, lines, result, chart, columns)")
    return out


TABLE_KEYS = {"from", "block", "rows", "columns", "max_rows", "title", "ids", "fit", "scale",
              "where", "through", "dense"}
# A schedule of dollars is shown at a declared scale, each dollar column headed with it:
# `64,143` under `$'000` (REPORT.md § 4).
SCALES = {"thousands": (1e3, "$'000"), "millions": (1e6, "$m")}
# What a scaled dollar column is formatted as once it is stated at a scale: whole units,
# negatives in parentheses, zero an en dash (REPORT.md § 4). Applied so the DISPLAYED
# precision is the column's own, whatever precision the tab held - and so `shown_unit`,
# which the footing check reads, is the unit the reader actually sees.
SCALED_FMT = '#,##0;(#,##0);"-"'.replace("-", "\u2013")
MONEY_FMT = re.compile(r"#,##0(?:\.00)?(?![.0-9])")


def is_money(fmt: str | None) -> bool:
    """A dollar column, whole or to the cent. A percent, ratio or day count is not one."""
    f = fmt or ""
    return "%" not in f and bool(MONEY_FMT.search(f))


# A count shares the whole-currency format on a tab (WORKBOOK_STYLE.md: `#,##0` serves
# both), so the header decides: a column it names as a count of things, holding whole
# numbers only, is never a dollar column and is never scaled.
COUNT_HEADER = re.compile(r"^(number|count|no\.) of\b|\b(count|customers|logos|contracts|"
                          r"invoices|lines|months|days)$", re.I)


def is_count(header: str, cells) -> bool:
    return bool(COUNT_HEADER.search(header.strip())) and all(
        float(c.value).is_integer() for c in cells)

CHART_KEYS = {"type", "from", "rows", "columns", "block", "title"}
STAT_KEYS = {"label", "value", "note"}


def prose(text: str, at: str) -> str:
    """A text block, references resolved, refused unless it reads as complete sentences."""
    why = incomplete(text)
    if why:
        raise SpecError(f"{at}: text {why}")
    return text


def known(v: dict, keys: set, at: str, what: str) -> None:
    extra = [k for k in v if k not in keys]
    if extra:
        raise SpecError(f"{at}: {what} has no key {', '.join(map(repr, extra))} — a comma inside a "
                        f"flow mapping splits a value; quote it (keys: {', '.join(sorted(keys))})")


DERIVED_ROW = re.compile(r"^\s*=")        # a schedule's own marker for a derived line


def row_labels(row: list[Cell]) -> list[str]:
    """The texts a row carries in the label columns — its id and its line label."""
    return [c.value.strip() for c in row[:LABEL_COLS + 1]
            if isinstance(c.value, str) and c.value.strip()]


ID_LIKE = re.compile(r"^[A-Z]{1,2}\.[A-Za-z0-9_.*-]+$")   # a ledger id (EVIDENCE.md § 0)


def row_label(row: list[Cell]) -> str:
    """A row's line label: the first text of its label columns that is not a ledger id."""
    texts = row_labels(row)
    return next((t for t in texts if not ID_LIKE.match(t)), texts[0] if texts else "")


def is_derived(row: list[Cell]) -> bool:
    """Whether a row is a schedule's own derived line — its label opens `= `."""
    return any(DERIVED_ROW.match(t) for t in row_labels(row))


def shown_unit(fmt: str) -> float:
    """The smallest difference the reader can see in a column formatted `fmt`."""
    m = re.search(r"0\.(0+)", fmt or "")
    return 10.0 ** -len(m.group(1)) if m else 1.0


def not_footing(rows: list[list[Cell]], picked: list[int], orig: list[list[Cell]],
                numeric: list[bool], kinds: list[str] | None = None) -> str | None:
    """Why a SELECTION of a schedule's rows does not foot as shown, or None (REPORT.md § 2).

    A derived row (`= …`) shown above contributing rows must equal the derived row above it
    plus the rows shown between them; signs are in the stored values (`less …` is negative
    on the tab). Two exemptions: the tab's own run of rows copied whole, and two derived
    rows with nothing shown between them. A subtotal shown beside the body rows it sums
    (a walk at item grain, `where:`) is not added a second time: where body rows and
    subtotal rows both stand between two derived rows, the body rows are the terms."""
    prev: int | None = None                  # position in `rows` of the derived row above
    acc: list[int] = []                      # positions shown since it
    label = lambda pos: row_label(orig[picked[pos]])                                  # noqa: E731
    for pos, row in enumerate(rows):
        if not is_derived(orig[picked[pos]]):
            acc.append(pos)
            continue
        start = picked[prev] if prev is not None else -1
        run = list(range(start + 1, picked[pos] + 1))
        if not acc or [picked[q] for q in acc] + [picked[pos]] == run:
            prev, acc = pos, []              # the two exemptions, above
            continue
        body = [q for q in acc if kinds is None or kinds[q] == "body"]
        terms_at = body if body and len(body) < len(acc) else acc
        for ci, is_num in enumerate(numeric):
            got = rows[pos][ci].value
            if not is_num or not isinstance(got, (int, float)) or isinstance(got, bool):
                continue
            terms = [rows[q][ci].value for q in terms_at
                     if isinstance(rows[q][ci].value, (int, float)) and not isinstance(rows[q][ci].value, bool)]
            base = rows[prev][ci].value if prev is not None and isinstance(rows[prev][ci].value, (int, float)) else 0
            expected = base + sum(terms)
            if abs(expected - got) > 0.55 * (len(terms) + 2) * shown_unit(rows[pos][ci].fmt):
                missing = [row_label(orig[j]) for j in run if j not in picked]
                # the rows nearest the derived line are the ones it is missing
                named = "; ".join(m[:44] for m in missing[-3:])
                return (f"the rows shown do not foot to `{label(pos)}`: they come to "
                        f"{expected:,.2f} against {got:,.2f} — the schedule derives it from "
                        + (f"{len(missing)} row(s) the slide drops "
                           f"({'…; ' if len(missing) > 3 else ''}{named})" if missing
                           else "rows shown in another order")
                        + ". Show the rows it derives from, or drop the derived row")
        prev, acc = pos, []
    return None


def squash(s: str) -> str:
    """A header word compared on its letters and digits only: `sub-group`, `Sub group`
    and `subgroup` are one header."""
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def column_index(t: Table, header: str) -> int | None:
    """The column headed `header`: the exact text, the squashed text, or the one header
    that contains it (`verdict` finds `verdict (sell-side)`)."""
    hits = [i for i, h in enumerate(t.headers) if h == header] or \
           [i for i, h in enumerate(t.headers) if squash(h) == squash(header)]
    if not hits:
        part = [i for i, h in enumerate(t.headers) if squash(header) in squash(h)]
        hits = part if len(part) == 1 else []
    return hits[0] if hits else None


def table_block(v, at: str, res: Resolver, book: Book) -> dict:
    if not isinstance(v, dict) or "from" not in v:
        raise SpecError(f"{at}: table needs `from: <tab>`")
    known(v, TABLE_KEYS, at, "table")
    tab = book.tab(v["from"])
    t = book.block(tab, v["block"]) if v.get("block") else book.primary(tab)
    if not isinstance(t, Table):
        raise SpecError(f"{at}: `{v.get('block')}` on `{tab}` is lines, not a table")
    t = deepcopy(t)
    if v.get("title"):
        t.title = res.resolve(v["title"], at)
    elif v.get("title") is None and "block" not in v:
        t.title = None
    keep_ids = bool(v.get("ids", False))
    cols = v.get("columns")
    if cols is not None:
        if not isinstance(cols, list) or not cols:
            raise SpecError(f"{at}: columns is a list of headers")
        idx = []
        for c in cols:
            hits = [i for i, h in enumerate(t.headers) if h == c] or \
                   [i for i, h in enumerate(t.headers) if normalize(h) == normalize(c)]
            if not hits:
                raise SpecError(f"{at}: `{tab}` has no column headed `{c}` "
                                f"(columns: {', '.join(t.headers)})")
            idx.append(hits[0])
    else:
        idx = [i for i, h in enumerate(t.headers) if keep_ids or h.strip().lower() not in ID_HEADERS]
        if not idx:
            idx = list(range(len(t.headers)))
    rows_sel = v.get("rows")
    orig_rows, picked_idx = list(t.rows), list(range(len(t.rows)))
    where = v.get("where")
    if where is not None:
        # `where: {verdict: supported}` — a row carrying a value in that column stays when
        # the value matches; a row with the column empty (walk mechanics, a subtotal, an
        # information line) is not a ruled item and stays regardless. `through` then ends
        # the schedule at a row's label, so the information lines under a walk drop.
        if not isinstance(where, dict) or not where:
            raise SpecError(f"{at}: where maps a column header to a value or a list of values")
        keep = list(range(len(t.rows)))
        for header, wanted in where.items():
            ci = column_index(t, str(header))
            if ci is None:
                raise SpecError(f"{at}: `{tab}` has no column headed `{header}` "
                                f"(columns: {', '.join(t.headers)})")
            allowed = {normalize(str(x)) for x in (wanted if isinstance(wanted, list) else [wanted])}
            keep = [j for j in keep if not str(t.rows[j][ci].value or "").strip()
                    or normalize(str(t.rows[j][ci].value)) in allowed]
        picked_idx = keep
    through = v.get("through")
    if through is not None:
        stop = next((n for n, j in enumerate(picked_idx)
                     if any(isinstance(c.value, str) and (c.value.strip() == str(through).strip()
                                                          or normalize(c.value) == normalize(str(through)))
                            for c in t.rows[j][:LABEL_COLS + 1])), None)
        if stop is None:
            raise SpecError(f"{at}: `{tab}` has no row labelled `{through}` to end the schedule at")
        picked_idx = picked_idx[:stop + 1]
    if where is not None or through is not None:
        t.rows = [orig_rows[j] for j in picked_idx]
        t.kinds = [t.kinds[j] for j in picked_idx]
    if rows_sel is not None:
        if not isinstance(rows_sel, list) or not rows_sel:
            raise SpecError(f"{at}: rows is a list of leading labels")
        picked, kinds, picked_idx = [], [], []
        for label in rows_sel:
            label = str(label).strip()
            hit = next((j for j, row in enumerate(t.rows)
                        if any(isinstance(c.value, str) and c.value.strip() == label for c in row[:LABEL_COLS + 1])), None)
            if hit is None:
                hit = next((j for j, row in enumerate(t.rows)
                            if any(isinstance(c.value, str) and normalize(c.value) == normalize(label)
                                   for c in row[:LABEL_COLS + 1])), None)
            if hit is None:
                raise SpecError(f"{at}: `{tab}` has no row labelled `{label}`")
            picked.append(t.rows[hit])
            kinds.append(t.kinds[hit])
            picked_idx.append(hit)
        t.rows, t.kinds = picked, kinds
    max_rows = v.get("max_rows")
    if max_rows is not None:
        max_rows = int(max_rows)
        if len(t.rows) > max_rows:
            t.more = len(t.rows) - max_rows
            t.rows, t.kinds = t.rows[:max_rows], t.kinds[:max_rows]
            picked_idx = picked_idx[:max_rows]
    t.headers = [t.headers[i] for i in idx]
    t.numeric = [t.numeric[i] for i in idx]
    t.rows = [[row[i] for i in idx] for row in t.rows]
    scale = v.get("scale")
    if scale is not None:
        if scale not in SCALES:
            raise SpecError(f"{at}: scale is {' or '.join(SCALES)}, not `{scale}`")
        factor, suffix = SCALES[scale]
        scaled = 0
        for i, is_num in enumerate(t.numeric):
            cells = [row[i] for row in t.rows if isinstance(row[i].value, (int, float))
                     and not isinstance(row[i].value, bool)]
            if not is_num or not cells or not all(is_money(c.fmt) for c in cells) \
                    or is_count(t.headers[i], cells):
                continue
            t.headers[i] = f"{t.headers[i]} ({suffix})"
            # Divide, never round: the stored value keeps full precision, and every check
            # downstream (not_footing below, and the deck gate's match against the workbook)
            # computes on these values. Rounding happens once, in fmt_value, on the way to
            # the slide - nothing reads a rendered number back.
            for row in t.rows:
                if isinstance(row[i].value, (int, float)) and not isinstance(row[i].value, bool):
                    row[i] = replace(row[i], value=row[i].value / factor, fmt=SCALED_FMT)
            scaled += 1
        if not scaled:
            raise SpecError(f"{at}: `{tab}` has no dollar column to state in {scale}")
    if rows_sel is not None or where is not None or through is not None:
        why = not_footing(t.rows, picked_idx, orig_rows, t.numeric, t.kinds)
        if why:
            raise SpecError(f"{at}: {why} (`{tab}`)")
    t.dense = bool(v.get("dense", False))
    res.use(tab)
    res.count += sum(1 for row in t.rows for c in row if isinstance(c.value, (int, float)))
    return {"t": "table", "table": t, "fit": v.get("fit", True)}


def chart_block(v, at: str, res: Resolver, book: Book) -> dict:
    if not isinstance(v, dict) or "from" not in v or "rows" not in v:
        raise SpecError(f"{at}: chart needs `from: <tab>` and `rows: [labels]`")
    known(v, CHART_KEYS, at, "chart")
    ctype = str(v.get("type", "column")).lower()
    if ctype not in ("column", "bar", "line"):
        raise SpecError(f"{at}: chart type is column, bar or line")
    tb = table_block({"from": v["from"], "rows": v["rows"], "ids": True,
                      **({"block": v["block"]} if v.get("block") else {})}, at, res, book)
    t: Table = tb["table"]
    cat_idx = [i for i, n in enumerate(t.numeric) if n]
    if v.get("columns"):
        want = [normalize(str(c)) for c in v["columns"]]
        cat_idx = [i for i in cat_idx if normalize(t.headers[i]) in want]
        missing = [c for c in v["columns"] if normalize(str(c)) not in {normalize(t.headers[i]) for i in cat_idx}]
        if missing:
            raise SpecError(f"{at}: `{t.source}` has no numeric column headed {', '.join(map(str, missing))}")
    if not cat_idx:
        raise SpecError(f"{at}: the chosen rows carry no numeric columns to chart")
    categories = [t.headers[i] for i in cat_idx]
    series = []
    for row in t.rows:
        label = next((str(c.value).strip() for j, c in enumerate(row)
                      if isinstance(c.value, str) and c.value.strip() and not t.numeric[j]
                      and t.headers[j].strip().lower() not in ID_HEADERS), "")
        vals = [row[i].value if isinstance(row[i].value, (int, float)) else 0 for i in cat_idx]
        series.append((label, vals))
    if len(series) > 4:
        raise SpecError(f"{at}: at most four series in one chart (WORKBOOK_STYLE.md § 6)")
    return {"t": "chart", "type": ctype, "categories": categories, "series": series,
            "title": res.resolve(v.get("title", ""), at), "source": t.source,
            "fmt": next((row[cat_idx[0]].fmt for row in t.rows), "")}


# --- layout: measure --------------------------------------------------------------
def col_widths(t: Table, width: float, pt: float) -> list[float]:
    n = len(t.headers)
    if n == 0:
        return []
    need = []
    for i in range(n):
        longest = max([len(t.headers[i])] + [len(fmt_value(r[i].value, r[i].fmt)) for r in t.rows] or [4])
        if t.numeric[i]:
            need.append(min(max(0.8, longest * pt * 0.55 / 72 + 0.2), 1.5))
        else:
            need.append(min(max(0.9, longest * pt * 0.52 / 72 + 0.2), 4.2))
    total = sum(need)
    if total > width:
        scale = width / total
        need = [max(w * scale, 0.55) for w in need]
    else:
        text_cols = [i for i in range(n) if not t.numeric[i]] or list(range(n))
        extra = (width - total) / len(text_cols)
        need = [w + (extra if i in text_cols else 0) for i, w in enumerate(need)]
    return need


def table_pt(t: Table) -> float:
    if t.dense:
        return PT["table_xs"]
    return PT["table_sm"] if len(t.headers) > 6 or len(t.rows) > 14 else PT["table"]


def header_pt(t: Table) -> float:
    return PT["heading_xs"] if t.dense else PT["heading"]


# A body row is its lines of text plus the row spacing — the cell's top and bottom margin
# — and nothing else. The deck declares every body row at one line; PowerPoint grows a
# row whose text wraps, so a row is never taller than its text needs. The line count
# estimated here (text_w leans generous) only decides how many rows a page takes.
ROW_SPACING = {True: 0.03, False: 0.045}     # inches, top and bottom margin; dense / not
LINE_H = 1.2                                 # a line of table text, in multiples of its size


def row_min(t: Table, pt: float) -> float:
    """A one-line body row: the height every body row is declared at."""
    return pt * LINE_H / 72 + 2 * ROW_SPACING[t.dense]


def row_heights(t: Table, widths: list[float], pt: float) -> tuple[float, list[float]]:
    pad = 0.10 if t.dense else 0.14            # cell margins 0.05 top and bottom, and slack
    hh = max(text_h(h.upper(), w - 0.14, header_pt(t), bold=True, spacing=1.15) for h, w in zip(t.headers, widths)) + pad + RULE_H
    rows = []
    for row in t.rows:
        n = max(lines_for(fmt_value(c.value, c.fmt), w - 0.14, pt) for c, w in zip(row, widths))
        rows.append(n * pt * LINE_H / 72 + 2 * ROW_SPACING[t.dense])
    return hh, rows



def header_h(title: str, message: str = "") -> float:
    """The headline's line box, and the message's beneath it when the page carries one."""
    h = text_h(title, BODY_W, PT["title"], bold=True, spacing=1.1)
    if message:
        h += MESSAGE_GAP + text_h(message, BODY_W, PT["message"], spacing=1.3)
    return h


def body_top(title: str, message: str = "") -> float:
    """Where the body starts under this header: the c4 body top under a one-line
    headline, lower under a two-line one or a message."""
    return max(BODY_Y, TITLE_Y + header_h(title, message) + HEADER_GAP)


def label_h(text: str, width: float) -> float:
    # tracked caps run ~10% wider than the plain estimate
    return text_h(text.upper(), width * 0.85, PT["heading"], bold=True) + 0.04


def measure(block: dict, width: float) -> float:
    t = block["t"]
    if t == "text":
        pt = PT["lead"] if block.get("lead") else PT["small"] if block.get("small") else PT["body"]
        return text_h(block["text"], width, pt, spacing=1.35 if block.get("lead") else 1.3)
    if t == "heading":
        return label_h(block["text"], width)
    if t == "note":
        return text_h(block["text"], width, PT["note"], spacing=1.25)
    if t == "bullets":
        pt = PT["small"] if block.get("small") else PT["body"]
        return sum(text_h(s, width - 0.26, pt, spacing=1.3) + 0.05 for s in block["items"])
    if t == "stats":
        n = len(block["tiles"])
        tile_w = (width - 0.4 * (n - 1)) / n
        return max(RULE_H + 0.16 + label_h(x["label"], tile_w) + 0.06 + text_h(x["value"], tile_w, PT["stat_value"], bold=True, spacing=1.1)
                   + (0.04 + text_h(x["note"], tile_w, PT["small"], spacing=1.25) if x["note"] else 0) for x in block["tiles"])
    if t == "kv":
        lw = min(2.4, width * 0.3)
        return sum(max(label_h(k, lw - 0.1), text_h(v, width - lw - 0.1, PT["body"], spacing=1.3)) + 0.08
                   for k, v in block["items"])
    if t == "table":
        tb: Table = block["table"]
        pt = table_pt(tb)
        widths = col_widths(tb, width, pt)
        hh, rows = row_heights(tb, widths, pt)
        h = hh + sum(rows)
        if tb.title:
            h += label_h(tb.title, width) + 0.06
        if tb.more:
            h += 0.24
        return h
    if t == "chart":
        return CHART_H + (label_h(block["title"], width) + 0.06 if block.get("title") else 0)
    if t == "columns":
        ws = [width * s - COL_GAP * (len(block["cols"]) - 1) / len(block["cols"]) for s in block["widths"]]
        return max(stack_h(col, w) for col, w in zip(block["cols"], ws))
    return 0.0


def stack_h(blocks: list[dict], width: float) -> float:
    if not blocks:
        return 0.0
    return sum(measure(b, width) for b in blocks) + GAP * (len(blocks) - 1)


def fit_table(block: dict, width: float, avail: float) -> None:
    """Trim a table's rows so the block fits `avail`, stating how many rows stay on the
    tab. Never below three rows: a schedule shown as fewer says nothing."""
    tb: Table = block["table"]
    if not block.get("fit", True):
        return
    while len(tb.rows) > 3 and measure(block, width) > avail:
        tb.rows.pop()
        tb.kinds.pop()
        tb.more += 1


# --- flow: split long blocks over continuation pages ----------------------------
def flow(page: Page) -> list[Page]:
    """The page, or the page and its continuations when its blocks measure past the
    body. Only tables and bullet lists split; anything else overflowing is a refusal.
    A continuation carries the headline marked `(continued)` and no message — the
    message was stated where the page began — so its body starts where that header ends."""
    pages: list[Page] = []
    pending = list(page.blocks)
    cur: list[dict] = []
    used = 0.0
    n = 0
    avail_first = BODY_BOTTOM - body_top(page.title, page.message)
    avail_more = BODY_BOTTOM - body_top(page.title + CONTINUED)

    def emit():
        nonlocal cur, used, n
        p = deepcopy(page)
        p.blocks = cur
        if n:
            p.title = page.title + CONTINUED
            p.message = ""
        pages.append(p)
        cur, used, n = [], 0.0, n + 1

    while pending:
        avail = avail_first if n == 0 else avail_more
        b = pending.pop(0)
        h = measure(b, BODY_W)
        gap = GAP if cur else 0
        if used + gap + h <= avail:
            cur.append(b)
            used += gap + h
            continue
        room = avail - used - gap
        # A split measures its rows and items exactly as measure() does, and always
        # leaves a remainder: a split that keeps everything would emit an empty
        # continuation page.
        if b["t"] == "table" and len(b["table"].rows) > 2:
            tb: Table = b["table"]
            pt = table_pt(tb)
            widths = col_widths(tb, BODY_W, pt)
            hh, rows = row_heights(tb, widths, pt)
            head = hh + (label_h(tb.title, BODY_W) + 0.06 if tb.title else 0)
            k, acc = 0, head
            while k < len(rows) and acc + rows[k] <= room:
                acc += rows[k]
                k += 1
            if 2 <= k < len(rows):
                first = deepcopy(b)
                first["table"].rows, first["table"].kinds = tb.rows[:k], tb.kinds[:k]
                first["table"].more = 0
                rest = deepcopy(b)
                rest["table"].rows, rest["table"].kinds = tb.rows[k:], tb.kinds[k:]
                rest["table"].title = (tb.title if tb.title.endswith(CONTINUED)
                                       else tb.title + CONTINUED) if tb.title else None
                cur.append(first)
                pending.insert(0, rest)
                emit()
                continue
        if b["t"] == "bullets" and len(b["items"]) > 1:
            items = b["items"]
            pt = PT["small"] if b.get("small") else PT["body"]
            k, acc = 0, 0.0
            while k < len(items) and acc + text_h(items[k], BODY_W - 0.26, pt, spacing=1.3) + 0.05 <= room:
                acc += text_h(items[k], BODY_W - 0.26, pt, spacing=1.3) + 0.05
                k += 1
            if 1 <= k < len(items):
                cur.append({**b, "items": items[:k]})
                pending.insert(0, {**b, "items": items[k:]})
                emit()
                continue
        if cur:
            pending.insert(0, b)
            emit()
            continue
        raise SpecError(f"page `{page.title[:60]}`: a {b['t']} block measures {h:.1f}in, "
                        f"over the {avail:.1f}in body — split it across pages or trim it")
    if cur or not pages:
        emit()
    return pages


# --- render -----------------------------------------------------------------------
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, name: str):
        s = self.prs.slides.add_slide(self.blank)
        s.name = name
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BONE)
        return s

    # primitives
    def rect(self, s, x, y, w, h, color, name="rect"):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.name = name
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def connector(self, s, x1, y1, x2, y2, color, name="line", pt=2.0):
        cx = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cx.name = name
        cx.line.color.rgb = rgb(color)
        cx.line.width = Pt(pt)
        return cx

    def dot(self, s, cx, cy, color, name="dot", d=CHART_DOT):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
        sh.name = name
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def text(self, s, x, y, w, h, runs, name="text", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=1.2, bullets=False, space_after=0, track=0.0):
        """`runs`: a list of paragraphs; a paragraph is a string or a list of
        (text, pt, bold, color[, italic]) tuples. `track` is letter-spacing in em."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = anchor
        for i, para in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            if space_after:
                p.space_after = Pt(space_after)
            parts = para if isinstance(para, list) else [para]
            for part in parts:
                if isinstance(part, str):
                    part = (part, PT["body"], False, BODY, False)
                txt, pt, bold, color, italic = (list(part) + [False])[:5]
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = FONT
                f.size = Pt(pt)
                f.bold = bold
                f.italic = italic
                f.color.rgb = rgb(color)
                if track:
                    r._r.get_or_add_rPr().set("spc", str(int(round(track * pt * 100))))
            if bullets:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(Inches(0.22)))
                pPr.set("indent", str(-Inches(0.22)))
                clr = pPr.makeelement(qn("a:buClr"), {})
                clr.append(clr.makeelement(qn("a:srgbClr"), {"val": DOT}))
                pPr.append(clr)
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        return tb

    def label(self, s, x, y, w, text, name="body-heading", color=MUTED):
        """A c4 label: 9pt bold caps, tracked, muted."""
        h = label_h(text, w)
        self.text(s, x, y, w, h, [[(text.upper(), PT["heading"], True, color)]], name, track=TRACK["heading"])
        return h

    def band(self, s, tagline: str, right_lines: list[str], names=("footer-left", "footer-source")):
        """The c4 footer band: a bold white tagline on the left, small white lines on
        the right. Each right line is its own shape, named, so the gate can read the
        source line."""
        self.rect(s, 0, FOOTER_Y, SLIDE_W, FOOTER_H, BAND, "footer-band")
        right_w = 4.4 if tagline else BODY_W
        rx = SLIDE_W - MARGIN - right_w
        if tagline:
            size = PT["tagline"]
            while size > 12 and text_w(tagline, size, True) > rx - MARGIN - 0.3:
                size -= 1
            self.text(s, MARGIN, FOOTER_Y + 0.23, rx - MARGIN - 0.3, 0.3, [[(tagline, size, True, WHITE)]],
                      "footer-tagline", anchor=MSO_ANCHOR.MIDDLE)
        n = len(right_lines)
        top = FOOTER_Y + (0.29 if n == 1 else 0.17)
        for i, (line, name) in enumerate(zip(right_lines, names)):
            self.text(s, rx, top + i * 0.19, right_w, 0.18, [[(line, PT["footer"], False, WHITE)]], name,
                      align=PP_ALIGN.RIGHT)

    def footer(self, s, company: str, sources: list[str], number: int, tagline: str = ""):
        lines = [f"Confidential · Prepared by Countz for {company} · {number}"]
        if sources:
            lines.append("Source: workbook.xlsx · " + " · ".join(sources))
        self.band(s, tagline, lines)

    def header(self, s, kicker: str, title: str, message: str = ""):
        """The c4 page header: kicker, the headline at the display size and, when the
        page carries one, its message as a muted sentence beneath (the cover-subtitle
        treatment). Each is its own named shape, so the gate can read them."""
        self.text(s, MARGIN, KICKER_Y, BODY_W, KICKER_H, [[(kicker.upper(), PT["kicker"], True, TEAL)]], "kicker",
                  track=TRACK["kicker"])
        th = text_h(title, BODY_W, PT["title"], bold=True, spacing=1.1)
        self.text(s, MARGIN, TITLE_Y, BODY_W, th, [[(title, PT["title"], True, INK)]], "title",
                  anchor=MSO_ANCHOR.TOP, spacing=1.1)
        if message:
            mh = text_h(message, BODY_W, PT["message"], spacing=1.3)
            self.text(s, MARGIN, TITLE_Y + th + MESSAGE_GAP, BODY_W, mh, [[(message, PT["message"], False, MUTED)]],
                      "message", spacing=1.3)

    # blocks
    def draw(self, s, block: dict, x: float, y: float, w: float) -> float:
        t = block["t"]
        if t == "text":
            lead = block.get("lead")
            pt = PT["lead"] if lead else PT["small"] if block.get("small") else PT["body"]
            sp = 1.35 if lead else 1.3
            h = text_h(block["text"], w, pt, spacing=sp)
            self.text(s, x, y, w, h, [[(block["text"], pt, False, INK if lead else BODY)]], "body-text", spacing=sp)
            return h
        if t == "heading":
            return self.label(s, x, y, w, block["text"])
        if t == "note":
            h = text_h(block["text"], w, PT["note"], spacing=1.25)
            self.text(s, x, y, w, h, [[(block["text"], PT["note"], False, MUTED)]], "body-note", spacing=1.25)
            return h
        if t == "bullets":
            pt = PT["small"] if block.get("small") else PT["body"]
            h = sum(text_h(i, w - 0.26, pt, spacing=1.3) + 0.05 for i in block["items"])
            self.text(s, x, y, w, h, [[(i, pt, False, BODY)] for i in block["items"]],
                      "body-bullets", bullets=True, space_after=4, spacing=1.3)
            return h
        if t == "stats":
            n = len(block["tiles"])
            tw = (w - 0.4 * (n - 1)) / n
            hmax = 0
            for i, tile in enumerate(block["tiles"]):
                tx = x + i * (tw + 0.4)
                self.rect(s, tx, y, tw, RULE_H, TEAL, "stat-rule")
                yy = y + RULE_H + 0.16
                yy += self.label(s, tx, yy, tw, tile["label"], "stat-label") + 0.06
                vh = text_h(tile["value"], tw, PT["stat_value"], bold=True, spacing=1.1)
                self.text(s, tx, yy, tw, vh, [[(tile["value"], PT["stat_value"], True, INK)]], "stat-value", spacing=1.1)
                yy += vh
                if tile["note"]:
                    nh = text_h(tile["note"], tw, PT["small"], spacing=1.25)
                    self.text(s, tx, yy + 0.04, tw, nh, [[(tile["note"], PT["small"], False, MUTED)]], "stat-note", spacing=1.25)
                    yy += nh + 0.04
                hmax = max(hmax, yy - y)
            return hmax
        if t == "kv":
            lw = min(2.4, w * 0.3)
            yy = y
            for k, v in block["items"]:
                h = max(label_h(k, lw - 0.1), text_h(v, w - lw - 0.1, PT["body"], spacing=1.3))
                self.label(s, x, yy + 0.03, lw - 0.1, k, "kv-label")
                self.text(s, x + lw, yy, w - lw, h, [[(v, PT["body"], False, BODY)]], "kv-value", spacing=1.3)
                yy += h + 0.08
            return yy - y
        if t == "table":
            return self.table(s, block["table"], x, y, w)
        if t == "chart":
            return self.chart(s, block, x, y, w)
        if t == "columns":
            n = len(block["cols"])
            widths = [(w - COL_GAP * (n - 1)) * share for share in block["widths"]]
            xx = x
            hmax = 0
            for col, cw in zip(block["cols"], widths):
                h = self.stack(s, col, xx, y, cw)
                hmax = max(hmax, h)
                xx += cw + COL_GAP
            return hmax
        return 0.0

    def stack(self, s, blocks: list[dict], x: float, y: float, w: float) -> float:
        yy = y
        for i, b in enumerate(blocks):
            if i:
                yy += GAP
            yy += self.draw(s, b, x, yy, w)
        return yy - y

    def table(self, s, tb: Table, x: float, y: float, w: float) -> float:
        yy = y
        if tb.title:
            yy += self.label(s, x, yy, w, tb.title, "table-title", color=INK) + 0.06
        pt = table_pt(tb)
        widths = col_widths(tb, w, pt)
        hh, rows = row_heights(tb, widths, pt)
        n_rows, n_cols = len(tb.rows) + 1, len(tb.headers)
        # the frame is the declared rows, never the estimate: a renderer stretches the rows
        # to fill a taller frame, which spreads the estimate's slack over every row
        gf = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(yy), Inches(sum(widths)),
                                Inches(hh + (n_rows - 1) * row_min(tb, pt)))
        gf.name = f"table:{tb.source}"
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = tblPr.makeelement(qn("a:tableStyleId"), {})
            tblPr.append(sid)
        sid.text = TABLE_STYLE_NONE
        tbl.first_row = False
        tbl.horz_banding = False
        for i, cw in enumerate(widths):
            tbl.columns[i].width = Inches(cw)
        tbl.rows[0].height = Inches(hh)
        for i in range(1, n_rows):
            tbl.rows[i].height = Inches(row_min(tb, pt))
        # one alignment rule, header and body alike: the first column — the row's label —
        # reads left, every other column reads right, so figures and their headers share
        # an edge and a text cell in a figure column does not break the rag
        align = [PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT for j in range(n_cols)]
        for j, h in enumerate(tb.headers):
            # the c4 label over a 3 px teal rule, a hairline beneath
            self.cell(tbl.cell(0, j), h.upper(), header_pt(tb), bold=True, color=MUTED, align=align[j],
                      top=(TEAL, 28575), bottom=(RULE, 9525), track=TRACK["heading"], header=True)
        last = len(tb.rows)
        for i, (row, kind) in enumerate(zip(tb.rows, tb.kinds), 1):
            for j, c in enumerate(row):
                txt = fmt_value(c.value, c.fmt)
                key = txt.strip().lower()
                color = STATUS_COLOR.get(key, INK if kind != "body" else BODY)
                self.cell(tbl.cell(i, j), txt, pt, bold=(kind != "body"), color=color, align=align[j],
                          top=(INK, 12700) if kind == "total" else None,
                          bottom=(RULE, 9525) if i < last or kind != "total" else (INK, 12700))
        yy += hh + sum(rows)
        if tb.more:
            self.text(s, x, yy + 0.06, w, 0.18,
                      [[(f"{tb.more} more row(s) on the {tb.source} tab of workbook.xlsx", PT["note"], False, MUTED)]],
                      "table-more")
            yy += 0.24
        return yy - y

    def cell(self, cell, text: str, pt: float, bold=False, color=BODY, align=PP_ALIGN.LEFT,
             top=None, bottom=None, track=0.0, header=False):
        """No fill (the ground shows through), no vertical rules: a horizontal hairline
        under each row, the c4 statements pattern. `top` / `bottom`: (colour, EMU) or
        None for no rule."""
        tcPr = cell._tc.get_or_add_tcPr()
        for tag, spec in (("a:lnL", None), ("a:lnR", None), ("a:lnT", top), ("a:lnB", bottom)):
            if spec is None:
                ln = tcPr.makeelement(qn(tag), {"w": "0"})
                ln.append(ln.makeelement(qn("a:noFill"), {}))
            else:
                colour, width = spec
                ln = tcPr.makeelement(qn(tag), {"w": str(width), "cap": "flat", "cmpd": "sng"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": colour}))
                ln.append(sf)
            tcPr.append(ln)
        cell.fill.background()
        cell.margin_left = cell.margin_right = Inches(0.06)
        if header:
            cell.margin_top = cell.margin_bottom = Inches(0.03 if pt <= PT["table_xs"] else 0.05)
        else:
            cell.margin_top = cell.margin_bottom = Inches(ROW_SPACING[pt <= PT["table_xs"]])
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        # a blank cell is a single space at the table's size: an empty paragraph is sized
        # by the renderer's default (18pt in Keynote, whatever the end mark says), and
        # that one cell grows its whole row
        r = p.add_run()
        r.text = text or " "
        r.font.name = FONT
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        if track:
            r._r.get_or_add_rPr().set("spc", str(int(round(track * pt * 100))))
        end = p._p.get_or_add_endParaRPr()
        end.set("sz", str(int(round(pt * 100))))
        end.set("lang", "en-US")

    def chart(self, s, block: dict, x: float, y: float, w: float) -> float:
        """The chart, drawn.

        Every plotted value rides on its shape's name — `chartval:<tab>:<value>` — which
        is how check_report.py audits a drawn chart against the workbook, the same audit
        it ran over a chart part's cached values. The furniture (ticks, category labels,
        legend) is named so the gate reads it as the scale it is, not as a figure.
        """
        yy = y
        if block.get("title"):
            yy += self.label(s, x, yy, w, block["title"], "chart-title", color=INK) + 0.06
        cats, series = block["categories"], block["series"]
        n_cat, n_ser = max(len(cats), 1), max(len(series), 1)
        horiz = block["type"] == "bar"
        pt = PT["stat_label"]
        tag = block["source"]
        vals = [v for _, vs in series for v in vs] or [0.0]
        lo, hi, ticks = nice_axis(min(vals), max(vals))
        step = (ticks[1] - ticks[0]) if len(ticks) > 1 else 1.0
        span = (hi - lo) or 1.0

        legend_h = CHART_LEGEND_H if len(series) > 1 else 0.0
        gutter_text = cats if horiz else [axis_text(t, step) for t in ticks]
        gut = min(max([text_w(t, pt) for t in gutter_text] + [0.35]) + 0.10, w * 0.35)
        px, pw = x + gut, w - gut
        py, ph = yy, CHART_H - legend_h - CHART_AXIS_H

        def vx(v):                                  # a value along a horizontal value axis
            return px + pw * (v - lo) / span

        def vy(v):                                  # a value along a vertical value axis
            return py + ph * (1 - (v - lo) / span)

        # the scale: a hairline at every tick, its label in the gutter or the strip below
        for t in ticks:
            if horiz:
                self.connector(s, vx(t), py, vx(t), py + ph, RULE, "chart-grid", GRID_PT)
                self.text(s, vx(t) - 0.6, py + ph + 0.07, 1.2, 0.18,
                          [[(axis_text(t, step), pt, False, MUTED)]], "chart-axis", align=PP_ALIGN.CENTER)
            else:
                self.connector(s, px, vy(t), px + pw, vy(t), RULE, "chart-grid", GRID_PT)
                self.text(s, x, vy(t) - 0.09, gut - 0.10, 0.18,
                          [[(axis_text(t, step), pt, False, MUTED)]], "chart-axis", align=PP_ALIGN.RIGHT)
        if horiz:
            self.connector(s, vx(0.0), py, vx(0.0), py + ph, MUTED, "chart-base", BASE_PT)
        else:
            self.connector(s, px, vy(0.0), px + pw, vy(0.0), MUTED, "chart-base", BASE_PT)

        # the series
        if block["type"] == "line":
            slot = pw / n_cat
            for si, (nm, vs) in enumerate(series):
                colour = SERIES[si % len(SERIES)]
                pts = [(px + ci * slot + slot / 2, vy(vs[ci] if ci < len(vs) else 0.0)) for ci in range(n_cat)]
                for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
                    self.connector(s, x1, y1, x2, y2, colour, "chart-line")
                for ci, (cx, cy) in enumerate(pts):
                    v = vs[ci] if ci < len(vs) else 0.0
                    self.dot(s, cx, cy, colour, f"chartval:{tag}:{float(v)!r}")
        else:
            extent = ph if horiz else pw
            slot = extent / n_cat
            band = slot * (1 - CHART_BAR_GAP)
            thick = band / n_ser
            for ci in range(n_cat):
                head = (py if horiz else px) + ci * slot + (slot - band) / 2
                for si, (nm, vs) in enumerate(series):
                    v = vs[ci] if ci < len(vs) else 0.0
                    colour = SERIES[si % len(SERIES)]
                    nmn = f"chartval:{tag}:{float(v)!r}"
                    if horiz:
                        a, b = sorted((vx(0.0), vx(v)))
                        self.rect(s, a, head + si * thick, max(b - a, 0.01), thick * 0.9, colour, nmn)
                    else:
                        a, b = sorted((vy(0.0), vy(v)))
                        self.rect(s, head + si * thick, a, thick * 0.9, max(b - a, 0.01), colour, nmn)

        # the categories
        for ci, c in enumerate(cats):
            if horiz:
                slot = ph / n_cat
                self.text(s, x, py + ci * slot, gut - 0.10, slot,
                          [[(str(c), pt, False, MUTED)]], "chart-cat",
                          align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            else:
                slot = pw / n_cat
                self.text(s, px + ci * slot, py + ph + 0.07, slot, 0.18,
                          [[(str(c), pt, False, MUTED)]], "chart-cat", align=PP_ALIGN.CENTER)

        # the legend, centred under the plot
        if legend_h:
            sw, gap, pad = 0.13, 0.07, 0.30
            tw = [text_w(str(nm), pt) + 0.05 for nm, _ in series]
            total = sum(sw + gap + t for t in tw) + pad * (len(series) - 1)
            lx = x + max((w - total) / 2, 0.0)
            ly = yy + CHART_H - legend_h + 0.06
            for si, (nm, _) in enumerate(series):
                self.rect(s, lx, ly + 0.02, sw, sw, SERIES[si % len(SERIES)], "chart-swatch")
                self.text(s, lx + sw + gap, ly, tw[si], 0.18, [[(str(nm), pt, False, MUTED)]], "chart-legend")
                lx += sw + gap + tw[si] + pad
        return yy - y + CHART_H


# --- assembling the pages -----------------------------------------------------------
def load_run(run_dir: pathlib.Path) -> dict:
    try:
        return json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def build_pages(spec: dict, book: Book, run: dict, res: Resolver) -> tuple[list[Page], dict]:
    """The author's sections and pages, references resolved, tables copied. Every defect
    is collected and the whole list refused at once, so one build names them all."""
    errors: list[str] = []
    meta = {
        "title": res.resolve(str(spec.get("title") or run.get("goal") or "Report"), "title").strip(),
        "company": str(spec.get("company") or (run.get("inputs") or {}).get("company") or "").strip(),
        "date": str(spec.get("date") or dt.date.today().strftime("%-d %B %Y")).strip(),
        "sections": [],
    }
    if not meta["company"]:
        errors.append("no `company` — run.json carries none and the spec names none")
    basis = str(book.cell(EXEC, "B2").value or "") if EXEC in book.tabs else ""
    meta["subtitle"] = (res.resolve(str(spec["subtitle"]), "subtitle").strip() if spec.get("subtitle")
                        else cover_subtitle(basis, meta["company"]))
    why = not_cover_title(meta["title"], meta["company"])
    if why:
        errors.append(f"title: {why}")
    if spec.get("subtitle"):
        why = not_cover_subtitle(meta["subtitle"], meta["title"], meta["company"])
        if why:
            errors.append(f"subtitle: {why}")
    sections = spec.get("sections")
    if not isinstance(sections, list) or not sections:
        errors.append("`sections` — a list of {title, pages} — is the document (REPORT.md § 2)")
        sections = []
    pages: list[Page] = []
    for i, sec in enumerate(sections):
        where = f"sections[{i}]"
        if not isinstance(sec, dict) or not sec.get("title") or not isinstance(sec.get("pages"), list) or not sec["pages"]:
            errors.append(f"{where}: a section has a `title` and a non-empty `pages` list")
            continue
        stitle = res.resolve(str(sec["title"]), f"{where}.title").strip()
        meta["sections"].append(stitle)
        for j, pg in enumerate(sec["pages"]):
            try:
                pages.append(author_page(pg, stitle, f"{where}.pages[{j}]", res, book))
            except SpecError as exc:
                errors.append(str(exc))
    errors.extend(res.errors)
    if errors:
        raise SpecError("\n".join(errors))
    return pages, meta


def author_page(pg, section: str, where: str, res: Resolver, book: Book) -> Page:
    if not isinstance(pg, dict) or not pg.get("title"):
        raise SpecError(f"{where}: a page has a `title` (its headline) and `blocks`")
    res.touched = set()
    title = res.resolve(str(pg["title"]), f"{where}.title").strip()
    why = not_headline(title)
    if why:
        raise SpecError(f"{where}.title: {why}")
    message = res.resolve(str(pg.get("message") or ""), f"{where}.message").strip()
    if message:
        why = incomplete(message)
        if why:
            raise SpecError(f"{where}.message: {why}")
        if len(message) > MESSAGE_MAX:
            raise SpecError(f"{where}.message: {len(message)} characters; one sentence, at most {MESSAGE_MAX}")
        if lines_for(message, BODY_W, PT["message"]) > 2:
            raise SpecError(f"{where}.message: runs past two lines at {PT['message']}pt — one sentence, shorter")
    kicker = res.resolve(str(pg.get("kicker") or section), f"{where}.kicker").strip()
    tagline = res.resolve(str(pg.get("tagline") or ""), f"{where}.tagline").strip()
    blocks = as_blocks(pg.get("blocks"), res, f"{where}.blocks", book)
    if not blocks:
        raise SpecError(f"{where}: no blocks")
    sources = [book.tab(t) for t in (pg.get("source") or [])]
    for t in sorted(res.touched):
        if t not in sources:
            sources.append(t)
    return Page(section, kicker, title, blocks, sources, "page", tagline=tagline, message=message)


# --- the deck ----------------------------------------------------------------------
def render(pages: list[Page], meta: dict, out: pathlib.Path) -> dict:
    """The cover, then the author's pages in order. No contents page and no section
    dividers: a section is the kicker its pages carry, nothing more."""
    deck = Deck()
    company = meta["company"]
    numbered = [Page("", "", meta["title"], [], [], "cover")] + list(pages)
    for i, p in enumerate(numbered, 1):
        p.number = i

    # cover — the c4 hero: kicker, display title, a short rule, a caption, the band
    s = deck.slide("cover")
    deck.text(s, MARGIN, 2.24, BODY_W, 0.2, [[(company.upper(), PT["kicker"], True, TEAL)]], "cover-company",
              track=TRACK["kicker"])
    th = text_h(meta["title"], BODY_W, PT["cover_title"], bold=True, spacing=1.1)
    deck.text(s, MARGIN, 2.6, BODY_W, th, [[(meta["title"], PT["cover_title"], True, INK)]], "cover-title", spacing=1.1)
    yy = 2.6 + th + 0.25
    deck.rect(s, MARGIN, yy, 1.0, RULE_H, TEAL, "cover-rule")
    yy += 0.22
    deck.text(s, MARGIN, yy, BODY_W, 0.4, [[(meta["subtitle"], PT["cover_sub"], False, MUTED)]], "cover-subtitle")
    yy += 0.55
    deck.text(s, MARGIN, yy, BODY_W, 0.6,
              [[(f"Prepared for {company} by Countz · {meta['date']}", PT["caption"], False, MUTED)],
               [("Prepared over the records the company supplied. This report states the procedures "
                 "performed, the figures they produced and the differences found.",
                 PT["small"], False, MUTED)]],
              "cover-prepared", spacing=1.35)
    deck.band(s, "Every figure in this report is copied from workbook.xlsx, where it re-performs from the source files.",
              ["Confidential · Copyright © Countz"], names=("footer-left",))

    counts = {"pages": len(numbered), "tables": 0, "charts": 0}
    for p in pages:
        s = deck.slide(f"page:{p.title[:40]}")
        deck.header(s, p.kicker, p.title, p.message)
        deck.stack(s, p.blocks, MARGIN, body_top(p.title, p.message), BODY_W)
        counts["tables"] += sum(1 for b in walk(p.blocks) if b["t"] == "table")
        counts["charts"] += sum(1 for b in walk(p.blocks) if b["t"] == "chart")
        deck.footer(s, company, p.sources, p.number, p.tagline)
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(str(out))
    return counts


def walk(blocks):
    for b in blocks:
        yield b
        if b["t"] == "columns":
            for col in b["cols"]:
                yield from walk(col)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("run_dir", type=pathlib.Path)
    ap.add_argument("--spec", type=pathlib.Path, help="default <run_dir>/out/.staging/report.yaml, else out/report.yaml")
    ap.add_argument("--workbook", type=pathlib.Path, help="default <run_dir>/out/.staging/workbook.xlsx, else out/workbook.xlsx")
    ap.add_argument("--out", type=pathlib.Path, help="default <run_dir>/out/.staging/report.pptx")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    rd = a.run_dir
    if not (rd / "run.json").is_file():
        print(f"{rd}: no run.json", file=sys.stderr)
        return 2
    staging = rd / "out" / ".staging"
    spec_path = a.spec or next((p for p in (staging / "report.yaml", rd / "out" / "report.yaml") if p.is_file()), staging / "report.yaml")
    wb_path = a.workbook or next((p for p in (staging / "workbook.xlsx", rd / "out" / "workbook.xlsx") if p.is_file()), staging / "workbook.xlsx")
    out = a.out or staging / "report.pptx"
    for p in (spec_path, wb_path):
        if not p.is_file():
            print(f"{p}: not a file", file=sys.stderr)
            return 2
    try:
        spec = yaml.safe_load(spec_path.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as exc:
        print(f"{spec_path.name}: not valid YAML — {exc}", file=sys.stderr)
        return 2
    if not isinstance(spec, dict):
        print(f"{spec_path.name}: the top level is a mapping", file=sys.stderr)
        return 2
    book = Book(wb_path)
    res = Resolver(book)
    run = load_run(rd)
    try:
        pages, meta = build_pages(spec, book, run, res)
        flowed: list[Page] = []
        for p in pages:
            flowed.extend(flow(p))
        counts = render(flowed, meta, out)
    except SpecError as exc:
        msgs = str(exc).splitlines()
        if a.json:
            print(json.dumps({"ok": False, "errors": msgs}, indent=2))
        else:
            print(f"{spec_path.name}: {len(msgs)} refusal(s) — no deck written.\n")
            for m in msgs:
                print(f"    {m}")
            print("\n  Fix: reference every figure as {tab | row | column} or {tab!B3} copied from "
                  "the workbook, name tabs and blocks as the workbook spells them, and split a page "
                  "that does not fit (REPORT.md § 2).")
        return 1
    rep = {"ok": True, "deck": str(out), "pages": counts["pages"], "tables": counts["tables"],
           "charts": counts["charts"], "figures_copied": res.count, "sources": sorted(res.sources)}
    if a.json:
        print(json.dumps(rep, indent=2))
    else:
        print(f"{out.name}: {counts['pages']} pages, {counts['tables']} tables, {counts['charts']} charts, "
              f"{res.count} figures copied from {len(res.sources)} tab(s) of {wb_path.name}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
