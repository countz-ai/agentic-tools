#!/usr/bin/env python3
"""Self-test for the report deck: build_report.py renders a deck from a spec-conformant
workbook, check_report.py passes it, and each refuses what it exists to refuse.

Builds, in a temporary directory, a run whose workbook is written to WORKBOOK.md +
WORKBOOK_STYLE.md (the kit of § 8) and passes link_workbook.py + check_workbook.py, so
the deck is tested over a workbook the plugin itself would seal. Then:

  1. a good report.yaml builds, and the deck passes check_report.py; on the deck as
     stored, a page's headline and its message are separate shapes, a page written with
     no message carries none, and every table reads its first column left and every
     other column right;
  2. a figure typed into a sentence builds (the builder cannot know) and is refused by
     the gate, named;
  3. a reference the workbook cannot resolve is refused by the builder, named;
  4. a deck whose figure was edited after the build is refused by the gate;
  4a. a chart is DRAWN, not placed as a chart part — the deck holds no `ppt/charts/`
     part — and its plotted values are still audited: a value edited on the shape that
     carries it is refused by the gate, named;
  5. a sentence where a headline belongs — a title over the cap, or ending in a full
     stop — is refused by the builder, named;
  6. a fragment where a sentence belongs — a text block with no full stop — is refused
     by the builder, named;
  7. a cover title carrying the company or the period is refused by the builder, named,
     and refused again by the gate on a deck built before the rule;
  8. a zero in a sentence reads `$0` — the en dash is the table's zero, and a sentence
     that trails off in one states nothing;
  9. a schedule condensed past its own arithmetic — a derived line shown over rows that
     do not make it — is refused by the builder, naming the rows it drops;
 10. the deck's own number conventions (REPORT.md § 4): a dollar figure in a sentence or
     a stat tile is scaled and rounded, a schedule declared `scale: thousands` is headed
     and divided, and the gate still finds the workbook cell behind each;
 11. a schedule declared `dense: true` is set at the dense size, and `where:` with
     `through:` keeps a walk's mechanics and its ruled rows and ends it at the closing
     line — the information line under it is not on the deck;
 13. the opening (REPORT.md § 1) is held by the gate: an executive summary page that is
     not first, or carries no figure block, a key-metrics page not headed as the recipe
     declares, and a first schedule two pages after it are each refused, named.
 12. the recipe's schedule (RECIPE_FORMAT.md § Report) is held by the gate: a deck whose
     walk is trimmed to `max_rows`, or carries no table from the walk's tab, is refused
     naming the schedule and the rows it lacks.

Run by check-plugin.py (8q) as `uv run --project <plugin> python3 scripts/report-selftest.py`.
Exit 0 when every case holds, 1 otherwise, naming the case.
"""
from __future__ import annotations

import html
import json
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent

# --- the kit: scripts/wbkit.py (WORKBOOK_STYLE.md § 9 + WORKBOOK.md § 7) ------------
sys.path.insert(0, str(HERE))
from wbkit import (ACCENT, BAND, FMT_AMOUNT, FMT_TEXT, MIST, S, SLATE, TINT,  # noqa: E402,F401
                   WIDTH, amount, band, finish, fit_rows, grid, header, ident, section,
                   status, text)


# --- the fixture run ----------------------------------------------------------------
PERIODS = ["FY2023", "FY2024", "LTM Jul 2025"]
SLUGS = ["fy2023", "fy2024", "ltm_2025-07"]
SUB = "Acme Corp · FY2023–LTM Jul 2025 · accrual · USD whole dollars · tolerance 1"
BRIDGE = [
    ("EBIT", [5_610_000, 6_120_000, 6_890_000], "Body"),
    ("+ D&A", [1_040_000, 1_120_000, 1_180_000], "Body"),
    ("= Reported EBITDA", [6_650_000, 7_240_000, 8_070_000], "Subtotal"),
    ("Non-recurring, supported", [420_000, 310_000, 264_000], "Body"),
    ("Normalization, supported", [-180_000, -95_000, 70_000], "Body"),
    ("Owner compensation, no adjustment taken", [0, 0, 0], "Body"),
    ("= Diligence adjusted EBITDA", [6_890_000, 7_455_000, 8_404_000], "Total"),
    ("Management adjusted EBITDA (information)", [7_800_000, 8_500_000, 9_650_000], "Body"),
]
POSITION = ("Diligence-adjusted EBITDA is 8,404,000 for LTM Jul 2025, 1,246,000 below "
            "management's adjusted figure.")


def make_run(rd: pathlib.Path) -> None:
    figures: dict[str, tuple] = {}
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Exec Summary")
    band(ws, "Exec Summary · quality of earnings", "Acme Corp · FY2023–LTM Jul 2025 · accrual · USD whole dollars",
         summary=POSITION)
    ws.freeze_panes = "B4"
    text(ws.cell(row=4, column=2), "EBITDA bridge — copied from the q6 EBITDA bridge tab", "Section")
    header(ws, 5, ["line"] + PERIODS, ["description", "period", "period", "period"], primary=False)
    r = 5
    for label, vals, st in BRIDGE:
        r += 1
        text(ws.cell(row=r, column=2), label, st)
        for j, v in enumerate(vals):
            amount(ws.cell(row=r, column=3 + j), v, style=st if st != "Body" else None)
    grid(ws, 5, r, 2, 5)
    fit_rows(ws, 4)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT

    ws = wb.create_sheet("q6 EBITDA bridge")
    band(ws, "q6 · EBIT walks to diligence adjusted EBITDA", SUB,
         summary="Diligence adjusted EBITDA is 8,404,000 LTM Jul 2025; management's figure is 1,246,000 higher.")
    header(ws, 4, ["id", "line", *PERIODS, "verdict", "note"],
           ["id", "description", "period", "period", "period", "status", "note"])
    r = 4
    for label, vals, st in BRIDGE:
        r += 1
        slug = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
        ident(ws.cell(row=r, column=2), f"F.q6.{slug}.ltm_2025-07")
        text(ws.cell(row=r, column=3), label, st)
        for j, v in enumerate(vals):
            amount(ws.cell(row=r, column=4 + j), v, style=st if st != "Body" else None)
            figures[f"F.q6.{slug}.{SLUGS[j]}"] = (label, v, "q6_bridge")
        status(ws.cell(row=r, column=7), "supported")
        text(ws.cell(row=r, column=8), "walk mechanics" if st != "Body" else "the rung's supported subtotal")
    last = r
    r += 2
    section(ws, r, "Exceptions")
    r += 1
    header(ws, r, ["id", "item", "amount", "owner", "what would clear it"],
           ["id", "description", "amount", "status", "note"], primary=False)
    hx = r
    r += 1
    ident(ws.cell(row=r, column=2), "X.q6.mgmt_residual")
    text(ws.cell(row=r, column=3), "Management adjusted EBITDA does not reconcile to its schedule")
    amount(ws.cell(row=r, column=4), 54_000)
    text(ws.cell(row=r, column=5), "CFO")
    text(ws.cell(row=r, column=6), "Management's revised schedule")
    grid(ws, hx, r, 2, 6)
    r += 2
    section(ws, r, "Notes")
    r += 1
    ident(ws.cell(row=r, column=2), "F.q6.ebit.ltm_2025-07")
    text(ws.cell(row=r, column=3), "Source: the income statement and the general ledger, as tied at q1.", "Note")
    r += 2
    section(ws, r, "To reperform")
    text(ws.cell(row=r + 1, column=2), "1. Read pl.xlsx · IS · rows 5-40, column E; sum to EBIT; add D&A.")
    finish(ws, last)

    ws = wb.create_sheet("Basis of Preparation")
    band(ws, "Basis of Preparation", "Acme Corp · FY2023–LTM Jul 2025 · accrual · USD whole dollars")
    header(ws, 4, ["source", "what it is", "class", "periods"], ["description", "description", "status", "description"])
    for i, row in enumerate([("gl", "General ledger export, gl.csv", "system-of-record", "Jan 2023 – Jul 2025"),
                             ("pl", "Income statements as presented, pl.xlsx", "management-prepared", "FY2023, FY2024, monthly")], 5):
        for j, v in enumerate(row):
            text(ws.cell(row=i, column=2 + j), v)
    grid(ws, 4, 6, 2, 5)
    section(ws, 8, "Procedures not performed")
    text(ws.cell(row=9, column=2), "The capex bridge was dropped: the room carries no capitalized-cost accounts.")
    section(ws, 11, "How to read this workbook")
    header(ws, 12, ["prefix", "meaning"], ["status", "description"], primary=False)
    text(ws.cell(row=13, column=2), "F.")
    text(ws.cell(row=13, column=3), "a figure — resolves on its Sources row")
    grid(ws, 12, 13, 2, 3)
    finish(ws, 6)

    ws = wb.create_sheet("q1 FY2023 statements")
    band(ws, "q1 · The FY2023 income statement agrees with the trial balance", SUB,
         summary="Net income and EBIT tie to the TB within tolerance; no exceptions.")
    header(ws, 4, ["id", "what is agreed", "side A", "source A", "amount A", "side B", "source B", "amount B", "difference", "status"],
           ["id", "description", "description", "description", "amount", "description", "description", "amount", "amount", "status"])
    r = 4
    for slug, what, a in [("net_income", "Net income FY2023", 3_912_000), ("ebit", "EBIT FY2023", 5_610_000)]:
        r += 1
        figures[f"F.q1.{slug}.fy2023"] = (what, a, "q1_fy2023")
        ident(ws.cell(row=r, column=2), f"F.q1.{slug}.fy2023")
        text(ws.cell(row=r, column=3), what)
        text(ws.cell(row=r, column=4), "income statement")
        text(ws.cell(row=r, column=5), "pl.xlsx · IS · C12")
        amount(ws.cell(row=r, column=6), a, hard_input=True)
        text(ws.cell(row=r, column=7), "trial balance")
        text(ws.cell(row=r, column=8), "tb.xlsx · TB · P&L close")
        amount(ws.cell(row=r, column=9), a, hard_input=True)
        amount(ws.cell(row=r, column=10), 0)
        status(ws.cell(row=r, column=11), "tied")
    last = r
    r += 2
    section(ws, r, "Notes")
    text(ws.cell(row=r + 1, column=2), "Population: every line of the income statement, 31 of 31.", "Note")
    r += 3
    section(ws, r, "To reperform")
    text(ws.cell(row=r + 1, column=2), "1. Read pl.xlsx · IS · column C; read tb.xlsx · TB; difference = A − B.")
    finish(ws, last)

    ws = wb.create_sheet("Coverage")
    band(ws, "Coverage", "Acme Corp · one row per rostered check")
    header(ws, 4, ["token · title", "kind", "step", "status", "what was examined", "what was not examined, and why"],
           ["description", "status", "status", "status", "note", "note"])
    for i, row in enumerate([("q1 · FY2023 statements", "tieout", "q1_fy2023", "complete", "31 income-statement lines, FY2023", "—"),
                             ("q6 · EBITDA bridge", "analysis", "q6_bridge", "complete", "the bridge over 3 periods", "the capex bridge")], 5):
        for j, v in enumerate(row):
            text(ws.cell(row=i, column=2 + j), v)
    finish(ws, 6)

    ws = wb.create_sheet("Open Items")
    band(ws, "Open Items", "Acme Corp · review calls, questions for management, data requests")
    header(ws, 4, ["id", "matter", "size", "what closes it", "owner", "raised by"],
           ["id", "description", "amount", "note", "status", "id"])
    ident(ws.cell(row=5, column=2), "Q.q6.mgmt_residual")
    text(ws.cell(row=5, column=3), "Q.q6.mgmt_residual. Which schedule version produced the published figure?")
    amount(ws.cell(row=5, column=4), 54_000)
    text(ws.cell(row=5, column=5), "The schedule version")
    text(ws.cell(row=5, column=6), "CFO")
    text(ws.cell(row=5, column=7), "")
    grid(ws, 4, 5, 2, 7)
    section(ws, 7, "Data requests")
    header(ws, 8, ["id", "matter", "size", "what closes it", "owner", "raised by"],
           ["id", "description", "amount", "note", "status", "id"], primary=False)
    ident(ws.cell(row=9, column=2), "D.q6.capex")
    text(ws.cell(row=9, column=3), "D.q6.capex. The capitalized-cost accounts, for the capex bridge.")
    amount(ws.cell(row=9, column=4), 0)
    text(ws.cell(row=9, column=5), "The rollforward")
    text(ws.cell(row=9, column=6), "Controller")
    text(ws.cell(row=9, column=7), "")
    grid(ws, 8, 9, 2, 7)
    finish(ws, 5)
    ws.sheet_properties.tabColor = "9A5B00"

    ws = wb.create_sheet("Sources")
    band(ws, "Sources", f"Acme Corp · {len(figures)} figure rows")
    header(ws, 4, ["id", "label", "value", "unit", "disposition", "expression", "inputs", "root source", "To reperform", "population", "caveats and carry-forwards", "from check"],
           ["id_ledger", "description", "amount", "status", "status", "note", "id", "id", "note", "note", "note", "status"])
    r = 4
    for fid, (label, value, check) in figures.items():
        r += 1
        ident(ws.cell(row=r, column=2), fid)
        text(ws.cell(row=r, column=3), label)
        amount(ws.cell(row=r, column=4), value)
        text(ws.cell(row=r, column=5), "USD")
        text(ws.cell(row=r, column=6), "measured")
        text(ws.cell(row=r, column=7), f"as read ({label})")
        ident(ws.cell(row=r, column=8), "E.q1.pl")
        ident(ws.cell(row=r, column=9), "E.q1.pl")
        text(ws.cell(row=r, column=10), "pl.xlsx · IS · rows 5-40 · col E · sum")
        text(ws.cell(row=r, column=11), "whole")
        text(ws.cell(row=r, column=12), "")
        text(ws.cell(row=r, column=13), check)
    finish(ws, r, ledger=True)

    ws = wb.create_sheet("Evidence")
    band(ws, "Evidence", "Acme Corp · 1 room read")
    header(ws, 4, ["id", "kind", "file", "source", "role", "sheet", "anchor", "rows", "columns", "filter", "row count", "control total", "value basis", "read date", "note"],
           ["id_ledger", "status", "description", "status", "status", "status", "status", "status", "note", "note", "amount", "amount", "status", "period", "note"])
    ident(ws.cell(row=5, column=2), "E.q1.pl")
    for j, v in enumerate(["range", "statements/pl.xlsx", "pl", "management_prepared", "IS", "B4", "5-40", "B:E", ""], start=3):
        text(ws.cell(row=5, column=j), v)
    amount(ws.cell(row=5, column=12), 36, fmt="#,##0")
    amount(ws.cell(row=5, column=13), 41_250_000)
    text(ws.cell(row=5, column=14), "as_stated")
    text(ws.cell(row=5, column=15), "2 Sep 2025")
    text(ws.cell(row=5, column=16), "income statement as presented")
    finish(ws, 5, ledger=True)

    staging = rd / "out" / ".staging"
    staging.mkdir(parents=True)
    (rd / "workpapers").mkdir()
    wb.save(staging / "workbook.xlsx")
    by_check: dict[str, list[str]] = {}
    for fid, (label, value, check) in figures.items():
        by_check.setdefault(check, []).append(
            f"- id: {fid}\n  label: {label}\n  value: {value}\n  unit: USD\n  expression: as read\n"
            f"  inputs:\n    - {{role: pl, source_type: room_file, citation_id: E.q1.pl}}\n")
    for check, lines in by_check.items():
        (rd / "workpapers" / f"figures-{check}.yaml").write_text("".join(lines))
    (rd / "workpapers" / "evidence-plan.yaml").write_text(
        "- id: E.q1.pl\n  kind: range\n  file: \"statements/pl.xlsx\"\n  control_total: 41250000\n  row_count: 36\n")
    recipe = rd / "QOE.md"
    recipe.write_text(RECIPE)
    (rd / "run.json").write_text(json.dumps({
        "schema": "countz-accounting/run@1", "run_id": "qoe-acme-corp.20250903-101500",
        "goal": "quality-of-earnings", "degraded": False,
        "inputs": {"run_dir": str(rd), "skill": "qoe", "company": "Acme Corp", "params": {}},
        "checks": [{"id": "q1_fy2023", "kind": "tieout", "params": {"family": "q1"}},
                   {"id": "q6_bridge", "kind": "analysis", "params": {"family": "q6"}}],
        "plan": {"recipe": str(recipe)}}))


# The recipe the fixture run pins: its `## Report` declares the walk the deck must carry
# (RECIPE_FORMAT.md § Report), which case 12 holds the gate to.
RECIPE = """---
name: quality-of-earnings
objective: o
headline: q6
lead: [q6]
---
# r

## Exec summary

x

## Report

```json
{"metrics": {"title": "Adjusted EBITDA"},
 "schedules": [
  {"title": "EBITDA walk", "from": "q6", "columns": ["line", "verdict"], "periods": "all",
   "where": {"verdict": "supported"}, "through": "= Diligence adjusted EBITDA", "dense": true}
]}
```

## What the plan notes rather than checks

y
"""

# The spec as REPORT.md § 2 has it: a title is a headline — the subject of a page of
# figures (`EBITDA bridge`), the conclusion of a page that argues one (`The walk foots at
# every rung`) — and the sentence, where a page needs one, is `message`.
GOOD_SPEC = """schema: countz-accounting/report@1
title: Quality of earnings review
sections:
  - title: Executive summary
    pages:
      - title: Executive summary
        message: "{Exec Summary!B3}"
        blocks:
          - stats:
              - {label: Reported EBITDA · LTM Jul 2025, value: "{q6 | = Reported EBITDA | LTM Jul 2025 | $}"}
              - {label: Diligence adjusted EBITDA · LTM Jul 2025, value: "{q6 | = Diligence adjusted EBITDA | LTM Jul 2025 | $}"}
          - heading: What this rests on
          - bullets: ["The FY2023 income statement ties to the trial balance with no exception (q1)."]
      - title: Adjusted EBITDA
        blocks:
          - table: {from: q6, rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025], title: "EBITDA bridge, USD"}
          - chart: {type: column, from: q6, rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [FY2023, FY2024, LTM Jul 2025]}
  - title: The bridge
    pages:
      - title: EBITDA walk
        blocks:
          - table: {from: q6, where: {verdict: supported}, through: "= Diligence adjusted EBITDA", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict], dense: true}
      - title: The walk foots at every rung
        blocks:
          - table: {from: q6, max_rows: 5}
          - table: {from: q6, block: Exceptions}
  - title: Appendix
    pages:
      - title: Coverage
        message: "Two checks ran; the capex bridge was dropped because the room carries no capitalized-cost accounts."
        blocks:
          - text: "Every rostered check is listed with what it examined and what it did not."
          - text: "No owner-compensation adjustment was taken: the amount is {q6 | Owner compensation, no adjustment taken | LTM Jul 2025 | $}."
          - table: {from: Coverage, columns: ["token · title", status, "what was examined", "what was not examined, and why"]}
          - lines: {from: Basis of Preparation, block: Procedures not performed}
"""
SENTENCE_TITLE = ("Reported EBITDA walks to diligence adjusted EBITDA through supported adjustments "
                  "at every rung of the bridge.")


def stored_slides(deck: pathlib.Path) -> list[str]:
    """The slide XML parts of a deck, in presentation order (the gate's own reading)."""
    sys.path.insert(0, str(HERE))
    from check_report import slide_parts  # noqa: E402
    with zipfile.ZipFile(deck) as z:
        return [z.read(p).decode("utf-8", "replace") for p in slide_parts(z)]


def structure(deck: pathlib.Path) -> list[str]:
    """What case 1 holds on the deck as stored: headline and message are separate
    named shapes, a page written with no message carries none, and every table reads
    its first column left and the rest right."""
    out: list[str] = []
    slides = stored_slides(deck)
    pages = [x for x in slides if 'name="page:' in x]
    with_msg = [x for x in pages if 'name="message"' in x]
    if len(pages) < 4:
        out.append(f"structure: expected four pages, found {len(pages)}")
    if len(with_msg) != 2:
        out.append(f"structure: two pages carry a message shape, found {len(with_msg)}")
    if not any('name="title"' in x and 'name="message"' not in x for x in pages):
        out.append("structure: a page written without a message must carry no message shape")
    titles = [html.unescape(m) for x in pages for m in re.findall(r'name="title".*?<a:t>(.*?)</a:t>', x, re.S)]
    if "Executive summary" not in titles or "Adjusted EBITDA" not in titles:
        out.append(f"structure: titles are the headlines as written, found {titles}")
    tables = 0
    for x in pages:
        for tbl in re.findall(r"<a:tbl>.*?</a:tbl>", x, re.S):
            rows = re.findall(r"<a:tr\b.*?</a:tr>", tbl, re.S)
            for tr in rows:
                cells = re.findall(r"<a:tc\b.*?</a:tc>", tr, re.S)
                if len(cells) < 2:
                    continue
                tables += 1
                algn = [(re.search(r'algn="(\w+)"', c) or [None, None])[1] for c in cells]
                if algn[0] != "l" or any(a != "r" for a in algn[1:]):
                    out.append(f"structure: a table row reads {algn}; the first column is left, the rest right")
                    break
    if not tables:
        out.append("structure: no table row with two or more columns was checked")
    return out


def run(cmd: list[str]) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True)


def main() -> int:
    fails: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        rd = pathlib.Path(td) / "run"
        make_run(rd)
        py = sys.executable
        wb = rd / "out" / ".staging" / "workbook.xlsx"
        r = run([py, str(HERE / "link_workbook.py"), str(wb)])
        if r.returncode != 0:
            fails.append(f"fixture: link_workbook exited {r.returncode}: {(r.stdout + r.stderr)[-200:]}")
        r = run([py, str(HERE / "check_workbook.py"), str(wb), "--run-dir", str(rd)])
        if r.returncode != 0:
            fails.append(f"fixture: the workbook must pass check_workbook.py (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[-400:]}")
        spec = rd / "out" / ".staging" / "report.yaml"
        deck = rd / "out" / ".staging" / "report.pptx"
        build = [py, str(HERE / "build_report.py"), str(rd)]
        gate = [py, str(HERE / "check_report.py"), str(deck), "--run-dir", str(rd)]

        # 1. good spec: builds and passes
        spec.write_text(GOOD_SPEC)
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"1: a good spec must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-400:]}")
        else:
            g = run(gate)
            if g.returncode != 0:
                fails.append(f"1: the built deck must pass check_report.py (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[-400:]}")
            fails.extend(f"1: {s}" for s in structure(deck))
            good = deck.read_bytes()

            # 4. a figure edited after the build
            tampered = deck.with_name("tampered.pptx")
            hits = 0
            with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(tampered, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("ppt/slides/slide") and b'name="page:' in data \
                            and not hits:
                        # edit a figure inside a text run, not the slide's name attribute
                        edited = re.sub(rb"(<a:t>[^<]*?)8,404,000", rb"\g<1>8,404,100", data, count=1)
                        if edited != data:
                            data, hits = edited, hits + 1
                    zout.writestr(item, data)
            g = run([py, str(HERE / "check_report.py"), str(tampered), "--workbook", str(wb), "--run-dir", str(rd)])
            if not hits or g.returncode != 1 or "8,404,100" not in g.stdout:
                fails.append(f"4: an edited figure must be refused, named (exit {g.returncode}, edited {hits}): "
                             f"{(g.stdout + g.stderr).strip()[:300]}")
            deck.write_bytes(good)

            # 4a. the chart is drawn, and its values are still audited. A chart part
            # rasterises on import in some viewers, so the builder draws the chart as
            # shapes and carries each plotted value on its shape's name; the audit then
            # rests on that name, and this case is what stops it being lost silently.
            with zipfile.ZipFile(deck) as z:
                parts = z.namelist()
            if any(n.startswith("ppt/charts/") for n in parts):
                fails.append("4a: the deck holds a chart part; a chart is drawn as shapes")
            drawn = deck.with_name("drawn.pptx")
            hits = 0
            with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(drawn, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename.startswith("ppt/slides/slide") and not hits:
                        edited = re.sub(rb'name="chartval:([^:"]+):[-\d.eE+]+"',
                                        rb'name="chartval:\g<1>:4242.5"', data, count=1)
                        if edited != data:
                            data, hits = edited, hits + 1
                    zout.writestr(item, data)
            g = run([py, str(HERE / "check_report.py"), str(drawn), "--workbook", str(wb), "--run-dir", str(rd)])
            if not hits:
                fails.append("4a: no chartval shape on the deck — a drawn chart carries its values "
                             "on its shape names, and the gate audits them from there")
            elif g.returncode != 1 or "4242.5" not in g.stdout:
                fails.append(f"4a: an edited chart value must be refused, named (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[:300]}")

        # 2. a typed figure in a sentence: builds, gate refuses it by token
        spec.write_text(GOOD_SPEC.replace(
            '"The FY2023 income statement ties to the trial balance with no exception (q1)."',
            '"The FY2023 income statement ties to the trial balance within $1,234,567 (q1)."'))
        r = run(build)
        g = run(gate)
        if r.returncode != 0 or g.returncode != 1 or "$1,234,567" not in g.stdout:
            fails.append(f"2: a typed figure must build and be refused by the gate, named (build {r.returncode}, "
                         f"gate {g.returncode}): {(g.stdout + g.stderr).strip()[:300]}")

        # 3. a reference nothing resolves
        spec.write_text(GOOD_SPEC.replace("{q6 | = Reported EBITDA | LTM Jul 2025 | $}", "{q6 | = Reported EBITDA | LTM Aug 2025}", 1))
        r = run(build)
        if r.returncode != 1 or "LTM Aug 2025" not in r.stdout:
            fails.append(f"3: an unresolved reference must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 5. a sentence where a headline belongs: over the cap, and ending in a full stop
        spec.write_text(GOOD_SPEC.replace("title: Adjusted EBITDA", f'title: "{SENTENCE_TITLE}"'))
        r = run(build)
        if r.returncode != 1 or "pages[1].title" not in r.stdout or "headline" not in r.stdout:
            fails.append(f"5a: a sentence title must be refused as not a headline, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace("title: Adjusted EBITDA", "title: The bridge foots."))
        r = run(build)
        if r.returncode != 1 or "pages[1].title" not in r.stdout or "full stop" not in r.stdout:
            fails.append(f"5b: a title ending in a full stop must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

        # 7. the cover: the title names the work, not the client and not the period
        spec.write_text(GOOD_SPEC.replace("title: Quality of earnings review",
                                          "title: Quality of earnings review, Acme Corp FY2023"))
        r = run(build)
        if r.returncode != 1 or "title:" not in r.stdout or "company" not in r.stdout:
            fails.append(f"7a: a cover title naming the company must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace("title: Quality of earnings review",
                                          "title: Quality of earnings review for FY2023"))
        r = run(build)
        if r.returncode != 1 or "period" not in r.stdout:
            fails.append(f"7b: a cover title carrying the period must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")
        spec.write_text(GOOD_SPEC.replace(
            "title: Quality of earnings review",
            "title: Quality of earnings review\nsubtitle: Acme Corp · FY2023, as of 31 July 2025"))
        r = run(build)
        if r.returncode != 1 or "subtitle:" not in r.stdout:
            fails.append(f"7c: a cover subtitle repeating the company must be refused, named "
                         f"(exit {r.returncode}): {(r.stdout + r.stderr).strip()[:300]}")
        # the gate refuses it on a deck the builder never saw: the cover title, edited after the build
        spec.write_text(GOOD_SPEC)
        r = run(build)
        edited = deck.with_name("cover.pptx")
        hits = 0
        with zipfile.ZipFile(deck) as zin, zipfile.ZipFile(edited, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("ppt/slides/slide") and b'name="cover-title"' in data and not hits:
                    e = data.replace(b"<a:t>Quality of earnings review</a:t>",
                                     b"<a:t>Quality of earnings review, Acme Corp FY2023</a:t>", 1)
                    if e != data:
                        data, hits = e, hits + 1
                zout.writestr(item, data)
        g = run([py, str(HERE / "check_report.py"), str(edited), "--workbook", str(wb), "--run-dir", str(rd)])
        if not hits or g.returncode != 1 or "cover title" not in g.stdout:
            fails.append(f"7d: the gate must refuse a cover title naming the company or the period "
                         f"(exit {g.returncode}, edited {hits}): {(g.stdout + g.stderr).strip()[:300]}")

        # 8. a zero in a sentence, and 9. a schedule condensed past its own arithmetic
        spec.write_text(GOOD_SPEC)
        r = run(build)
        if r.returncode == 0 and deck.is_file():
            text_ = "\n".join(stored_slides(deck))
            if "the amount is $0." not in html.unescape(text_).replace("</a:t><a:t>", ""):
                fails.append("8: a zero in a sentence must read `$0` — the en dash is the table's zero")
        spec.write_text(GOOD_SPEC.replace(
            'rows: ["= Reported EBITDA", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025]',
            'rows: ["EBIT", "= Diligence adjusted EBITDA"], columns: [line, FY2023, FY2024, LTM Jul 2025]'))
        r = run(build)
        if r.returncode != 1 or "do not foot" not in r.stdout or "Non-recurring" not in r.stdout:
            fails.append(f"9: a schedule condensed past its own arithmetic must be refused, the dropped "
                         f"rows named (exit {r.returncode}): {(r.stdout + r.stderr).strip()[:400]}")

        # 10. the deck's number conventions
        spec.write_text(GOOD_SPEC.replace(
            "columns: [line, FY2023, FY2024, LTM Jul 2025], title:",
            "columns: [line, FY2023, FY2024, LTM Jul 2025], scale: thousands, title:"))
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"10: a scaled schedule must build (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[-300:]}")
        else:
            body = html.unescape("\n".join(stored_slides(deck))).replace("</a:t><a:t>", "")
            for want in ("$8.4m", "($'000)", "8,404"):
                if want not in body:
                    fails.append(f"10: the deck must carry `{want}` — REPORT.md § 4")
            g = run(gate)
            if g.returncode != 0:
                fails.append(f"10: a scaled deck must pass the gate (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[-300:]}")

        # 11. a dense schedule, filtered and ended
        spec.write_text(GOOD_SPEC)
        r = run(build)
        if r.returncode != 0 or not deck.is_file():
            fails.append(f"11: the good spec must build (exit {r.returncode})")
        else:
            body = html.unescape("\n".join(stored_slides(deck))).replace("</a:t><a:t>", "")
            if 'sz="825"' not in body:
                fails.append("11: a table declared `dense: true` is set at 8.25pt (`sz=\"825\"`)")
            walk = next((x for x in stored_slides(deck) if "EBITDA walk" in x and "table:q6" in x), "")
            if "Management adjusted EBITDA (information)" in walk:
                fails.append("11: `through:` ends the walk at the closing line — the information "
                             "line under it must not be on the walk's page")
            if "Owner compensation, no adjustment taken" not in walk or "+ D&amp;A" not in walk:
                fails.append("11: `where:` keeps the walk's mechanics and every ruled row")

        # 12. the recipe's schedule is held by the gate
        spec.write_text(GOOD_SPEC.replace(
            'where: {verdict: supported}, through: "= Diligence adjusted EBITDA", columns: [line, FY2023, FY2024, LTM Jul 2025, verdict], dense: true',
            'max_rows: 3, columns: [line, FY2023, FY2024, LTM Jul 2025, verdict]'))
        r = run(build)
        if r.returncode != 0:
            fails.append(f"12: a trimmed walk must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            g = run(gate)
            if g.returncode != 1 or "schedule `EBITDA walk`" not in g.stdout or "Non-recurring" not in g.stdout:
                fails.append(f"12: a walk trimmed to max_rows must be refused naming the schedule and the "
                             f"rows it lacks (exit {g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")
        spec.write_text(re.sub(r"      - title: EBITDA walk\n        blocks:\n          - table: .*\n", "", GOOD_SPEC)
                        .replace("          - table: {from: q6, max_rows: 5}\n", "")
                        .replace("          - table: {from: q6, block: Exceptions}\n", "          - text: \"No exception stands.\"\n")
                        .replace("          - table: {from: q6, rows: [\"= Reported EBITDA\", \"= Diligence adjusted EBITDA\"], "
                                 "columns: [line, FY2023, FY2024, LTM Jul 2025], title: \"EBITDA bridge, USD\"}\n", ""))
        r = run(build)
        if r.returncode != 0:
            fails.append(f"12: a deck without the walk must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
        else:
            g = run(gate)
            if g.returncode != 1 or "no page carries a table from `q6" not in g.stdout:
                fails.append(f"12: a deck with no table from the walk's tab must be refused naming the "
                             f"schedule (exit {g.returncode}): {(g.stdout + g.stderr).strip()[:400]}")

        # 13. the opening is held by the gate (REPORT.md § 1): the executive summary page
        #     first, with a message and a figure block; the key-metrics page headed as the
        #     recipe declares; the first schedule at most one page after it.
        opening = {
            "a": (GOOD_SPEC.replace("      - title: Executive summary\n", "      - title: Diligence-adjusted EBITDA\n", 1),
                  "headed `Executive summary`"),
            "b": (GOOD_SPEC.replace(
                '          - stats:\n'
                '              - {label: Reported EBITDA · LTM Jul 2025, value: "{q6 | = Reported EBITDA | LTM Jul 2025 | $}"}\n'
                '              - {label: Diligence adjusted EBITDA · LTM Jul 2025, value: "{q6 | = Diligence adjusted EBITDA | LTM Jul 2025 | $}"}\n',
                ''), "no stat tile, table or chart"),
            "c": (GOOD_SPEC.replace("      - title: Adjusted EBITDA\n", "      - title: EBITDA\n", 1),
                  "key-metrics page is headed `Adjusted EBITDA`"),
            "d": (GOOD_SPEC.replace(
                "      - title: EBITDA walk\n",
                '      - title: Context first\n        blocks:\n          - text: "We set the scene here."\n'
                '      - title: Context second\n        blocks:\n          - text: "We set more of the scene here."\n'
                "      - title: EBITDA walk\n", 1), "2 pages after the key-metrics page"),
        }
        for case, (text_, want) in opening.items():
            assert text_ != GOOD_SPEC, f"13{case}: the case did not change the spec"
            spec.write_text(text_)
            r = run(build)
            if r.returncode != 0:
                fails.append(f"13{case}: the spec must build (exit {r.returncode}): {(r.stdout + r.stderr).strip()[-300:]}")
                continue
            g = run(gate)
            if g.returncode != 1 or want not in g.stdout:
                fails.append(f"13{case}: the gate must refuse the opening naming `{want}` (exit {g.returncode}): "
                             f"{(g.stdout + g.stderr).strip()[:400]}")

        # 6. a fragment where a sentence belongs: a text block with no full stop
        spec.write_text(GOOD_SPEC.replace(
            '- text: "Every rostered check is listed with what it examined and what it did not."',
            '- text: "Every rostered check, what it examined and what it did not"'))
        r = run(build)
        if r.returncode != 1 or "blocks[0]" not in r.stdout or "full stop" not in r.stdout:
            fails.append(f"6: a fragment text block must be refused, named (exit {r.returncode}): "
                         f"{(r.stdout + r.stderr).strip()[:300]}")

    if fails:
        print(f"report-selftest: {len(fails)} failure(s)")
        for f in fails:
            print(f"  {f}")
        return 1
    print("report-selftest: ok — a spec-conformant workbook builds a deck that passes, headlines and "
          "messages as separate shapes and tables first-column-left, rest-right; a typed figure, an "
          "unresolved reference, an edited figure, a sentence title, a fragment, a cover title "
          "carrying the company or the period and a schedule condensed past its own arithmetic "
          "are refused; a zero in a sentence reads $0, a dollar figure reads $8.4m and a "
          "declared schedule reads in thousands; a dense, filtered walk builds, and the recipe's "
          "schedule trimmed or absent is refused; the opening — the executive summary with its message "
          "and a figure block, the key-metrics page headed as the recipe declares, the first schedule "
          "at most one page after it — is held.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
